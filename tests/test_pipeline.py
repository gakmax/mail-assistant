import datetime
import json
import poplib
import sqlite3
import time
import subprocess
import sys
import tempfile
from mail_assistant import money, notify
import unittest
import types
from contextlib import contextmanager
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import call, patch, MagicMock

from mail_assistant.core import (ANALYZING, FAILED, HANDLED, HEADERS, NO_RETRY,
                                 PRIORITY_ORDER, ROOM_MARK, Store,
                                 account_key, STATES, STATE_SQL, state_where, wait_cutoff, KST,
                                 SKIPPED, parse_mail, row_view, state_of,
                                 WAITING, PROGRESS, day_bounds, filter_rows,
                                 local_text, parse_mail, row_view, sql_text, state_of,
                                 thread_key, message_ids, address_of, display_name,
                                 safe_name, NAME_LIMIT, attachments_of, attachment_bytes,
                                 table_text, text_of_html, workbook_rows)
from mail_assistant.excel import (Excel, ExcelUpdateError, append_missing, ensure_table,
                                  error_detail, first_column, repair_generated_header,
                                  row_text)
from mail_assistant.calendar_sheet import (SHEET, Entry, cell_text, collect, month_grid,
                                           next_month, overdue, parse_day, sheet_events,
                                           signature, update_calendar)
from mail_assistant.dashboard import (PRIORITIES, SHEET as DASHBOARD, UPCOMING_ROWS,
                                      UPCOMING_TOP, blocks, describe, update_dashboard, upcoming)
from mail_assistant.excel import address_of, link_to_draft, link_to_mail, mail_rows, mailto
from mail_assistant.overview import (briefing_input, failures, overview, past_due,
                                     waiting_replies)
from mail_assistant.rules import (AD_REASON, AUTO_REASON, BULK_HEADERS, BULK_REASON,
                                  LETTER_REASON, skip_bulk, skip_reason, skipped_text)
from mail_assistant.settings import (GRADES, field_errors, model_choices, model_label,
                                     model_rows, model_traits, normalize, read_models)
from mail_assistant.services import check_connection, connection_steps, login_state
from mail_assistant.style import STYLE_VERSION, URGENT, apply_style
from mail_assistant.services import (BODY_LIMIT, DRAFT_BODY_LIMIT, DRAFT_TONES,
                                      DRAFT_WAYS, EFFORT_READ, EFFORT_THINK,
                                      Unanalyzable,
                                      TRANSLATE_LIMIT, THREAD_TURNS, MONEY_RULES,
                                      ANALYSIS_RULES, analyze, briefing,
                                      thread_context, context_size, chat_reply,
                                      chat_schema, draft, draft_schema, fetch_mail,
                                      read_password, save_password, translate,
                                      translate_schema, BATCH_CHARS, BATCH_MAILS,
                                      analyze_many, group_mails, prepare, squeeze_body)


CONFIG = {'host': 'pop3s.hiworks.com', 'port': 995, 'email': 'test@example.com', 'model': ''}
RESULT = {'category': '업무 요청', 'summary': '견적 회신 요청', 'requests': '견적 확인',
          'events': [{'title': '회신', 'start': '', 'deadline': '2026-09-11', 'evidence': '9월 11일까지', 'needs_review': False}],
          'priority': '높음', 'priority_reason': '회신 마감 명시', 'next_action': '견적 검토',
          'reply_needed': True, 'reply_subject': 'Re: 견적', 'reply_draft': '확인 후 회신드리겠습니다.',
          'money': [{'kind': '견적', 'amount': '1234000', 'currency': 'KRW',
                     'label': '9월 견적', 'evidence': '공급가액 1,234,000원',
                     'needs_review': False}],
          'order_no': 'A26090135'}


def mail():
    msg = EmailMessage()
    msg['Subject'] = '견적 요청'
    msg['From'] = 'sender@example.com'
    msg['Date'] = 'Thu, 10 Sep 2026 10:00:00 +0900'
    msg.set_content('9월 11일까지 회신 부탁드립니다.')
    return msg.as_bytes()


class FakePOP:
    messages = {'old': mail()}
    def __init__(self, *args, **kwargs):
        self.deleted = False
    def user(self, user): pass
    def pass_(self, password): pass
    def quit(self): pass
    def uidl(self):
        return b'+OK', [f'{i} {uid}'.encode() for i, uid in enumerate(self.messages, 1)], 0
    def retr(self, number):
        return b'+OK', list(self.messages.values())[number - 1].splitlines(), 0
    def dele(self, number):
        raise AssertionError('mail must never be deleted')


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.temp.name) / 'mail.db')
        FakePOP.messages = {'old': mail()}

    def tearDown(self):
        self.store.db.close()
        self.temp.cleanup()

    def test_baseline_new_mail_and_restart_dedup(self):
        fetch_mail(CONFIG, self.store, 'secret', FakePOP)
        self.assertEqual(self.store.counts(account_key(CONFIG))['total'], 0)
        FakePOP.messages['new'] = mail()
        fetch_mail(CONFIG, self.store, 'secret', FakePOP)
        self.store.db.close()
        self.store = Store(Path(self.temp.name) / 'mail.db')
        fetch_mail(CONFIG, self.store, 'secret', FakePOP)
        self.assertEqual(self.store.counts(account_key(CONFIG))['total'], 1)

    def test_empty_mailbox_baseline_and_account_isolation(self):
        FakePOP.messages = {}
        fetch_mail(CONFIG, self.store, 'secret', FakePOP)
        FakePOP.messages = {'first': mail()}
        fetch_mail(CONFIG, self.store, 'secret', FakePOP)
        other = {**CONFIG, 'email': 'other@example.com'}
        fetch_mail(other, self.store, 'secret', FakePOP)
        self.assertEqual(self.store.counts(account_key(CONFIG))['total'], 1)
        self.assertEqual(self.store.counts(account_key(other))['total'], 0)

    def test_retr_failure_is_not_marked_seen(self):
        fetch_mail(CONFIG, self.store, 'secret', FakePOP)
        FakePOP.messages['new'] = mail()
        with patch.object(FakePOP, 'retr', side_effect=OSError):
            with self.assertRaises(OSError):
                fetch_mail(CONFIG, self.store, 'secret', FakePOP)
        self.assertNotIn('new', self.store.seen(account_key(CONFIG)))

    def test_result_survives_failed_export(self):
        ident = self.store.add(account_key(CONFIG), 'new', mail())
        self.store.analyzed(ident, parse_mail(mail()), RESULT)
        rows = self.store.unexported(account_key(CONFIG))
        self.assertEqual(len(rows), 1)
        self.assertEqual(workbook_rows(rows[0])['일정'][0][0], ident + ':0')
        self.store.exported([ident])
        self.assertEqual(len(self.store.unexported(account_key(CONFIG))), 0)

    def test_backoff(self):
        ident = self.store.add(account_key(CONFIG), 'new', mail())
        self.store.failed(ident, 'error', 100)
        self.assertEqual(len(self.store.pending(account_key(CONFIG), 99)), 0)
        self.assertEqual(len(self.store.pending(account_key(CONFIG), 100)), 1)

    def test_html_and_attachments(self):
        msg = EmailMessage()
        msg['Subject'] = '한글 제목'
        msg.set_content('<p>본문</p><script>숨김</script>', subtype='html')
        msg.add_attachment(b'abc', maintype='application', subtype='pdf', filename='견적.pdf')
        data = parse_mail(msg.as_bytes())
        self.assertIn('본문', data['body'])
        self.assertNotIn('숨김', data['body'])
        self.assertEqual(data['attachments'], ['견적.pdf'])

    def test_table_keeps_its_cells_apart(self):
        """Hiworks sends 수주 mail as a table; the figures must not run together."""
        msg = EmailMessage()
        msg['Subject'] = '수주채번'
        msg.set_content('<div>완료했습니다.</div>'
                        '<table><tr><th>수주번호</th><th>공급가액</th><th>세액</th></tr>'
                        '<tr><td>A26090135</td><td>90,000</td><td>9,000</td></tr></table>',
                        subtype='html')
        body = parse_mail(msg.as_bytes())['body']
        self.assertNotIn('A2609013590,000', body)
        self.assertIn('| 수주번호 | 공급가액 | 세액 |', body)
        self.assertIn('| A26090135 | 90,000 | 9,000 |', body)
        # The separator row is what makes the column a cell belongs to readable.
        self.assertIn('| --- | --- | --- |', body)

    def test_layout_table_is_not_drawn_as_one(self):
        """Mail HTML wraps bodies and signatures in tables; a grid round those is noise."""
        text = text_of_html('<table><tr><td><p>안녕하세요</p><p>김규림 드림</p></td></tr></table>')
        self.assertNotIn('---', text)
        self.assertEqual(text.split(), ['안녕하세요', '김규림', '드림'])
        self.assertIn('\n', text.strip())      # the cell keeps its own line breaks

    def test_table_inside_a_layout_wrapper_survives(self):
        text = text_of_html('<table><tr><td>앞말'
                            '<table><tr><th>항목</th><th>금액</th></tr>'
                            '<tr><td>총액</td><td>99,000</td></tr></table>'
                            '뒷말</td></tr></table>')
        self.assertIn('| 항목 | 금액 |', text)
        self.assertIn('| 총액 | 99,000 |', text)

    def test_ragged_and_unclosed_table(self):
        """A short row is padded and a table the mail never closed is still emitted."""
        text = text_of_html('<table><tr><td>a|b</td><td>줄1<br>줄2</td></tr><tr><td>x</td>')
        self.assertIn(r'| a\|b | 줄1 줄2 |', text)     # a pipe in a cell cannot open one
        self.assertIn('| x |  |', text)

    def test_empty_table_draws_nothing(self):
        self.assertEqual(table_text([[['\xa0'], ['  ']]]), '')

    def test_codex_stdin_schema_and_no_api_key(self):
        def execute(command, **kwargs):
            self.assertNotIn('OPENAI_API_KEY', kwargs['env'])
            self.assertNotIn('CODEX_API_KEY', kwargs['env'])
            self.assertNotIn('9월 11일', ' '.join(command))
            self.assertIn('9월 11일', kwargs['input'])
            self.assertIn('--ignore-user-config', command)
            Path(command[command.index('-o') + 1]).write_text(json.dumps(RESULT), encoding='utf-8')
            return subprocess.CompletedProcess(command, 0)
        ident = self.store.add(account_key(CONFIG), 'new', mail())
        row = self.store.pending(account_key(CONFIG), 0)[0]
        with patch('mail_assistant.services.codex_command', return_value=['codex']), patch('mail_assistant.services.subprocess.run', side_effect=execute), patch.dict('os.environ', {'OPENAI_API_KEY': 'secret', 'CODEX_API_KEY': 'secret'}):
            _, result = analyze(row, CONFIG)
        self.assertEqual(result, RESULT)

    def test_invalid_codex_output_not_accepted(self):
        def execute(command, **kwargs):
            Path(command[command.index('-o') + 1]).write_text('{}', encoding='utf-8')
            return subprocess.CompletedProcess(command, 0)
        self.store.add(account_key(CONFIG), 'new', mail())
        row = self.store.pending(account_key(CONFIG), 0)[0]
        from jsonschema import ValidationError
        with patch('mail_assistant.services.codex_command', return_value=['codex']), patch('mail_assistant.services.subprocess.run', side_effect=execute):
            with self.assertRaises(ValidationError):
                analyze(row, CONFIG)


OLD_SCHEMA = '''
    CREATE TABLE meta (key TEXT PRIMARY KEY, value TEXT NOT NULL);
    CREATE TABLE seen (account TEXT, uid TEXT, PRIMARY KEY(account,uid));
    CREATE TABLE mail (
        id TEXT PRIMARY KEY, account TEXT NOT NULL, uid TEXT NOT NULL,
        raw BLOB NOT NULL, received TEXT NOT NULL, parsed TEXT,
        result TEXT, exported INTEGER NOT NULL DEFAULT 0,
        attempts INTEGER NOT NULL DEFAULT 0, retry_at REAL NOT NULL DEFAULT 0,
        error TEXT NOT NULL DEFAULT '', UNIQUE(account,uid));
'''


class MigrationTests(unittest.TestCase):
    def old_database(self, folder):
        """A database written by the installed version, with one analysed mail."""
        path = Path(folder) / 'mail.db'
        db = sqlite3.connect(path)
        db.executescript(OLD_SCHEMA)
        db.execute('INSERT INTO mail(id,account,uid,raw,received,parsed,result,exported) VALUES (?,?,?,?,?,?,?,1)',
                   ('old-1', 'acct', 'uid-1', mail(), '2026-09-10T01:00:00+00:00',
                    json.dumps({'sender': 'sender@example.com', 'subject': '견적 요청', 'attachments': []}),
                    json.dumps(RESULT)))
        db.execute("INSERT INTO meta VALUES ('baseline:acct', '2026-09-10T00:00:00+00:00')")
        db.commit()
        db.close()
        return path

    def test_rows_survive_and_new_columns_appear(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.old_database(folder)
            store = Store(path)
            columns = {row['name'] for row in store.db.execute('PRAGMA table_info(mail)')}
            self.assertLessEqual({'handled', 'draft_edit', 'notified', 'analyzed_at', 'exported_at',
                                  'subject', 'sender', 'analyzing', 'todo_hidden'}, columns)
            row = store.detail('old-1')
            self.assertEqual(row['handled'], '')
            self.assertEqual(store.get_meta('baseline:acct'), '2026-09-10T00:00:00+00:00')
            # subject/sender are backfilled from the analysis that was already stored
            self.assertEqual((row['subject'], row['sender']), ('견적 요청', 'sender@example.com'))
            store.db.close()

    def test_opening_twice_is_safe(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.old_database(folder)
            Store(path).db.close()
            store = Store(path)
            self.assertEqual(store.counts('acct')['total'], 1)
            store.db.close()


class StoreViewTests(unittest.TestCase):
    def store(self, folder):
        store = Store(Path(folder) / 'mail.db')
        first = store.add('acct', 'uid-1', mail())
        second = store.add('acct', 'uid-2', mail())
        store.analyzed(first, {'sender': 'kim@example.com', 'subject': '견적 요청', 'attachments': []}, RESULT)
        return store, first, second

    def test_add_stores_the_headers_for_the_list(self):
        with tempfile.TemporaryDirectory() as folder:
            store, first, _ = self.store(folder)
            view = row_view(store.page('acct')[-1])
            self.assertEqual(view['subject'], '견적 요청')
            self.assertEqual(view['sender'], 'sender@example.com')
            store.db.close()

    def test_state_reflects_analysis_and_handling(self):
        with tempfile.TemporaryDirectory() as folder:
            store, first, second = self.store(folder)
            states = {row['id']: state_of(row) for row in store.page('acct')}
            self.assertEqual(states[first], '미처리')
            self.assertEqual(states[second], '분석 대기')
            store.failed(second, '분석 실패', 0)
            self.assertEqual(state_of(store.detail(second)), '1회 실패')
            store.set_handled(first, HANDLED)
            self.assertEqual(state_of(store.detail(first)), HANDLED)
            self.assertEqual(store.open_count('acct'), 0)
            store.db.close()

    def test_filter_by_text_and_state(self):
        views = [{'subject': '견적 요청', 'sender': 'kim@a.com', 'state': '미처리'},
                 {'subject': 'Weekly report', 'sender': 'admin@b.com', 'state': '처리'}]
        self.assertEqual(len(filter_rows(views, query='견적')), 1)
        self.assertEqual(len(filter_rows(views, query='WEEKLY')), 1)      # case insensitive
        self.assertEqual(len(filter_rows(views, query='@b.com')), 1)      # sender too
        self.assertEqual(len(filter_rows(views, state='처리')), 1)
        self.assertEqual(len(filter_rows(views, query='견적', state='처리')), 0)
        self.assertEqual(len(filter_rows(views)), 2)

    def test_filter_reaches_failed_mail_whatever_the_retry_count(self):
        views = [{'subject': 'a', 'sender': 'a@a', 'state': '1회 실패'},
                 {'subject': 'b', 'sender': 'b@b', 'state': '4회 실패'},
                 {'subject': 'c', 'sender': 'c@c', 'state': '분석 대기'}]
        self.assertEqual([view['subject'] for view in filter_rows(views, state=FAILED)], ['a', 'b'])
        self.assertEqual([view['subject'] for view in filter_rows(views, state='분석 대기')], ['c'])
        self.assertEqual(len(filter_rows(views, state='없는 상태')), 0)

    def test_reset_clears_backoff_and_optionally_the_result(self):
        with tempfile.TemporaryDirectory() as folder:
            store, first, second = self.store(folder)
            store.failed(second, '분석 실패', time.time() + 3600)
            store.reset([second])
            row = store.detail(second)
            self.assertEqual((row['attempts'], row['retry_at'], row['error']), (0, 0, ''))
            self.assertEqual([r['id'] for r in store.pending('acct', time.time())], [second])
            store.reset([first], reanalyze=True)
            self.assertIsNone(store.detail(first)['result'])
            self.assertEqual({r['id'] for r in store.retryable('acct')}, {first, second})
            store.db.close()

    def test_draft_edit_and_notified_flags(self):
        with tempfile.TemporaryDirectory() as folder:
            store, first, _ = self.store(folder)
            store.set_draft(first, '사람이 고친 초안')
            self.assertEqual(store.detail(first)['draft_edit'], '사람이 고친 초안')
            self.assertEqual([r['id'] for r in store.unnotified('acct')], [first])
            store.mark_notified([first])
            self.assertEqual(store.unnotified('acct'), [])
            store.db.close()


class ThreadKeyTests(unittest.TestCase):
    """대화를 잇는 규칙. References의 첫 항목이 뿌리라는 것이 전부다."""

    def test_the_root_of_the_references_chain_is_the_key(self):
        parsed = {'message_id': '<c@x>', 'references': '<a@x> <b@x>', 'in_reply_to': '<b@x>'}
        self.assertEqual(thread_key(parsed), '<a@x>')

    def test_the_first_mail_of_a_thread_keys_itself(self):
        """뿌리 메일의 Message-ID가 바로 자식들의 References[0]이다 — 그래서 만난다."""
        root = {'message_id': '<a@x>', 'references': '', 'in_reply_to': ''}
        child = {'message_id': '<b@x>', 'references': '<a@x>', 'in_reply_to': '<a@x>'}
        self.assertEqual(thread_key(root), thread_key(child))

    def test_order_does_not_matter(self):
        """POP3는 순서를 약속하지 않고, 기준점 때문에 앞부분을 못 볼 수도 있다."""
        child = {'message_id': '<b@x>', 'references': '<a@x>'}
        grandchild = {'message_id': '<c@x>', 'references': '<a@x> <b@x>'}
        self.assertEqual(thread_key(grandchild), thread_key(child))

    def test_in_reply_to_alone_asks_what_we_already_know(self):
        """References 없이 In-Reply-To만 보내는 클라이언트가 있다 — 그때만 조회한다."""
        parsed = {'message_id': '<c@x>', 'references': '', 'in_reply_to': '<b@x>'}
        self.assertEqual(thread_key(parsed, lookup=lambda ref: '<a@x>'), '<a@x>')
        # 아는 부모가 없으면 부모의 id가 키다. 부모가 나중에 오면 그때 만난다.
        self.assertEqual(thread_key(parsed, lookup=lambda ref: ''), '<b@x>')

    def test_a_mail_with_no_headers_at_all_gets_nothing_here(self):
        """add()가 제 메일 id를 준다. ''로 두면 헤더 없는 메일끼리 한 대화가 된다."""
        self.assertEqual(thread_key({}), '')

    def test_only_bracketed_ids_are_read(self):
        self.assertEqual(message_ids('<a@x> junk <b@x>'), ['<a@x>', '<b@x>'])
        self.assertEqual(message_ids(''), [])


class ThreadStoreTests(unittest.TestCase):
    """대화가 데이터베이스에서 실제로 묶이는가."""

    ACCOUNT = 'acct'

    def mail(self, subject, ident, refs='', parent=''):
        message = EmailMessage()
        message['From'] = 'kim@buyer.example'
        message['Subject'] = subject
        message['Message-ID'] = ident
        if refs:
            message['References'] = refs
        if parent:
            message['In-Reply-To'] = parent
        message.set_content('본문')
        return message.as_bytes()

    def test_a_reply_lands_in_the_same_conversation_as_its_original(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                root = store.add(self.ACCOUNT, 'u1', self.mail('견적 요청', '<a@x>'))
                reply = store.add(self.ACCOUNT, 'u2',
                                  self.mail('Re: 견적 요청', '<b@x>', refs='<a@x>', parent='<a@x>'))
                other = store.add(self.ACCOUNT, 'u3', self.mail('무관한 메일', '<z@x>'))
                rows = store.thread_rows(self.ACCOUNT, store.detail(reply)['thread'])
                self.assertEqual([row['id'] for row in rows], [root, reply])
                self.assertEqual(
                    [row['id'] for row in store.thread_rows(
                        self.ACCOUNT, store.detail(other)['thread'])], [other])
            finally:
                store.db.close()

    def test_the_reply_may_arrive_first(self):
        """POP3 순서를 믿지 않는다는 것이 References[0] 규칙을 고른 이유다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                reply = store.add(self.ACCOUNT, 'u2',
                                  self.mail('Re: 견적', '<b@x>', refs='<a@x>'))
                root = store.add(self.ACCOUNT, 'u1', self.mail('견적', '<a@x>'))
                rows = store.thread_rows(self.ACCOUNT, store.detail(root)['thread'])
                self.assertEqual({row['id'] for row in rows}, {root, reply})
            finally:
                store.db.close()

    def test_older_mail_is_threaded_by_the_backfill(self):
        """컬럼이 생기기 전에 수집된 메일도 같은 대화가 되어야 한다 — raw가 남아 있다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                root = store.add(self.ACCOUNT, 'u1', self.mail('견적', '<a@x>'))
                reply = store.add(self.ACCOUNT, 'u2',
                                  self.mail('Re: 견적', '<b@x>', refs='<a@x>'))
                with store.db:
                    store.db.execute("UPDATE mail SET thread='', message_id=''")
                store.backfill_threads()
                rows = store.thread_rows(self.ACCOUNT, store.detail(reply)['thread'])
                self.assertEqual([row['id'] for row in rows], [root, reply])
                left = store.db.execute("SELECT COUNT(*) FROM mail WHERE thread=''").fetchone()[0]
                self.assertEqual(left, 0)
            finally:
                store.db.close()

    ANALYSIS = {'category': '견적·계약', 'summary': '9월 견적 검토 요청',
                'requests': '단가 회신', 'events': [], 'priority': '높음',
                'priority_reason': '', 'next_action': '', 'reply_needed': True,
                'reply_subject': '', 'reply_draft': '', 'money': [], 'order_no': ''}

    def threaded_pair(self, store):
        """분석까지 끝난 원 메일과, 그 답장. 답장의 id를 돌려준다."""
        root = store.add(self.ACCOUNT, 'u1', self.mail('견적 요청', '<a@x>'))
        store.analyzed(root, {'sender': '', 'subject': '', 'body': '비밀 본문',
                              'attachments': []}, dict(self.ANALYSIS))
        reply = store.add(self.ACCOUNT, 'u2',
                          self.mail('Re: 견적 요청', '<b@x>', refs='<a@x>'))
        return root, reply

    def test_the_analysis_context_is_summaries_and_never_a_body(self):
        """분석에 실려 가는 것은 analyze()가 이미 값을 치른 답이지 본문이 아니다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                root, reply = self.threaded_pair(store)
                row = store.detail(reply)
                turns = thread_context(store.thread_before(
                    self.ACCOUNT, row['thread'], reply))
                self.assertEqual(len(turns), 1)
                self.assertEqual(turns[0]['summary'], '9월 견적 검토 요청')
                self.assertNotIn('비밀 본문', json.dumps(turns, ensure_ascii=False))
                # 자기 자신은 맥락이 아니다 — 이제 시각이 아니라 제 행이 그것을 자른다.
                self.assertEqual(
                    store.thread_before(self.ACCOUNT, row['thread'], root), [])
            finally:
                store.db.close()

    def test_a_thread_collected_in_one_cycle_keeps_its_context(self):
        """한 번의 수집이 원 메일과 답장을 같이 가져오는 것은 흔한 일이고, 그때 둘의
        `received`는 같은 문자열이다 — Windows 시계는 15ms쯤마다 한 번 움직인다.

        `received<?` 하나로 자르면 원 메일이 제 대화에서 빠지고, 답장은 맥락 없이
        분석된다. 조용히: 맥락이 빈 것은 첫 메일도 마찬가지이기 때문이다. 시계를
        묶어 두면 Windows 밖에서도 그대로 재현되고, CI의 Windows 다리에서 실제로
        이렇게 무너졌다.
        """
        with tempfile.TemporaryDirectory() as folder:
            with patch('mail_assistant.core.now',
                       return_value='2026-09-11T01:00:00+00:00'):
                store = Store(Path(folder) / 'mail.db')
                try:
                    root, reply = self.threaded_pair(store)
                    rows = [store.detail(root), store.detail(reply)]
                    self.assertEqual(rows[0]['received'], rows[1]['received'])
                    turns = thread_context(store.thread_before(
                        self.ACCOUNT, rows[1]['thread'], reply))
                    self.assertEqual([turn['summary'] for turn in turns],
                                     ['9월 견적 검토 요청'])
                    # 그리고 그 한 틱 안에서도 순서는 한 방향이다.
                    self.assertEqual(
                        store.thread_before(self.ACCOUNT, rows[0]['thread'], root), [])
                finally:
                    store.db.close()


class MoneyTests(unittest.TestCase):
    """금액 — 이 앱에서 유일하게 합계를 그리는 기능이고, 규칙은 전부 '세지 않는 쪽'이다.

    다른 답이 틀리면 사람이 읽다가 알아본다. 합계는 틀렸다는 것을 스스로 말하지 않는다.
    """

    def item(self, amount='1234000', currency='KRW', kind='견적', review=False, **extra):
        return {'kind': kind, 'amount': amount, 'currency': currency, 'label': '9월 견적',
                'evidence': '공급가액 1,234,000원', 'needs_review': review, **extra}

    def row(self, ident='m1', items=(), received='2026-09-14T00:00:00+00:00', order_no=''):
        return {'id': ident, 'subject': '견적', 'sender': 'kim@x', 'received': received,
                'handled': '', 'result': json.dumps({'money': list(items),
                                                     'order_no': order_no})}

    def test_a_number_is_a_number_and_everything_else_is_not(self):
        """None과 0은 다르다. 0은 '영 원'이고 None은 '나는 이것을 읽지 못했다'이다."""
        self.assertEqual(money.money_value('1,234,000'), 1234000)
        self.assertEqual(money.money_value(' 90000 '), 90000)
        self.assertEqual(money.money_value('12.50'), 12.5)
        self.assertEqual(money.money_value('0'), 0)
        for guess in ('약 90,000', '90,000~100,000', '9만', '1,234,000원', '', None, '-500'):
            self.assertIsNone(money.money_value(guess), guess)

    def test_two_currencies_are_never_added(self):
        """$5,000과 ₩5,000,000을 더한 5,005,000은 숫자가 아니라 사고다."""
        found = money.entries([self.row('m1', [self.item('5000', 'USD'),
                                               self.item('5000000', 'KRW')])])
        book = money.totals(found)
        self.assertEqual(book['sums']['USD']['견적'], 5000)
        self.assertEqual(book['sums']['KRW']['견적'], 5000000)
        self.assertEqual(set(book['sums']), {'USD', 'KRW'})

    def test_an_unreadable_amount_is_dropped_not_counted_as_zero(self):
        """못 읽은 것을 0으로 세면 합계는 조용히 작아지고, 작아진 합계는 말이 없다."""
        found = money.entries([self.row('m1', [self.item('1000'),
                                               self.item('약 90,000')])])
        book = money.totals(found)
        self.assertEqual(book['sums']['KRW']['견적'], 1000)
        self.assertEqual(book['counted'], 1)
        self.assertEqual(book['skipped'], 1)

    def test_a_model_that_is_unsure_is_taken_at_its_word(self):
        found = money.entries([self.row('m1', [self.item('1000', review=True)])])
        book = money.totals(found)
        self.assertEqual(book['sums'], {})
        self.assertEqual(book['skipped'], 1)

    def test_an_unknown_currency_never_joins_a_total(self):
        """무엇인지 모르는 돈끼리 더한 수는 아무 질문에도 답하지 않는다."""
        found = money.entries([self.row('m1', [self.item('1000', '기타'),
                                               self.item('1000', 'ZZZ')])])
        book = money.totals(found)
        self.assertEqual(book['sums'], {})
        self.assertEqual(book['skipped'], 2)

    def test_the_total_cannot_be_read_without_the_number_it_left_out(self):
        """합계와 제외 건수를 한 dict에 담는 것이 이 모듈의 요점이다."""
        book = money.totals(money.entries([self.row('m1', [self.item('1000'),
                                                           self.item('bad')])]))
        self.assertIn('skipped', book)
        self.assertEqual(money.skipped_text(book), '확인 필요 1건은 합계에서 뺐어요')
        clean = money.totals(money.entries([self.row('m1', [self.item('1000')])]))
        self.assertEqual(money.skipped_text(clean), '')

    def test_a_mail_with_no_money_contributes_nothing(self):
        self.assertEqual(money.entries([self.row('m1', [])]), [])
        self.assertEqual(money.entries([{'id': 'm', 'subject': '', 'sender': '',
                                         'received': '', 'handled': '', 'result': None}]), [])
        self.assertEqual(money.entries([{'id': 'm', 'subject': '', 'sender': '',
                                         'received': '', 'handled': '',
                                         'result': '{not json'}]), [])

    def test_the_order_number_falls_through_from_the_mail(self):
        found = money.entries([self.row('m1', [self.item()], order_no='A26090135')])
        self.assertEqual(found[0]['order_no'], 'A26090135')

    def test_months_and_filters_narrow_the_same_list(self):
        rows = [self.row('m1', [self.item('100')], received='2026-09-14T00:00:00+00:00'),
                self.row('m2', [self.item('200', kind='입금')],
                         received='2026-08-14T00:00:00+00:00')]
        found = money.entries(rows)
        self.assertEqual(money.months(found), ['2026-09', '2026-08'])
        self.assertEqual(len(money.in_month(found, '2026-09')), 1)
        self.assertEqual(len(money.in_month(found, '')), 2)
        self.assertEqual(len(money.by_kind(found, '입금')), 1)
        self.assertEqual(money.month_title('2026-09'), '2026년 9월')
        self.assertEqual(money.month_title(''), '전체 기간')

    def test_the_currency_shown_is_the_one_most_of_the_mail_used(self):
        found = money.entries([self.row('m1', [self.item('1', 'USD'), self.item('2', 'USD'),
                                               self.item('3', 'KRW')])])
        self.assertEqual(money.main_currency(found), 'USD')
        self.assertEqual(money.main_currency([]), 'KRW')

    def test_a_person_can_correct_what_the_model_read(self):
        """고칠 수 없으면 합계는 있어서는 안 되는 기능이다 — 틀린 채로 영영 선다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                message = EmailMessage()
                message['From'] = 'kim@x.example'
                message['Subject'] = '견적'
                message.set_content('본문')
                ident = store.add('acct', 'u1', message.as_bytes())
                store.analyzed(ident, {'sender': '', 'subject': '', 'body': '',
                                       'attachments': []},
                               {**RESULT, 'money': [self.item('약 90,000', review=True)]})
                before = money.totals(money.entries(store.page('acct')))
                self.assertEqual(before['sums'], {})
                self.assertTrue(store.set_money(ident, 0, '900000', 'KRW'))
                after = money.totals(money.entries(store.page('acct')))
                self.assertEqual(after['sums']['KRW']['견적'], 900000)
                self.assertEqual(after['skipped'], 0)
                entry = money.entries(store.page('acct'))[0]
                self.assertTrue(entry['edited'])
                # 근거는 그대로 남는다: 사람이 고쳤어도 원문이 무엇이었는지가 판단의 근거다.
                self.assertEqual(entry['evidence'], '공급가액 1,234,000원')
            finally:
                store.db.close()

    def test_a_correction_that_names_nothing_real_changes_nothing(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                message = EmailMessage()
                message['From'] = 'kim@x.example'
                message['Subject'] = '견적'
                message.set_content('본문')
                ident = store.add('acct', 'u1', message.as_bytes())
                store.analyzed(ident, {'sender': '', 'subject': '', 'body': '',
                                       'attachments': []},
                               {**RESULT, 'money': [self.item()]})
                self.assertFalse(store.set_money(ident, 5, '1', 'KRW'))
                self.assertFalse(store.set_money('nosuch', 0, '1', 'KRW'))
            finally:
                store.db.close()

    def test_다시_분석은_고친_금액도_함께_지운다(self):
        """money는 result 안에 산다 — 답이 사라지면 고친 값도 같이 사라져야 한다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                message = EmailMessage()
                message['From'] = 'kim@x.example'
                message['Subject'] = '견적'
                message.set_content('본문')
                ident = store.add('acct', 'u1', message.as_bytes())
                store.analyzed(ident, {'sender': '', 'subject': '', 'body': '',
                                       'attachments': []},
                               {**RESULT, 'money': [self.item()]})
                store.set_money(ident, 0, '999', 'KRW')
                store.reset([ident], reanalyze=True)
                self.assertEqual(money.entries(store.page('acct')), [])
            finally:
                store.db.close()

    def test_the_prompt_forbids_arithmetic_and_guessing(self):
        """지어낸 숫자가 합계에 들어가는 것이 이 기능의 유일한 진짜 사고다."""
        self.assertIn('적혀 있는', MONEY_RULES)
        self.assertIn('빈 배열', MONEY_RULES)
        self.assertIn('곱하지', MONEY_RULES)
        self.assertIn('needs_review=true', MONEY_RULES)
        self.assertIn(MONEY_RULES, ANALYSIS_RULES)


class NotifyTests(unittest.TestCase):
    """창 밖 알림 — 무엇을 알리고, 무엇을 알리지 않고, 실패하면 어떻게 되는가."""

    ACCOUNT = 'acct'

    def row(self, ident, priority, subject='제목', handled='', action='다음 행동'):
        return {'id': ident, 'handled': handled, 'subject': subject,
                'result': json.dumps({'priority': priority, 'next_action': action})}

    def test_only_긴급과_높음이_창_밖으로_나간다(self):
        """'보통'까지 알리면 알림은 '메일이 왔다'와 같아지고, 그것은 메일함이 할 말이다."""
        rows = [self.row('a', '긴급'), self.row('b', '높음'),
                self.row('c', '보통'), self.row('d', '낮음')]
        self.assertEqual([item['id'] for item in notify.worth_telling(rows)], ['a', 'b'])

    def test_a_mail_already_finished_is_not_news(self):
        """알림이 도착하기 전에 읽고 완료한 메일까지 알리면, 방금 한 일을 모르는 알림이다."""
        self.assertEqual(notify.worth_telling([self.row('a', '긴급', handled=HANDLED)]), [])

    def test_긴급이_먼저_선다(self):
        rows = [self.row('a', '높음'), self.row('b', '긴급')]
        self.assertEqual([item['id'] for item in notify.worth_telling(rows)], ['b', 'a'])

    def test_one_mail_names_itself_and_several_are_counted(self):
        one = notify.worth_telling([self.row('a', '긴급', subject='서버 점검')])
        title, body = notify.summarise(one)
        self.assertEqual(title, '긴급 메일: 서버 점검')
        self.assertEqual(body, '다음 행동')
        many = notify.worth_telling([self.row('a', '긴급', subject='서버 점검'),
                                     self.row('b', '긴급'), self.row('c', '높음')])
        title, body = notify.summarise(many)
        self.assertEqual(title, '긴급 메일 3건')
        self.assertIn('외 2건', body)
        self.assertEqual(notify.summarise([]), (None, None))

    def test_a_broken_result_is_skipped_not_raised(self):
        self.assertEqual(notify.worth_telling(
            [{'id': 'a', 'handled': '', 'subject': 's', 'result': '{not json'}]), [])

    def test_the_setting_defaults_to_on_for_a_config_that_predates_it(self):
        self.assertTrue(notify.enabled({}))
        self.assertTrue(notify.enabled({'notify': '1'}))
        self.assertFalse(notify.enabled({'notify': ''}))
        self.assertFalse(notify.enabled({'notify': '0'}))

    def test_a_pc_that_cannot_toast_says_so_rather_than_raising(self):
        """알림이 안 뜨는 것은 불편이고, 알림 때문에 수집이 멈추는 것은 고장이다."""
        with patch.dict('sys.modules', {'win32gui': None, 'win32con': None}):
            self.assertFalse(notify.send('제목', '본문'))

    def test_the_balloon_goes_out_through_shell_notifyicon(self):
        calls = []
        fake_gui = types.SimpleNamespace(
            GetModuleHandle=lambda _: 1,
            WNDCLASS=lambda: types.SimpleNamespace(),
            RegisterClass=lambda klass: 7,
            CreateWindow=lambda *args: 99,
            LoadIcon=lambda *args: 5,
            ExtractIconEx=lambda *args: ([], []),
            DestroyIcon=lambda handle: None,
            Shell_NotifyIcon=lambda action, data: calls.append((action, data)))
        fake_con = types.SimpleNamespace(WS_OVERLAPPED=0, IDI_APPLICATION=32512,
                                         IMAGE_ICON=1, LR_LOADFROMFILE=16, LR_DEFAULTSIZE=64)
        notify._tray.clear()
        try:
            with patch.dict('sys.modules', {'win32gui': fake_gui, 'win32con': fake_con}):
                self.assertTrue(notify.send('긴급 메일 2건', '외 1건'))
                # 두 번째 알림은 창을 다시 만들지 않는다: 등록이 두 번 되면 실패하고,
                # 창을 지우면 풍선도 같이 사라진다.
                self.assertTrue(notify.send('또', '하나'))
            self.assertEqual([action for action, _ in calls],
                             [notify.NIM_ADD, notify.NIM_MODIFY, notify.NIM_MODIFY])
            self.assertIn('긴급 메일 2건', calls[1][1])
        finally:
            notify._tray.clear()

    def test_tell_marks_every_mail_it_looked_at_even_the_quiet_ones(self):
        """다음 주기에 세 시간 전의 긴급 메일을 처음인 양 띄우는 것은 잔소리다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                seen = []
                for uid, priority in (('u1', '긴급'), ('u2', '낮음')):
                    message = EmailMessage()
                    message['From'] = 'kim@x.example'
                    message['Subject'] = f'{uid} 제목'
                    message.set_content('본문')
                    ident = store.add(self.ACCOUNT, uid, message.as_bytes())
                    store.analyzed(ident, {'sender': '', 'subject': '', 'body': '',
                                           'attachments': []},
                                   {'category': '문의', 'summary': '', 'requests': '',
                                    'events': [], 'priority': priority, 'priority_reason': '',
                                    'next_action': '', 'reply_needed': False,
                                    'reply_subject': '', 'reply_draft': '', 'money': [], 'order_no': ''})
                    seen.append(ident)
                with patch.object(notify, 'send', lambda *args, **kwargs: True):
                    self.assertEqual(notify.tell(store, self.ACCOUNT, {'notify': '1'}), 1)
                    # 두 번째 주기에는 알릴 것이 없다.
                    self.assertEqual(notify.tell(store, self.ACCOUNT, {'notify': '1'}), 0)
                self.assertEqual(store.unnotified(self.ACCOUNT), [])
            finally:
                store.db.close()

    def test_a_switched_off_notification_still_marks_the_mail(self):
        """끈 동안 쌓인 것을 다시 켠 날 한꺼번에 띄우면, 그것은 알림이 아니라 사고다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                message = EmailMessage()
                message['From'] = 'kim@x.example'
                message['Subject'] = '긴급'
                message.set_content('본문')
                ident = store.add(self.ACCOUNT, 'u1', message.as_bytes())
                store.analyzed(ident, {'sender': '', 'subject': '', 'body': '',
                                       'attachments': []},
                               {'category': '문의', 'summary': '', 'requests': '', 'events': [],
                                'priority': '긴급', 'priority_reason': '', 'next_action': '',
                                'reply_needed': False, 'reply_subject': '', 'reply_draft': '', 'money': [], 'order_no': ''})
                sent = []
                with patch.object(notify, 'send', lambda *a, **k: sent.append(a) or True):
                    self.assertEqual(notify.tell(store, self.ACCOUNT, {'notify': ''}), 0)
                self.assertEqual(sent, [])
                self.assertEqual(store.unnotified(self.ACCOUNT), [])
            finally:
                store.db.close()


class AttachmentTests(unittest.TestCase):
    """첨부 꺼내기 — 바이트는 처음부터 raw 안에 있었고, 이름만 화면에 있었다."""

    ACCOUNT = 'acct'

    def test_a_name_from_a_header_can_never_reach_a_directory(self):
        """첨부 이름은 이 앱에서 바깥 사람이 고르는 유일한 문자열이고, 파일시스템에 닿는다."""
        for hostile in ('../../../etc/passwd', r'..\..\windows\system32\x.dll',
                        'C:/Windows/x.dll', '..', '.', '/etc/shadow'):
            got = safe_name(hostile)
            self.assertNotIn('/', got, hostile)
            self.assertNotIn('\\', got, hostile)
            self.assertNotIn(':', got, hostile)
            self.assertFalse(got.startswith('.'), hostile)
            self.assertEqual(Path(got).name, got, hostile)

    def test_a_windows_device_name_is_moved_out_of_the_way(self):
        """CON.txt로 저장하면 확장자와 무관하게 파일이 생기지 않는다."""
        self.assertEqual(safe_name('CON.txt'), '_CON.txt')
        self.assertEqual(safe_name('nul'), '_nul')
        self.assertEqual(safe_name('content.txt'), 'content.txt')

    def test_an_unusable_name_becomes_a_usable_one_rather_than_an_error(self):
        """적대적인 이름을 붙인 메일도 읽는 사람이 열고 싶어 하는 메일이다."""
        self.assertEqual(safe_name(''), '첨부파일')
        self.assertEqual(safe_name('...'), '첨부파일')
        self.assertEqual(safe_name(None), '첨부파일')

    def test_a_long_name_keeps_its_extension(self):
        got = safe_name('가' * 300 + '.xlsx')
        self.assertTrue(got.endswith('.xlsx'))
        self.assertLessEqual(len(got), NAME_LIMIT)

    def mail(self, files):
        message = EmailMessage()
        message['From'] = 'kim@buyer.example'
        message['Subject'] = '견적 첨부'
        message.set_content('확인 바랍니다')
        for name, payload in files:
            message.add_attachment(payload, maintype='application',
                                   subtype='octet-stream', filename=name)
        return message.as_bytes()

    def test_the_list_carries_size_and_order_not_just_names(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                ident = store.add(self.ACCOUNT, 'u1',
                                  self.mail([('견적서.xlsx', b'x' * 82000),
                                             ('발주서.pdf', b'y' * 512)]))
                items = store.attachments(ident)
                self.assertEqual([item['name'] for item in items],
                                 ['견적서.xlsx', '발주서.pdf'])
                self.assertEqual([item['index'] for item in items], [0, 1])
                self.assertEqual([item['size'] for item in items], [82000, 512])
            finally:
                store.db.close()

    def test_an_attachment_is_written_under_its_own_mail(self):
        """두 메일의 '견적서.xlsx'가 서로를 덮어쓰지 않는 유일한 방법이다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                first = store.add(self.ACCOUNT, 'u1',
                                  self.mail([('견적서.xlsx', '첫번째'.encode())]))
                second = store.add(self.ACCOUNT, 'u2',
                                   self.mail([('견적서.xlsx', '두번째'.encode())]))
                out = Path(folder) / 'attachments'
                one = store.save_attachment(first, 0, out)
                two = store.save_attachment(second, 0, out)
                self.assertNotEqual(one, two)
                self.assertEqual(one.read_bytes(), '첫번째'.encode())
                self.assertEqual(two.read_bytes(), '두번째'.encode())
                self.assertEqual(one.parent.name, first)
            finally:
                store.db.close()

    def test_a_hostile_filename_is_written_inside_the_folder_it_was_given(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                ident = store.add(self.ACCOUNT, 'u1',
                                  self.mail([('../../../escaped.txt', b'no')]))
                out = (Path(folder) / 'attachments').resolve()
                path = store.save_attachment(ident, 0, out).resolve()
                self.assertTrue(str(path).startswith(str(out)), path)
            finally:
                store.db.close()

    def test_a_number_nobody_sent_is_not_a_file(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                ident = store.add(self.ACCOUNT, 'u1', self.mail([('a.txt', b'x')]))
                self.assertIsNone(store.save_attachment(ident, 7, Path(folder) / 'out'))
                self.assertIsNone(store.save_attachment('nosuchmail', 0, Path(folder) / 'out'))
            finally:
                store.db.close()

    def test_a_mail_with_no_attachment_says_so_quietly(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                ident = store.add(self.ACCOUNT, 'u1', self.mail([]))
                self.assertEqual(store.attachments(ident), [])
            finally:
                store.db.close()


class SenderTests(unittest.TestCase):
    """거래처 — 한 주소로 묶고, 세는 쪽과 여는 쪽이 같은 컬럼을 쓴다."""

    ACCOUNT = 'acct'
    TODAY = datetime.date(2026, 9, 15)

    def test_the_address_is_the_key_and_the_display_name_is_not(self):
        """표시 이름은 보내는 쪽 클라이언트가 이번 주에 쓰기로 한 것이다."""
        self.assertEqual(address_of('김과장 <Kim@Buyer.example>'), 'kim@buyer.example')
        self.assertEqual(address_of('plain@x.com'), 'plain@x.com')
        self.assertEqual(address_of('이름만 있고 주소 없음'), '')
        self.assertEqual(display_name('김과장 <kim@x>'), '김과장')
        self.assertEqual(display_name('"Kim, J" <kim@x>'), 'Kim, J')
        self.assertEqual(display_name('plain@x.com'), '')

    def store(self, folder, plan):
        """plan: (uid, From 헤더, 받은 날짜, 분석함, reply_needed, handled)"""
        store = Store(Path(folder) / 'mail.db')
        self.ids = {}
        for uid, sender, day, analysed, needed, handled in plan:
            message = EmailMessage()
            message['From'] = sender
            message['Subject'] = f'{uid} 제목'
            message['Message-ID'] = f'<{uid}@x>'
            message.set_content('본문')
            ident = store.add(self.ACCOUNT, uid, message.as_bytes())
            self.ids[uid] = ident
            if analysed:
                store.analyzed(ident, {'sender': sender, 'subject': '', 'body': '',
                                       'attachments': []},
                               {'category': '문의', 'summary': '', 'requests': '',
                                'events': [], 'priority': '보통', 'priority_reason': '',
                                'next_action': '', 'reply_needed': needed,
                                'reply_subject': '', 'reply_draft': '', 'money': [], 'order_no': ''})
            if handled:
                store.set_handled(ident, handled)
            with store.db:
                store.db.execute('UPDATE mail SET received=? WHERE id=?',
                                 (day_bounds(day)[0], ident))
        return store

    PLAN = [('a', '김과장 <kim@buyer.example>', datetime.date(2026, 9, 14), True, True, ''),
            ('b', 'KIM <Kim@Buyer.Example>', datetime.date(2026, 9, 5), True, True, ''),
            ('c', 'kim@buyer.example', datetime.date(2026, 9, 1), True, False, HANDLED),
            ('d', '이대리 <lee@corp.example>', datetime.date(2026, 9, 12), True, False, '')]

    def test_one_company_is_one_row_however_the_name_was_written(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, self.PLAN)
            try:
                rows = {row['addr']: row for row in store.senders(self.ACCOUNT, self.TODAY)}
                self.assertEqual(set(rows), {'kim@buyer.example', 'lee@corp.example'})
                kim = rows['kim@buyer.example']
                self.assertEqual(kim['total'], 3)
                self.assertEqual(kim['open'], 2)      # c는 처리 완료
                self.assertEqual(kim['done'], 1)
                # a는 하루 전이라 아직 대기가 아니고, b는 열흘 전이라 대기다.
                self.assertEqual(kim['waiting'], 1)
            finally:
                store.db.close()

    @contextmanager
    def frozen(self):
        """달력을 양쪽에 똑같이 물려 둔다.

        senders()는 today를 받고 search()는 받지 않는다 — state_where()가 제 기본값으로
        시계를 읽기 때문이다. 앱에서는 둘 다 진짜 오늘이라 어긋날 수 없지만, 테스트가
        한쪽만 TODAY로 얼리면 그 사이에 걸친 메일 하나가 하루가 지날 때마다 편을 바꾼다.
        이 테스트는 2026-09-15에는 통과하고 16일에는 실패했다 — 재던 것이 아니라 날짜를
        재고 있었던 것이다.
        """
        stamp = datetime.datetime.combine(self.TODAY, datetime.time(9, 0), tzinfo=KST)
        with patch('mail_assistant.core.local_now', return_value=stamp):
            yield

    def test_the_card_and_the_list_it_opens_count_the_same_mail(self):
        """카드는 GROUP BY로 세고 목록은 sender_addr로 거른다. 둘은 같아야 한다."""
        with tempfile.TemporaryDirectory() as folder, self.frozen():
            store = self.store(folder, self.PLAN)
            try:
                for row in store.senders(self.ACCOUNT, self.TODAY):
                    _, total = store.search(self.ACCOUNT, sender=row['addr'])
                    self.assertEqual(total, row['total'], row['addr'])
                    _, waiting = store.search(self.ACCOUNT, sender=row['addr'], state=WAITING)
                    self.assertEqual(waiting, row['waiting'], row['addr'])
                    _, open_count = store.search(self.ACCOUNT, sender=row['addr'], state='미처리')
                    self.assertEqual(open_count, row['open'], row['addr'])
            finally:
                store.db.close()

    def test_older_mail_is_grouped_by_the_backfill(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, self.PLAN)
            try:
                with store.db:
                    store.db.execute("UPDATE mail SET sender_addr=''")
                store.backfill_senders()
                self.assertEqual(len(store.senders(self.ACCOUNT, self.TODAY)), 2)
            finally:
                store.db.close()

    def test_the_newest_contact_comes_first(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, self.PLAN)
            try:
                rows = store.senders(self.ACCOUNT, self.TODAY)
                self.assertEqual([row['addr'] for row in rows],
                                 ['kim@buyer.example', 'lee@corp.example'])
            finally:
                store.db.close()


class WaitingReplyTests(unittest.TestCase):
    """답장 대기 — 분석은 답장이 필요하다 했는데 완료 표시 없이 며칠이 지난 메일.

    이 앱은 보낸 메일을 볼 수 없다. 그래서 이 기능이 말할 수 있는 것은 '완료 표시가
    없다'까지이고, 여기서 지키는 것은 그 판정을 내리는 자리가 하나뿐이라는 것이다 —
    카드는 waiting_replies()로 세고 그 카드가 여는 목록은 SQL로 거르므로, 둘이
    갈라지면 같은 질문에 두 답이 생긴다.
    """

    ACCOUNT = 'acct'
    TODAY = datetime.date(2026, 9, 15)

    def store(self, folder, plan):
        """plan: (uid, 받은 날짜, reply_needed, handled)"""
        store = Store(Path(folder) / 'mail.db')
        self.ids = {}
        for uid, day, needed, handled in plan:
            message = EmailMessage()
            message['From'] = f'{uid}@corp.example'
            message['Subject'] = f'{uid} 제목'
            message.set_content('본문')
            ident = store.add(self.ACCOUNT, uid, message.as_bytes())
            self.ids[uid] = ident
            store.analyzed(ident, {'sender': '', 'subject': '', 'body': '', 'attachments': []},
                           {'category': '문의', 'summary': '', 'requests': '', 'events': [],
                            'priority': '보통', 'priority_reason': '', 'next_action': '',
                            'reply_needed': needed, 'reply_subject': '', 'reply_draft': '', 'money': [], 'order_no': ''})
            if handled:
                store.set_handled(ident, handled)
            # `received` is written by now(); the age is the whole point, so it is set
            # to the Korean day the fixture asked for.
            with store.db:
                store.db.execute('UPDATE mail SET received=? WHERE id=?',
                                 (day_bounds(day)[0], ident))
        return store

    def waiting(self, store):
        rows = list(store.page(self.ACCOUNT))
        handled = {row['id'] for row in rows if row['handled'] == HANDLED}
        return waiting_replies(rows, self.TODAY, handled)

    def test_a_mail_younger_than_the_threshold_is_not_yet_late(self):
        plan = [('today', self.TODAY, True, ''),
                ('yesterday', self.TODAY - datetime.timedelta(days=1), True, ''),
                ('twodays', self.TODAY - datetime.timedelta(days=2), True, '')]
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, plan)
            try:
                self.assertEqual([row['id'] for row in self.waiting(store)],
                                 [self.ids['twodays']])
            finally:
                store.db.close()

    def test_the_list_stands_oldest_first(self):
        plan = [('old', self.TODAY - datetime.timedelta(days=9), True, ''),
                ('mid', self.TODAY - datetime.timedelta(days=4), True, ''),
                ('new', self.TODAY - datetime.timedelta(days=2), True, '')]
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, plan)
            try:
                rows = self.waiting(store)
                self.assertEqual([row['days'] for row in rows], [9, 4, 2])
            finally:
                store.db.close()

    def test_a_reply_nobody_asked_for_and_a_finished_mail_are_both_out(self):
        plan = [('needs', self.TODAY - datetime.timedelta(days=5), True, ''),
                ('no_reply', self.TODAY - datetime.timedelta(days=5), False, ''),
                ('done', self.TODAY - datetime.timedelta(days=5), True, HANDLED)]
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, plan)
            try:
                self.assertEqual([row['id'] for row in self.waiting(store)],
                                 [self.ids['needs']])
            finally:
                store.db.close()

    def test_진행_중은_빠지지_않고_그렇게_적힌다(self):
        """열흘째 진행 중인 메일이야말로 이 패널이 있는 이유다."""
        plan = [('moving', self.TODAY - datetime.timedelta(days=10), True, PROGRESS)]
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, plan)
            try:
                rows = self.waiting(store)
                self.assertEqual(len(rows), 1)
                self.assertTrue(rows[0]['progress'])
            finally:
                store.db.close()

    def test_the_card_and_the_list_it_opens_count_the_same_mail(self):
        """카드는 waiting_replies()로 세고 목록은 SQL로 거른다. 둘은 같아야 한다."""
        plan = [('a', self.TODAY - datetime.timedelta(days=9), True, ''),
                ('b', self.TODAY - datetime.timedelta(days=3), True, ''),
                ('c', self.TODAY - datetime.timedelta(days=1), True, ''),
                ('d', self.TODAY - datetime.timedelta(days=6), False, ''),
                ('e', self.TODAY - datetime.timedelta(days=6), True, HANDLED),
                ('f', self.TODAY - datetime.timedelta(days=6), True, PROGRESS)]
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, plan)
            try:
                counted = {row['id'] for row in self.waiting(store)}
                clause, params = state_where(WAITING, self.TODAY)
                listed = {row['id'] for row in store.db.execute(
                    f'SELECT id FROM mail WHERE account=? AND ({clause})',
                    (self.ACCOUNT, *params))}
                self.assertEqual(counted, listed)
                self.assertEqual(counted, {self.ids['a'], self.ids['b'], self.ids['f']})
            finally:
                store.db.close()

    def test_다시_분석은_판정까지_되돌린다(self):
        """result를 지우면서 reply_needed를 두고 가면, 없는 분석을 근거로 서 있게 된다."""
        plan = [('again', self.TODAY - datetime.timedelta(days=5), True, '')]
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, plan)
            try:
                self.assertEqual(len(self.waiting(store)), 1)
                store.reset([self.ids['again']], reanalyze=True)
                self.assertEqual(self.waiting(store), [])
                row = store.db.execute('SELECT reply_needed FROM mail WHERE id=?',
                                       (self.ids['again'],)).fetchone()
                self.assertEqual(row['reply_needed'], -1)
            finally:
                store.db.close()

    def test_older_mail_gets_its_verdict_backfilled(self):
        """컬럼이 생기기 전에 분석된 메일도 같은 답을 내야 한다."""
        plan = [('old', self.TODAY - datetime.timedelta(days=5), True, ''),
                ('quiet', self.TODAY - datetime.timedelta(days=5), False, '')]
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, plan)
            try:
                with store.db:
                    store.db.execute('UPDATE mail SET reply_needed=-1')
                store.backfill_replies()
                self.assertEqual([row['id'] for row in self.waiting(store)],
                                 [self.ids['old']])
                # 두 번째 호출이 집어 들 행이 없어야 한다: -1이 '아직 안 채움'의
                # 정확한 표시라는 것이 이 backfill이 끝을 아는 유일한 방법이다.
                left = store.db.execute('SELECT COUNT(*) FROM mail WHERE result IS NOT NULL '
                                        'AND reply_needed = -1').fetchone()[0]
                self.assertEqual(left, 0)
            finally:
                store.db.close()

    def test_the_fallback_window_filters_on_the_same_judgement(self):
        """app.py는 SQL이 아니라 filter_rows()로 거른다 — 드롭다운에 값만 생기고
        거르는 쪽이 모르면, 그 화면은 아무것도 없는 목록을 조용히 보여 준다."""
        # 세 줄 모두 self.TODAY 기준이고, 거르는 쪽 시계도 같은 날에 물려 둔다.
        # 'fresh'가 datetime.date.today()였을 때는 고정된 TODAY와 진짜 오늘과
        # filter_rows()가 읽는 시계, 셋이 서로 다른 날을 가리켰다 — 오늘과 내일은
        # 우연히 맞았고 일주일 뒤에는 fresh가 대기 쪽으로 넘어갔다.
        plan = [('late', self.TODAY - datetime.timedelta(days=6), True, ''),
                ('fresh', self.TODAY, True, ''),
                ('quiet', self.TODAY - datetime.timedelta(days=6), False, '')]
        stamp = datetime.datetime.combine(self.TODAY, datetime.time(9, 0), tzinfo=KST)
        with tempfile.TemporaryDirectory() as folder, \
                patch('mail_assistant.core.local_now', return_value=stamp):
            store = self.store(folder, plan)
            try:
                views = [row_view(row) for row in store.page(self.ACCOUNT)]
                shown = filter_rows(views, state=WAITING)
                self.assertEqual([view['id'] for view in shown], [self.ids['late']])
                # 전체는 그대로 전부다.
                self.assertEqual(len(filter_rows(views)), 3)
            finally:
                store.db.close()

    def test_the_briefing_is_told_what_is_waiting_and_how_long(self):
        """브리핑이 '오늘 무엇이 왔나' 말고 '내가 무엇을 붙잡고 있나'를 말할 수 있는 자리."""
        plan = [('a', self.TODAY - datetime.timedelta(days=8), True, ''),
                ('b', self.TODAY - datetime.timedelta(days=3), True, '')]
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder, plan)
            try:
                payload = briefing_input(list(store.page(self.ACCOUNT)), self.TODAY)
                self.assertEqual([row['days'] for row in payload['waiting_reply']], [8, 3])
                self.assertEqual(payload['counts']['답장 대기'], 2)
                self.assertEqual(payload['truncated']['waiting_reply'],
                                 {'shown': 2, 'total': 2})
            finally:
                store.db.close()

    def test_the_cutoff_is_a_calendar_day_not_48_hours(self):
        """'이틀 지났다'는 달력 두 장이지 48시간이 아니다 — 화면의 N일째와 같은 계산."""
        self.assertEqual(wait_cutoff(self.TODAY, days=2),
                         day_bounds(self.TODAY - datetime.timedelta(days=1))[0])
        self.assertEqual(wait_cutoff(self.TODAY, days=1), day_bounds(self.TODAY)[0])


class SearchTests(unittest.TestCase):
    """The list query. Filtering, sorting and paging all happen in SQL now."""

    ACCOUNT = 'acct'

    def store(self, folder):
        store = Store(Path(folder) / 'mail.db')
        plan = [
            # uid, subject, sender, body, category, priority, analysed, handled, attempts
            ('uid-1', '견적 검토 요청', 'kim@buyer.example', '9월 견적서를 검토해 주세요',
             '견적·계약', '긴급', True, '', 0),
            ('uid-2', 'Weekly report', 'admin@corp.example', '주간 보고 첨부합니다',
             '공지', '낮음', True, HANDLED, 0),
            ('uid-3', '회의 일정 조정', 'lee@corp.example', '수요일로 옮길까요',
             '회의·일정', '보통', True, '', 0),
            ('uid-4', '아직 분석 안 됨', 'new@corp.example', '본문', '', '', False, '', 0),
            ('uid-5', '분석 실패한 메일', 'bad@corp.example', '본문', '', '', False, '', 3),
        ]
        self.ids = {}
        for uid, subject, sender, body, category, priority, analysed, handled, attempts in plan:
            message = EmailMessage()
            message['From'] = sender
            message['Subject'] = subject
            message['Date'] = 'Fri, 11 Sep 2026 10:00:00 +0900'
            message.set_content(body)
            ident = store.add(self.ACCOUNT, uid, message.as_bytes())
            self.ids[uid] = ident
            if analysed:
                store.analyzed(ident, {'sender': sender, 'subject': subject, 'body': body,
                                       'attachments': []},
                               {'category': category, 'summary': f'{subject} 요약',
                                'requests': '회신 필요', 'events': [],
                                'priority': priority, 'priority_reason': '근거',
                                'next_action': '담당자 확인', 'reply_needed': False,
                                'reply_subject': '', 'reply_draft': '', 'money': [], 'order_no': ''})
            if handled:
                store.set_handled(ident, handled)
            if attempts:
                store.failed(ident, '분석 실패', 0)
                for _ in range(attempts - 1):
                    store.failed(ident, '분석 실패', 0)
        return store

    def subjects(self, store, **kwargs):
        rows, _ = store.search(self.ACCOUNT, **kwargs)
        return [row['subject'] for row in rows]

    def test_text_matches_subject_sender_body_and_summary(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            self.assertEqual(self.subjects(store, query='견적 검토'), ['견적 검토 요청'])
            self.assertEqual(self.subjects(store, query='WEEKLY'), ['Weekly report'])   # ascii case
            self.assertEqual(self.subjects(store, query='@buyer.example'), ['견적 검토 요청'])
            self.assertEqual(self.subjects(store, query='수요일로'), ['회의 일정 조정'])   # body
            self.assertEqual(self.subjects(store, query='회신 필요'),
                             ['회의 일정 조정', 'Weekly report', '견적 검토 요청'])         # requests
            store.db.close()

    def test_mail_stored_in_the_same_tick_still_has_one_order(self):
        """`received` comes from now(), and a Windows clock ticks about every 15ms.

        A cycle that stores several mails therefore writes the same string for all of
        them, and an ORDER BY with no tie-break may hand back a different order every
        call — which is how the list reshuffled itself between two refreshes.
        """
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.db.execute("UPDATE mail SET received='2026-09-11T01:00:00+00:00'")
            store.db.commit()
            first = self.subjects(store, query='회신 필요')
            self.assertEqual(first, ['회의 일정 조정', 'Weekly report', '견적 검토 요청'])
            for _ in range(5):
                self.assertEqual(self.subjects(store, query='회신 필요'), first)
            self.assertEqual([row['subject'] for row in store.page(self.ACCOUNT)][:3],
                             ['분석 실패한 메일', '아직 분석 안 됨', '회의 일정 조정'])
            store.db.close()

    def test_each_state_filter_returns_its_own_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            self.assertEqual(self.subjects(store, state=HANDLED), ['Weekly report'])
            self.assertEqual(self.subjects(store, state='미처리'),
                             ['회의 일정 조정', '견적 검토 요청'])
            self.assertEqual(self.subjects(store, state='분석 대기'), ['아직 분석 안 됨'])
            self.assertEqual(self.subjects(store, state=FAILED), ['분석 실패한 메일'])
            self.assertEqual(len(self.subjects(store, state='없는 상태')), 5)   # ignored
            store.db.close()

    def test_a_kind_and_an_urgency_narrow_the_list(self):
        """The two columns the 대시보드 bars are drawn from, as filters on the list."""
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            self.assertEqual(self.subjects(store, category='공지'), ['Weekly report'])
            self.assertEqual(self.subjects(store, priority='긴급'), ['견적 검토 요청'])
            # Together with each other and with the filters that were already there.
            self.assertEqual(self.subjects(store, category='공지', priority='긴급'), [])
            self.assertEqual(self.subjects(store, category='공지', state=HANDLED),
                             ['Weekly report'])
            # '' is 전체, and a value nothing carries is simply empty — never everything.
            self.assertEqual(len(self.subjects(store, category='')), 5)
            self.assertEqual(self.subjects(store, category='없는 종류'), [])
            store.db.close()

    def test_failures_are_reachable_which_they_were_not_before(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            rows, total = store.search(self.ACCOUNT, state=FAILED)
            self.assertEqual(total, 1)
            self.assertEqual(state_of(rows[0]), '3회 실패')
            store.db.close()

    def test_priority_sorts_by_urgency_not_by_spelling(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            rows, _ = store.search(self.ACCOUNT, state='미처리', sort='priority', desc=False)
            self.assertEqual([row['subject'] for row in rows],
                             ['견적 검토 요청', '회의 일정 조정'])          # 긴급 then 보통
            store.db.close()

    def test_state_sorts_by_progress(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            rows, _ = store.search(self.ACCOUNT, sort='state', desc=False)
            self.assertEqual(state_of(rows[0]), '분석 대기')
            self.assertEqual(state_of(rows[-1]), HANDLED)
            store.db.close()

    def test_sort_by_subject_and_an_unknown_key_falls_back(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            self.assertEqual(self.subjects(store, sort='subject', desc=False)[0], 'Weekly report')
            newest = self.subjects(store)
            self.assertEqual(self.subjects(store, sort='없는 컬럼'), newest)
            store.db.close()

    def test_paging_reports_the_total_not_the_page(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            first, total = store.search(self.ACCOUNT, limit=2, offset=0)
            second, again = store.search(self.ACCOUNT, limit=2, offset=2)
            self.assertEqual((total, again), (5, 5))
            self.assertEqual(len(first), 2)
            self.assertEqual(len(second), 2)
            self.assertFalse({row['id'] for row in first} & {row['id'] for row in second})
            store.db.close()

    def test_text_and_state_narrow_together(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            self.assertEqual(self.subjects(store, query='견적', state=HANDLED), [])
            self.assertEqual(self.subjects(store, query='견적', state='미처리'),
                             ['견적 검토 요청'])
            store.db.close()

    def test_the_verdict_columns_track_the_result(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            row = store.detail(self.ids['uid-1'])
            self.assertEqual((row['category'], row['priority']), ('견적·계약', '긴급'))
            store.reset([self.ids['uid-1']], reanalyze=True)
            row = store.detail(self.ids['uid-1'])
            self.assertEqual((row['category'], row['priority']), ('', ''))
            store.db.close()

    def test_migrate_backfills_the_verdicts_of_older_mail(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.db.execute("UPDATE mail SET category='', priority=''")
            store.db.commit()
            store.db.close()
            again = Store(Path(folder) / 'mail.db')          # migrate() runs on open
            row = again.detail(self.ids['uid-1'])
            self.assertEqual((row['category'], row['priority']), ('견적·계약', '긴급'))
            again.db.close()

    def test_a_broken_result_does_not_stop_the_backfill(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.db.execute("UPDATE mail SET category='', priority=''")
            store.db.execute("UPDATE mail SET result='{망가진' WHERE uid='uid-2'")
            store.db.commit()
            store.db.close()
            again = Store(Path(folder) / 'mail.db')
            self.assertEqual(again.detail(self.ids['uid-1'])['priority'], '긴급')
            self.assertEqual(again.detail(self.ids['uid-2'])['priority'], '')
            again.db.close()

    def test_handling_many_at_once(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.set_handled_many([self.ids['uid-1'], self.ids['uid-3']], HANDLED)
            self.assertEqual(store.search(self.ACCOUNT, state=HANDLED)[1], 3)
            store.set_handled_many([self.ids['uid-1']], '')
            self.assertEqual(store.search(self.ACCOUNT, state=HANDLED)[1], 2)
            store.db.close()

    def test_deleting_mail_takes_its_chat_and_leaves_the_rest(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.add_chat(self.ACCOUNT, 'user', '이 건 어떻게 할까요', self.ids['uid-1'])
            store.add_chat(self.ACCOUNT, 'user', '일반 상담')
            store.delete([self.ids['uid-1'], self.ids['uid-4']])
            self.assertEqual(self.subjects(store), ['분석 실패한 메일', '회의 일정 조정',
                                                    'Weekly report'])
            self.assertEqual(store.chat(self.ACCOUNT, self.ids['uid-1']), [])
            self.assertEqual(len(store.chat(self.ACCOUNT)), 1)
            store.db.close()

    def test_a_thread_is_listed_for_every_room_that_has_one(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.add_chat(self.ACCOUNT, 'user', '이 건 어떻게 할까요', self.ids['uid-1'])
            store.add_chat(self.ACCOUNT, 'codex', '이렇게 하세요', self.ids['uid-1'])
            store.add_chat(self.ACCOUNT, 'user', '일반 상담')
            rooms = {row['key']: row for row in store.rooms(self.ACCOUNT)}
            self.assertEqual(set(rooms), {'', self.ids['uid-1']})
            self.assertEqual(rooms[self.ids['uid-1']]['turns'], 2)
            # The mail thread is named after the mail, in one query for all of them.
            self.assertEqual(rooms[self.ids['uid-1']]['name'], '견적 검토 요청')
            self.assertEqual(rooms[self.ids['uid-1']]['role'], 'codex')
            self.assertEqual(rooms[self.ids['uid-1']]['last'], '이렇게 하세요')
            store.db.close()

    def test_the_general_thread_is_listed_even_when_it_is_empty(self):
        """It is where the page lands when nothing else is asked for."""
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            self.assertEqual([row['key'] for row in store.rooms(self.ACCOUNT)], [''])
            store.db.close()

    def test_rooms_come_back_newest_activity_first(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.add_chat(self.ACCOUNT, 'user', '먼저', self.ids['uid-1'])
            store.add_chat(self.ACCOUNT, 'user', '나중', self.ids['uid-2'])
            self.assertEqual([row['key'] for row in store.rooms(self.ACCOUNT)][:2],
                             [self.ids['uid-2'], self.ids['uid-1']])
            store.db.close()

    def test_a_free_room_shows_up_before_anyone_has_written_in_it(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            ident = store.new_room(self.ACCOUNT)
            rooms = {row['key']: row for row in store.rooms(self.ACCOUNT)}
            self.assertIn(ident, rooms)
            self.assertEqual((rooms[ident]['turns'], rooms[ident]['name']), (0, ''))
            store.db.close()

    def test_a_room_names_itself_once_and_a_rename_always_wins(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            ident = store.new_room(self.ACCOUNT)
            store.name_room(ident, '첫 질문', only_if_unnamed=True)
            # A second question must not rename the room out from under the first.
            store.name_room(ident, '둘째 질문', only_if_unnamed=True)
            rooms = {row['key']: row for row in store.rooms(self.ACCOUNT)}
            self.assertEqual(rooms[ident]['name'], '첫 질문')
            store.name_room(ident, '손으로 바꾼 이름')
            rooms = {row['key']: row for row in store.rooms(self.ACCOUNT)}
            self.assertEqual(rooms[ident]['name'], '손으로 바꾼 이름')
            store.db.close()

    def test_a_free_room_id_can_never_be_a_mail_id(self):
        """Mail ids are 24 hex characters, so the marker is what tells the two apart."""
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            ident = store.new_room(self.ACCOUNT)
            self.assertTrue(ident.startswith(ROOM_MARK))
            self.assertNotIn(ident, self.ids.values())
            store.db.close()

    def test_dropping_a_free_room_takes_its_turns_and_nothing_else(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            ident = store.new_room(self.ACCOUNT)
            store.add_chat(self.ACCOUNT, 'user', '자유 질문', ident)
            store.add_chat(self.ACCOUNT, 'user', '메일 질문', self.ids['uid-1'])
            self.assertTrue(store.drop_room(self.ACCOUNT, ident))
            self.assertEqual(store.chat(self.ACCOUNT, ident), [])
            self.assertEqual(len(store.chat(self.ACCOUNT, self.ids['uid-1'])), 1)
            self.assertNotIn(ident, [row['key'] for row in store.rooms(self.ACCOUNT)])
            store.db.close()

    def test_drop_room_refuses_a_mail_thread(self):
        """A mail's thread goes with the mail; this button must not be able to reach it."""
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.add_chat(self.ACCOUNT, 'user', '메일 질문', self.ids['uid-1'])
            self.assertFalse(store.drop_room(self.ACCOUNT, self.ids['uid-1']))
            self.assertFalse(store.drop_room(self.ACCOUNT, ''))
            self.assertEqual(len(store.chat(self.ACCOUNT, self.ids['uid-1'])), 1)
            store.db.close()

    def test_deleted_mail_is_not_collected_again(self):
        """`seen` keeps the uid: the mail is still on the server, and the user threw
        away the copy on purpose."""
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.delete([self.ids['uid-1']])
            self.assertIn('uid-1', store.seen(self.ACCOUNT))
            store.db.close()

    def test_the_priority_order_matches_the_dashboard(self):
        """Two lists of the same four words; a rename must not silently reorder one."""
        self.assertEqual(PRIORITY_ORDER, tuple(name for name, _ in PRIORITIES))

    def test_the_offered_filters_are_exactly_the_ones_sql_knows(self):
        """A value in the dropdown with no SQL behind it would silently show everything.

        Held against state_where() rather than STATE_SQL because that is what search()
        calls: 답장 대기 carries a cutoff that moves with the calendar and so is built
        per call, and a test on the static table alone would not have seen it.
        """
        for state in STATES[1:]:
            clause, params = state_where(state, datetime.date(2026, 9, 15))
            self.assertTrue(clause, state)
            self.assertEqual(clause.count('?'), len(params), state)
        self.assertEqual(state_where('', datetime.date(2026, 9, 15)), ('', ()))
        self.assertEqual(STATES[0], '')
        self.assertLessEqual(set(STATE_SQL), set(STATES))

    def test_a_constant_with_a_quote_is_refused_rather_than_interpolated(self):
        self.assertEqual(sql_text('처리'), "'처리'")
        with self.assertRaises(ValueError):
            sql_text("' OR 1=1 --")


class AnalyzingTests(unittest.TestCase):
    """'분석 중': what the worker is holding right now, as a state the screens can read."""

    ACCOUNT = 'acct'

    def store(self, folder):
        store = Store(Path(folder) / 'mail.db')
        self.ids = {uid: store.add(self.ACCOUNT, uid, mail()) for uid in ('uid-1', 'uid-2')}
        return store

    def test_a_mail_in_codex_says_so_instead_of_waiting(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            self.assertEqual(state_of(store.detail(self.ids['uid-1'])), '분석 대기')
            store.mark_analyzing(self.ACCOUNT, self.ids['uid-1'])
            self.assertEqual(state_of(store.detail(self.ids['uid-1'])), ANALYZING)
            self.assertEqual(state_of(store.detail(self.ids['uid-2'])), '분석 대기')
            store.db.close()

    def test_a_retry_in_flight_is_being_analysed_not_failing(self):
        """'2회 실패' of a mail Codex is reading right now is a week-old fact."""
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.failed(self.ids['uid-1'], '분석 실패', 0)
            store.failed(self.ids['uid-1'], '분석 실패', 0)
            self.assertEqual(state_of(store.detail(self.ids['uid-1'])), '2회 실패')
            store.mark_analyzing(self.ACCOUNT, self.ids['uid-1'])
            self.assertEqual(state_of(store.detail(self.ids['uid-1'])), ANALYZING)
            store.db.close()

    def test_only_one_mail_carries_the_marker(self):
        """One Codex process at a time, so a second mark clears the first."""
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.mark_analyzing(self.ACCOUNT, self.ids['uid-1'])
            store.mark_analyzing(self.ACCOUNT, self.ids['uid-2'])
            self.assertEqual(store.analyzing(self.ACCOUNT), self.ids['uid-2'])
            self.assertEqual(state_of(store.detail(self.ids['uid-1'])), '분석 대기')
            store.mark_analyzing(self.ACCOUNT, '')
            self.assertEqual(store.analyzing(self.ACCOUNT), '')
            store.db.close()

    def test_the_filter_finds_it_and_the_other_two_do_not_claim_it(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.failed(self.ids['uid-1'], '분석 실패', 0)
            store.mark_analyzing(self.ACCOUNT, self.ids['uid-1'])
            for state, expected in ((ANALYZING, [self.ids['uid-1']]),
                                    (FAILED, []),
                                    ('분석 대기', [self.ids['uid-2']])):
                rows, total = store.search(self.ACCOUNT, state=state)
                self.assertEqual([row['id'] for row in rows], expected, state)
                self.assertEqual(total, len(expected), state)
            store.db.close()

    def test_an_analysed_mail_drops_the_marker_from_every_state(self):
        """analyzed() does not clear it; the worker's finally does. The state still
        reads from the result, because a mail with an answer is not being analysed."""
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.mark_analyzing(self.ACCOUNT, self.ids['uid-1'])
            store.analyzed(self.ids['uid-1'], {'sender': 'a@b.c', 'subject': 's',
                                               'attachments': []}, RESULT)
            self.assertEqual(state_of(store.detail(self.ids['uid-1'])), '미처리')
            rows, _ = store.search(self.ACCOUNT, state=ANALYZING)
            self.assertEqual(rows, [])
            store.db.close()

    def test_a_mail_in_codex_is_not_reset_and_the_count_says_so(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.analyzed(self.ids['uid-1'], {'sender': 'a@b.c', 'subject': 's',
                                               'attachments': []}, RESULT)
            store.mark_analyzing(self.ACCOUNT, self.ids['uid-1'])
            self.assertEqual(store.reset([self.ids['uid-1']], reanalyze=True), 0)
            self.assertIsNotNone(store.detail(self.ids['uid-1'])['result'])
            store.mark_analyzing(self.ACCOUNT, '')
            self.assertEqual(store.reset([self.ids['uid-1']], reanalyze=True), 1)
            self.assertIsNone(store.detail(self.ids['uid-1'])['result'])
            store.db.close()

    def test_a_mixed_selection_resets_what_it_can(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.mark_analyzing(self.ACCOUNT, self.ids['uid-1'])
            self.assertEqual(store.reset(list(self.ids.values())), 1)
            store.db.close()

    def test_sorting_by_state_puts_it_between_waiting_and_failed(self):
        with tempfile.TemporaryDirectory() as folder:
            store = self.store(folder)
            store.failed(self.ids['uid-2'], '분석 실패', 0)
            store.mark_analyzing(self.ACCOUNT, self.ids['uid-1'])
            rows, _ = store.search(self.ACCOUNT, sort='state', desc=False)
            self.assertEqual([state_of(row) for row in rows], [ANALYZING, '1회 실패'])
            store.db.close()


class BoardHiddenTests(unittest.TestCase):
    """A mail card swept off the 할 일 판. The mail itself is untouched."""

    def test_the_flag_is_the_mails_own_and_survives_a_reopen(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'mail.db'
            store = Store(path)
            ident = store.add('acct', 'uid-1', mail())
            self.assertEqual(store.detail(ident)['todo_hidden'], 0)
            store.set_todo_hidden(ident)
            store.db.close()
            store = Store(path)
            self.assertEqual(store.detail(ident)['todo_hidden'], 1)
            self.assertEqual([row['todo_hidden'] for row in store.page('acct')], [1])
            store.set_todo_hidden(ident, False)
            self.assertEqual(store.detail(ident)['todo_hidden'], 0)
            store.db.close()


class NoteStoreTests(unittest.TestCase):
    """메모. The table is new, so nothing here is a migration — only the rules."""

    @contextmanager
    def opened(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                yield store
            finally:
                store.db.close()

    def test_pinned_first_then_most_recently_written(self):
        with self.opened() as store:
            first = store.add_note('acct', '하나')
            second = store.add_note('acct', '둘')
            third = store.add_note('acct', '셋')
            store.set_note_pinned(first)
            # Every id in one burst shares a `now()` string, so the order here is
            # decided by the id tie-break and not by whatever sqlite felt like.
            self.assertEqual([row['id'] for row in store.notes('acct')],
                             [first, third, second])

    def test_only_the_text_moves_a_memo_to_the_front(self):
        """The clock is pinned: a Windows one ticks about every 15ms.

        Left to the real clock this asserts the tick resolution rather than the
        ordering rule, and fails on Windows and nowhere else — which is exactly what
        it did. The stamps below are a second apart because that is what a person
        editing a memo actually produces.
        """
        stamps = ['2026-09-11T01:00:00+00:00', '2026-09-11T01:00:01+00:00',
                  '2026-09-11T01:00:02+00:00']
        with self.opened() as store:
            with patch('mail_assistant.core.now', side_effect=stamps):
                first = store.add_note('acct', '하나')
                second = store.add_note('acct', '둘')
                store.set_note_color(first, 'green')      # no stamp: colour moves nothing
                self.assertEqual([row['id'] for row in store.notes('acct')],
                                 [second, first])
                store.set_note_text(first, '하나 고침')
            self.assertEqual([row['id'] for row in store.notes('acct')][0], first)

    def test_memos_written_in_one_tick_fall_back_on_the_id(self):
        """Two stamps can be equal, and then something still has to decide.

        `now()` is a string off a clock that may not have moved between two writes, so
        `updated DESC` alone has nothing left to order by and sqlite may answer
        differently on each call — the same hazard every mail list ends with `rowid`
        for. Newest id first, which is the order they were written in.
        """
        with self.opened() as store:
            with patch('mail_assistant.core.now', return_value='2026-09-11T01:00:00+00:00'):
                first = store.add_note('acct', '하나')
                second = store.add_note('acct', '둘')
                store.set_note_text(first, '하나 고침')
            for _ in range(3):
                self.assertEqual([row['id'] for row in store.notes('acct')],
                                 [second, first])

    def test_a_mail_filter_narrows_to_that_mails_own(self):
        with self.opened() as store:
            ident = store.add('acct', 'uid-1', mail())
            attached = store.add_note('acct', '붙임', mail_id=ident)
            store.add_note('acct', '자유')
            self.assertEqual([row['id'] for row in store.notes('acct', ident)], [attached])
            self.assertEqual([row['id'] for row in store.notes('acct', '')],
                             [row['id'] for row in store.notes('acct')
                              if not row['mail_id']])

    def test_deleting_a_mail_keeps_the_memo_and_cuts_it_loose(self):
        """The analysis and the draft came from the mail; a memo is the user's own."""
        with self.opened() as store:
            ident = store.add('acct', 'uid-1', mail())
            note = store.add_note('acct', '박 과장님 견적 건', mail_id=ident)
            store.delete([ident])
            rows = list(store.notes('acct'))
            self.assertEqual([row['id'] for row in rows], [note])
            self.assertEqual(rows[0]['text'], '박 과장님 견적 건')
            self.assertEqual(rows[0]['mail_id'], '')

    def test_a_blank_memo_is_swept_and_a_written_one_is_not(self):
        with self.opened() as store:
            blank = store.add_note('acct')
            spaces = store.add_note('acct', '   \n ')
            written = store.add_note('acct', '적은 것')
            store.delete_empty_notes('acct')
            self.assertEqual([row['id'] for row in store.notes('acct')], [written])
            self.assertNotIn(blank, [row['id'] for row in store.notes('acct')])
            self.assertNotIn(spaces, [row['id'] for row in store.notes('acct')])

    def test_the_sweep_spares_the_one_just_made(self):
        with self.opened() as store:
            store.add_note('acct')
            keep = store.add_note('acct')
            store.delete_empty_notes('acct', keep=keep)
            self.assertEqual([row['id'] for row in store.notes('acct')], [keep])

    def test_the_sweep_stops_at_the_account(self):
        with self.opened() as store:
            mine = store.add_note('acct')
            theirs = store.add_note('other')
            store.delete_empty_notes('acct')
            self.assertEqual([row['id'] for row in store.notes('acct')], [])
            self.assertEqual([row['id'] for row in store.notes('other')], [theirs])
            self.assertNotEqual(mine, theirs)

    def test_subjects_arrive_in_one_query_and_skip_what_is_gone(self):
        with self.opened() as store:
            ident = store.add('acct', 'uid-1', mail())
            self.assertEqual(store.mail_subjects([ident, 'gone', '', ident]),
                             {ident: '견적 요청'})
            self.assertEqual(store.mail_subjects([]), {})

    def test_hover_cards_arrive_the_same_way_and_carry_the_analysis(self):
        with self.opened() as store:
            ident = store.add('acct', 'uid-1', mail())
            found = store.mail_cards([ident, 'gone', '', ident])
            self.assertEqual(list(found), [ident])
            self.assertEqual(found[ident]['subject'], '견적 요청')
            # The columns the list itself reads, so a card needs no second query.
            self.assertIn('result', found[ident].keys())
            self.assertEqual(store.mail_cards([]), {})


class DayWindowTests(unittest.TestCase):
    def test_day_bounds_are_korean_midnight_in_utc(self):
        start, end = day_bounds(datetime.date(2026, 9, 11))
        self.assertEqual(start, '2026-09-10T15:00:00+00:00')
        self.assertEqual(end, '2026-09-11T15:00:00+00:00')

    def test_local_text_shifts_to_korean_time(self):
        self.assertEqual(local_text('2026-09-10T15:30:00+00:00'), '09-11 00:30')
        self.assertEqual(local_text(''), '')
        self.assertEqual(local_text('망가진 값'), '망가진 값')

    def test_day_counts_use_the_korean_day(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            ident = store.add('acct', 'uid-1', mail())
            # Collected 2026-09-11 00:30 KST, which is the previous UTC day.
            store.db.execute('UPDATE mail SET received=? WHERE id=?', ('2026-09-10T15:30:00+00:00', ident))
            store.db.commit()
            store.analyzed(ident, {'sender': 'a@b.c', 'subject': 's', 'attachments': []}, RESULT)
            store.db.execute('UPDATE mail SET analyzed_at=? WHERE id=?', ('2026-09-10T16:00:00+00:00', ident))
            store.db.commit()
            counts = store.day_counts('acct', *day_bounds(datetime.date(2026, 9, 11)))
            self.assertEqual((counts['collected'], counts['analyzed'], counts['exported']), (1, 1, 0))
            empty = store.day_counts('acct', *day_bounds(datetime.date(2026, 9, 10)))
            self.assertEqual(empty['collected'], 0)
            store.db.close()


class OverviewTests(unittest.TestCase):
    TODAY = datetime.date(2026, 9, 11)

    def rows(self, folder):
        store = Store(Path(folder) / 'mail.db')
        urgent = dict(RESULT, priority='긴급', category='견적·계약',
                      events=[{'title': '견적 회신', 'start': '', 'deadline': '2026-09-14',
                               'evidence': 'e', 'needs_review': False}])
        far = dict(RESULT, priority='낮음', category='공지', reply_needed=False,
                   reply_subject='', reply_draft='',
                   events=[{'title': '정기 점검', 'start': '', 'deadline': '2026-10-20',
                            'evidence': 'e', 'needs_review': False}])
        ids = {}
        for uid, result, stamp in (('uid-1', urgent, '2026-09-11T01:00:00+00:00'),
                                   ('uid-2', far, '2026-09-11T02:00:00+00:00')):
            ids[uid] = store.add('acct', uid, mail())
            store.analyzed(ids[uid], {'sender': 'a@b.c', 'subject': 's', 'attachments': []}, result)
            store.db.execute('UPDATE mail SET received=? WHERE id=?', (stamp, ids[uid]))
        ids['uid-3'] = store.add('acct', 'uid-3', mail())   # collected, not analysed yet
        # Stamped like the others: store.add() uses the real clock, so leaving it made
        # the 최근 7일 assertion below pass only when run on TODAY itself.
        store.db.execute('UPDATE mail SET received=? WHERE id=?',
                         ('2026-09-11T03:00:00+00:00', ids['uid-3']))
        store.db.commit()
        return store, ids

    def test_cards_count_what_the_screen_shows(self):
        with tempfile.TemporaryDirectory() as folder:
            store, _ = self.rows(folder)
            data = overview(store.page('acct'), self.TODAY)
            self.assertEqual(data['cards']['미처리 메일'], 2)      # analysed and not handled
            self.assertEqual(data['cards']['긴급·높음'], 1)
            self.assertEqual(data['cards']['7일 내 마감'], 1)      # October deadline excluded
            self.assertEqual(data['cards']['검토 전 초안'], 1)     # the 공지 needs no reply
            self.assertEqual(data['waiting'], 1)
            store.db.close()

    def test_handled_mail_leaves_the_cards_and_the_deadlines(self):
        with tempfile.TemporaryDirectory() as folder:
            store, ids = self.rows(folder)
            store.set_handled(ids['uid-1'], HANDLED)   # the one with the 09-14 deadline
            data = overview(store.page('acct'), self.TODAY)
            self.assertEqual(data['cards']['미처리 메일'], 1)
            self.assertEqual(data['cards']['7일 내 마감'], 0)
            self.assertEqual(data['upcoming'], [])
            store.db.close()

    def test_missed_deadlines_stay_visible(self):
        with tempfile.TemporaryDirectory() as folder:
            store, _ = self.rows(folder)
            late = dict(RESULT, events=[{'title': '지난 회신', 'start': '', 'deadline': '2026-09-04',
                                         'evidence': 'e', 'needs_review': False}])
            ident = store.add('acct', 'uid-4', mail())
            store.analyzed(ident, {'sender': 'a@b.c', 'subject': 's', 'attachments': []}, late)
            data = overview(store.page('acct'), self.TODAY)
            self.assertEqual([entry.label for _, entry in data['past_due']], ['■ 지난 회신'])
            self.assertEqual(data['cards']['7일 내 마감'], 1)   # a missed deadline is not 임박
            store.set_handled(ident, HANDLED)
            self.assertEqual(overview(store.page('acct'), self.TODAY)['past_due'], [])
            store.db.close()

    def test_past_due_lists_the_nearest_miss_first(self):
        events = collect([['1:0', 'mail-old', '오래된 마감', '', '2026-08-01', '', '', ''],
                          ['2:0', 'mail-new', '어제 마감', '', '2026-09-10', '', '', ''],
                          ['3:0', 'mail-next', '앞으로', '', '2026-09-14', '', '', '']])
        self.assertEqual([entry.label for _, entry in past_due(events, self.TODAY)],
                         ['■ 어제 마감', '■ 오래된 마감'])
        self.assertEqual(past_due(events, self.TODAY, limit=1)[0][0], datetime.date(2026, 9, 10))

    def test_edited_draft_is_no_longer_review_pending(self):
        with tempfile.TemporaryDirectory() as folder:
            store, ids = self.rows(folder)
            store.set_draft(ids['uid-1'], '사람이 고친 초안')   # the only one needing a reply
            self.assertEqual(overview(store.page('acct'), self.TODAY)['cards']['검토 전 초안'], 0)
            store.db.close()

    def test_distributions_and_recent_week(self):
        with tempfile.TemporaryDirectory() as folder:
            store, _ = self.rows(folder)
            data = overview(store.page('acct'), self.TODAY)
            self.assertEqual(data['categories']['견적·계약'], 1)
            self.assertEqual(data['categories']['공지'], 1)
            self.assertEqual(data['priorities'], {'긴급': 1, '높음': 0, '보통': 0, '낮음': 1})
            self.assertEqual(len(data['recent']), 7)
            self.assertEqual(data['recent'][-1], ('2026-09-11', 3))   # includes the unanalysed one
            store.db.close()

    def test_broken_result_json_does_not_break_the_screen(self):
        with tempfile.TemporaryDirectory() as folder:
            store, _ = self.rows(folder)
            store.db.execute("UPDATE mail SET result='{망가진' WHERE uid='uid-1'")
            store.db.commit()
            data = overview(store.page('acct'), self.TODAY)
            self.assertEqual(data['cards']['긴급·높음'], 0)
            self.assertEqual(data['categories']['견적·계약'], 0)
            store.db.close()

    def test_calendar_reuses_the_schedule_shape(self):
        with tempfile.TemporaryDirectory() as folder:
            store, _ = self.rows(folder)
            events = overview(store.page('acct'), self.TODAY)['events']
            text, _ = month_grid(2026, 9, events)
            self.assertIn('견적 회신', ''.join(''.join(line) for line in text))
            store.db.close()


class CredentialTests(unittest.TestCase):
    def test_unicode_password_roundtrip(self):
        credentials = {}
        def write(value, flags):
            blob = value['CredentialBlob']
            if not isinstance(blob, str):
                raise TypeError('Objects of type bytes can not be converted to Unicode')
            credentials[value['TargetName']] = {'CredentialBlob': blob.encode('utf-16-le')}
        api = types.SimpleNamespace(
            CRED_TYPE_GENERIC=1, CRED_PERSIST_LOCAL_MACHINE=2,
            CredWrite=write, CredRead=lambda target, kind: credentials[target],
        )
        with patch.dict('sys.modules', {'win32cred': api}):
            for password in ('ascii-password!123', '한글 비밀번호🔐', ' leading and trailing '):
                with self.subTest(password_type='unicode'):
                    save_password('Test@Example.com', password)
                    self.assertEqual(read_password('test@example.com'), password)

    def api(self, failure):
        def read(target, kind):
            raise failure
        return types.SimpleNamespace(CRED_TYPE_GENERIC=1, CredRead=read)

    def test_missing_credential_is_a_lookup_error(self):
        failure = OSError('Element not found.')
        failure.winerror = 1168
        with patch.dict('sys.modules', {'win32cred': self.api(failure)}):
            with self.assertRaises(LookupError):
                read_password('test@example.com')

    def test_other_windows_failures_propagate(self):
        # '없음' would send the user to re-enter a password that is already stored.
        failure = OSError('The Credential Manager service is disabled.')
        failure.winerror = 1058
        with patch.dict('sys.modules', {'win32cred': self.api(failure)}):
            with self.assertRaises(OSError) as caught:
                read_password('test@example.com')
        self.assertEqual(caught.exception.winerror, 1058)


class Cell:
    def __init__(self, sheet, row, col):
        self.sheet, self.Row, self.col = sheet, row, col
    def End(self, direction):
        return Cell(self.sheet, max(r for r, c in self.sheet.data if c == 1), 1)


class Range:
    def __init__(self, sheet, first, last):
        self.sheet, self.first, self.last = sheet, first, last
        self.NumberFormat = None
    @property
    def Value(self):
        return tuple(tuple(self.sheet.data.get((r, c)) for c in range(self.first.col, self.last.col + 1)) for r in range(self.first.Row, self.last.Row + 1))
    @Value.setter
    def Value(self, values):
        if self.NumberFormat != '@':
            raise AssertionError('text format required before setting values')
        for r, row in enumerate(values, self.first.Row):
            for c, value in enumerate(row, self.first.col):
                self.sheet.data[r, c] = value


class Row:
    def __init__(self, sheet, row):
        self.sheet, self.row = sheet, row
    def Delete(self):
        self.sheet.data = {(r - 1 if r > self.row else r, c): value
                           for (r, c), value in self.sheet.data.items() if r != self.row}


class Rows:
    Count = 1048576
    def __init__(self, sheet): self.sheet = sheet
    def __call__(self, row): return Row(self.sheet, row)


class Sheet:
    def __init__(self, data=None):
        self.data = {(1, 1): 'ID'} if data is None else dict(data)
        self.Rows = Rows(self)
        self.ListObjects = MagicMock()
    def Cells(self, row, col): return Cell(self, row, col)
    def Range(self, first, last): return Range(self, first, last)


class ExcelTests(unittest.TestCase):
    def test_retry_dedup_and_preserve_user_edit(self):
        sheet = Sheet()
        append_missing(sheet, [['a', '생성 초안']], 2)
        sheet.data[2, 2] = '사용자 수정'
        append_missing(sheet, [['a', '생성 초안'], ['b', '=HYPERLINK("bad")']], 2)
        self.assertEqual(sheet.data[2, 2], '사용자 수정')
        self.assertEqual(sheet.data[3, 2], '=HYPERLINK("bad")')
        self.assertEqual(len([k for k in sheet.data if k[1] == 1]), 3)

    def adapter(self, folder, running, readonly=False):
        path = Path(folder) / 'sample.xlsx'
        path.touch()
        pythoncom = MagicMock()
        pythoncom.com_error = OSError
        client = MagicMock()
        app, book = MagicMock(), MagicMock()
        app.Ready = True
        book.ReadOnly = readonly
        book.Application = app
        app.Workbooks.Open.return_value = book
        client.DispatchEx.return_value = app
        client.Dispatch.return_value = book
        moniker = MagicMock()
        moniker.GetDisplayName.return_value = str(path)
        pythoncom.GetRunningObjectTable.return_value.EnumRunning.return_value = [moniker] if running else []
        modules = {'pythoncom': pythoncom, 'win32com': types.ModuleType('win32com'), 'win32com.client': client}
        modules['win32com'].client = client
        return path, modules, app, book

    def test_running_workbook_saved_but_never_closed(self):
        with tempfile.TemporaryDirectory() as folder:
            path, modules, app, book = self.adapter(folder, True)
            with patch.dict('sys.modules', modules), patch.object(Excel, '_write'):
                Excel(path).update([], {})
            book.Save.assert_called_once()
            book.Close.assert_not_called()
            app.Quit.assert_not_called()

    def test_owned_excel_cleanup_after_save_failure(self):
        with tempfile.TemporaryDirectory() as folder:
            path, modules, app, book = self.adapter(folder, False)
            book.Save.side_effect = OSError('locked')
            with patch.dict('sys.modules', modules), patch.object(Excel, '_write'):
                with self.assertRaisesRegex(ExcelUpdateError, '엑셀 파일 저장.*locked'):
                    Excel(path).update([], {})
            book.Close.assert_called_once_with(SaveChanges=False)
            app.Quit.assert_called_once()

    def test_readonly_workbook_is_not_written(self):
        with tempfile.TemporaryDirectory() as folder:
            path, modules, app, book = self.adapter(folder, True, True)
            with patch.dict('sys.modules', modules), patch.object(Excel, '_write') as write:
                with self.assertRaises(RuntimeError):
                    Excel(path).update([], {})
            write.assert_not_called()
            book.Save.assert_not_called()
            app.Quit.assert_not_called()

    def test_table_created_with_explicit_header_flag(self):
        sheet, head = MagicMock(), MagicMock()
        sheet.ListObjects.Count = 0
        ensure_table(sheet, head, 9)
        source = sheet.Range.return_value
        sheet.ListObjects.Add.assert_called_once_with(SourceType=1, Source=source, XlListObjectHasHeaders=1)
        sheet.ListObjects.Count = 1
        ensure_table(sheet, head, 9)
        self.assertEqual(sheet.ListObjects.Add.call_count, 1)

    def test_named_arguments_fall_back_to_positional_add(self):
        sheet, head = MagicMock(), MagicMock()
        sheet.ListObjects.Count = 0
        sheet.ListObjects.Add.side_effect = [TypeError('named arguments unsupported'), None]
        ensure_table(sheet, head, 9)
        self.assertEqual(sheet.ListObjects.Add.call_args, call(1, sheet.Range.return_value))

    def test_generated_column_headers_are_repaired(self):
        headers = HEADERS['메일 목록']
        data = {(1, c): f'Column{c}' for c in range(1, len(headers) + 1)}
        data.update({(2, c): value for c, value in enumerate(headers, 1)})
        data.update({(3, 1): 'mail-1', (3, 2): '2026-09-10T00:00:00+00:00'})
        sheet = Sheet(data)
        head = sheet.Range(sheet.Cells(1, 1), sheet.Cells(1, len(headers)))
        self.assertTrue(repair_generated_header(sheet, head, headers))
        self.assertEqual(row_text(sheet, 1, len(headers)), headers)
        self.assertEqual(sheet.data[2, 1], 'mail-1')  # data row moved up, header copy removed
        self.assertNotIn((4, 1), sheet.data)

    def test_generated_header_above_real_data_keeps_the_data(self):
        headers = HEADERS['메일 목록']
        data = {(1, c): f'Column{c}' for c in range(1, len(headers) + 1)}
        data[2, 1] = 'mail-1'
        sheet = Sheet(data)
        head = sheet.Range(sheet.Cells(1, 1), sheet.Cells(1, len(headers)))
        self.assertTrue(repair_generated_header(sheet, head, headers))
        self.assertEqual(row_text(sheet, 1, len(headers)), headers)
        self.assertEqual(sheet.data[2, 1], 'mail-1')

    def test_correct_headers_are_not_touched(self):
        headers = HEADERS['메일 목록']
        sheet = Sheet({(1, c): value for c, value in enumerate(headers, 1)})
        head = sheet.Range(sheet.Cells(1, 1), sheet.Cells(1, len(headers)))
        self.assertFalse(repair_generated_header(sheet, head, headers))
        self.assertEqual(row_text(sheet, 1, len(headers)), headers)

    def test_header_spacing_difference_is_rewritten_not_rejected(self):
        headers = HEADERS['메일 목록']
        sheets = {}
        for name, columns in HEADERS.items():
            sheet = MagicMock()
            spaced = [columns[0] + ' '] + list(columns[1:])
            sheet.Range.return_value.Value = (tuple(spaced),)
            sheet.ListObjects.Count = 0  # skip the resize pass, which needs a live sheet
            sheets[name] = sheet
        book = MagicMock()
        book.Worksheets.side_effect = sheets.__getitem__
        Excel._write(book, [], {})
        self.assertEqual(sheets['메일 목록'].Range.return_value.Value, (tuple(headers),))

    def test_renamed_column_names_every_mismatch(self):
        headers = list(HEADERS['메일 목록'])
        sheets = {}
        for name, columns in HEADERS.items():
            sheet = MagicMock()
            wrong = list(columns)
            if len(wrong) > 2:
                wrong[2] = '보낸사람'
            sheet.Range.return_value.Value = (tuple(wrong),)
            sheet.ListObjects.Count = 0
            sheets[name] = sheet
        book = MagicMock()
        book.Worksheets.side_effect = sheets.__getitem__
        with self.assertRaisesRegex(RuntimeError, 'C1: 현재 "보낸사람" / 기대 "발신자"'):
            Excel._write(book, [], {})

    def test_header_only_sheet_recovers_missing_table(self):
        from mail_assistant.core import HEADERS
        sheets = {}
        for name, headers in HEADERS.items():
            sheet = MagicMock()
            sheet.Range.return_value.Value = (tuple(headers),)
            sheet.ListObjects.Count = 0
            sheets[name] = sheet
        book = MagicMock()
        book.Worksheets.side_effect = sheets.__getitem__
        Excel._write(book, [], {})
        for name, sheet in sheets.items():
            if name != '실행 상태':
                sheet.ListObjects.Add.assert_called_once_with(
                    SourceType=1, Source=sheet.Range.return_value, XlListObjectHasHeaders=1)
            else:
                sheet.ListObjects.Add.assert_not_called()

    def test_cleanup_failure_does_not_hide_save_failure(self):
        with tempfile.TemporaryDirectory() as folder:
            path, modules, app, book = self.adapter(folder, False)
            book.Save.side_effect = OSError('save failed')
            app.Quit.side_effect = OSError('quit failed')
            with patch.dict('sys.modules', modules), patch.object(Excel, '_write'):
                with self.assertRaisesRegex(ExcelUpdateError, '엑셀 파일 저장.*save failed'):
                    Excel(path).update([], {})
            modules['pythoncom'].CoUninitialize.assert_called_once()

    def test_com_error_description_and_codes(self):
        error = RuntimeError('generic message')
        error.hresult = -2147352567
        error.excepinfo = (0, 'Microsoft Excel', 'LinkSource is invalid', None, 0, -2146827284)
        detail = error_detail(ExcelUpdateError('표 만들기', error))
        self.assertIn('표 만들기', detail)
        self.assertIn('LinkSource is invalid', detail)
        self.assertIn('0x80020009', detail)
        self.assertIn('0x800A03EC', detail)


class SettingsTests(unittest.TestCase):
    VALID = {'email': ' user@example.com ', 'host': 'pop3s.hiworks.com', 'port': '995',
             'interval': '180', 'workbook': 'C:\\Users\\a\\메일.xlsx', 'model': '', 'password': 'secret'}

    def test_values_are_trimmed_and_typed(self):
        result = normalize(self.VALID)
        self.assertEqual(result['email'], 'user@example.com')
        self.assertEqual((result['port'], result['interval']), (995, 180))
        self.assertNotIn('password', result)      # never written to config.json

    def test_rejects_bad_settings(self):
        for field, value, message in (('email', 'nope', '메일 계정'), ('host', '', '수신 서버'),
                                      ('interval', '30', '60초 이상'), ('port', '70000', '1~65535'),
                                      ('port', 'abc', '숫자'), ('workbook', '메일.xlsx', '절대 경로'),
                                      ('workbook', 'C:\\a\\메일.xls', '.xlsx')):
            with self.subTest(field=field, value=value):
                with self.assertRaisesRegex(ValueError, message):
                    normalize({**self.VALID, field: value})


class FieldErrorTests(unittest.TestCase):
    """The inline form and the save gate must never disagree about what is valid."""

    VALID = SettingsTests.VALID

    def test_valid_settings_have_no_field_errors(self):
        self.assertEqual(field_errors(self.VALID), {})

    def test_each_bad_field_is_named(self):
        for field, value in (('email', 'nope'), ('host', ''), ('interval', '30'),
                             ('port', '70000'), ('port', 'abc'), ('workbook', '메일.xlsx'),
                             ('workbook', 'C:\\a\\메일.xls')):
            with self.subTest(field=field, value=value):
                errors = field_errors({**self.VALID, field: value})
                self.assertIn(field, errors)

    def test_the_two_checks_agree_on_every_case(self):
        cases = [self.VALID,
                 {**self.VALID, 'email': 'nope'},
                 {**self.VALID, 'host': ''},
                 {**self.VALID, 'port': 'abc'},
                 {**self.VALID, 'port': '0'},
                 {**self.VALID, 'interval': '59'},
                 {**self.VALID, 'interval': ''},
                 {**self.VALID, 'workbook': ''},
                 {**self.VALID, 'workbook': 'C:\\a\\b.xls'},
                 {**self.VALID, 'model': 'gpt-5.6-luna'},
                 {**self.VALID, 'model': '두 단어'},
                 {}]
        for values in cases:
            with self.subTest(values=values):
                try:
                    normalize(values)
                except ValueError:
                    refused = True
                else:
                    refused = False
                self.assertEqual(refused, bool(field_errors(values)))


class ModelChoiceTests(unittest.TestCase):
    """The dropdown is built from Codex's own cache, so a retired slug cannot linger."""

    CACHE = {'models': [
        {'slug': 'gpt-6-astra', 'display_name': 'GPT-6-Astra', 'visibility': 'list',
         'description': '가장 뛰어난 모델', 'priority': 1},
        {'slug': 'gpt-reserve', 'display_name': 'GPT-Reserve', 'visibility': 'hide',
         'description': '숨김'},
        {'slug': 'gpt-5.6-luna', 'display_name': 'GPT-5.6-Luna', 'visibility': 'list',
         'priority': 8},
    ]}

    def test_only_the_models_the_account_may_pick_are_offered(self):
        self.assertEqual([slug for slug, _, _ in model_choices(self.CACHE)],
                         ['gpt-6-astra', 'gpt-5.6-luna'])

    def test_the_hint_leads_with_the_slug_that_goes_on_the_command_line(self):
        self.assertEqual(model_choices(self.CACHE)[0][2], 'gpt-6-astra · 가장 뛰어난 모델')

    def test_a_model_with_no_description_is_still_named_by_its_slug(self):
        self.assertEqual(model_choices(self.CACHE)[1][1:], ('GPT-5.6-Luna', 'gpt-5.6-luna'))

    def test_a_cache_of_another_shape_is_no_models_rather_than_a_crash(self):
        for data in (None, {}, {'models': None}, {'models': ['gpt-6-astra']}, 'text'):
            with self.subTest(data=data):
                self.assertEqual(model_choices(data), [])

    def test_the_rows_open_with_the_default_and_offer_nothing_to_type_into(self):
        rows = model_rows(model_choices(self.CACHE))
        self.assertEqual(rows[0]['value'], '')
        # 직접 입력 is gone: every row is a slug Codex itself just listed.
        self.assertEqual([row['value'] for row in rows],
                         ['', 'gpt-6-astra', 'gpt-5.6-luna'])

    def test_a_saved_model_codex_no_longer_lists_keeps_its_own_row(self):
        rows = model_rows(model_choices(self.CACHE), 'gpt-5.1-codex')
        self.assertIn('gpt-5.1-codex', [row['value'] for row in rows])

    def test_a_saved_model_codex_still_lists_is_not_repeated(self):
        rows = model_rows(model_choices(self.CACHE), 'gpt-6-astra')
        self.assertEqual([row['value'] for row in rows].count('gpt-6-astra'), 1)

    def test_codex_own_priority_decides_the_order_the_traits_are_read_off(self):
        cache = {'models': [
            {'slug': 'light', 'visibility': 'list', 'priority': 9},
            {'slug': 'heavy', 'visibility': 'list', 'priority': 2},
        ]}
        self.assertEqual([slug for slug, _, _ in model_choices(cache)], ['heavy', 'light'])

    def test_a_cache_with_no_priority_at_all_keeps_the_order_it_was_written_in(self):
        cache = {'models': [{'slug': 'first', 'visibility': 'list'},
                            {'slug': 'second', 'visibility': 'list'}]}
        self.assertEqual([slug for slug, _, _ in model_choices(cache)], ['first', 'second'])

    def test_the_traits_run_from_the_top_of_the_list_to_the_bottom(self):
        self.assertEqual(model_traits(0, 3), GRADES[0])
        self.assertEqual(model_traits(1, 3), GRADES[1])
        self.assertEqual(model_traits(2, 3), GRADES[-1])
        # Two models are the two ends, not the top and the middle.
        self.assertEqual((model_traits(0, 2), model_traits(1, 2)), (GRADES[0], GRADES[-1]))

    def test_one_model_is_given_the_middle_because_there_is_no_comparison(self):
        self.assertEqual(model_traits(0, 1), GRADES[len(GRADES) // 2])

    def test_every_offered_model_says_what_it_costs_and_the_default_does_not(self):
        rows = model_rows(model_choices(self.CACHE))
        self.assertEqual(model_label(rows[0]), '기본값')
        self.assertEqual(model_label(rows[1]), 'GPT-6-Astra · 성능 높음 · 사용량 많음')
        self.assertEqual(model_label(rows[-1]), 'GPT-5.6-Luna · 가볍고 빠름 · 사용량 적음')

    def test_a_model_codex_stopped_listing_carries_no_trait_it_cannot_support(self):
        row = model_rows(model_choices(self.CACHE), 'gpt-5.1-codex')[-1]
        self.assertEqual((row['power'], row['cost']), ('', ''))
        self.assertEqual(model_label(row), 'gpt-5.1-codex')

    def test_a_missing_cache_is_an_empty_list_and_never_an_exception(self):
        with tempfile.TemporaryDirectory() as folder:
            self.assertEqual(read_models(Path(folder) / 'models_cache.json'), [])
            broken = Path(folder) / 'broken.json'
            broken.write_text('{not json', encoding='utf-8')
            self.assertEqual(read_models(broken), [])

    def test_a_real_cache_is_read_from_disk(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'models_cache.json'
            path.write_text(json.dumps(self.CACHE, ensure_ascii=False), encoding='utf-8')
            self.assertEqual([slug for slug, _, _ in read_models(path)],
                             ['gpt-6-astra', 'gpt-5.6-luna'])


class TestClient:
    def __init__(self, count=2, fail_pass=False):
        self.count, self.fail_pass, self.quit_called = count, fail_pass, False
    def user(self, user): pass
    def pass_(self, password):
        if self.fail_pass:
            raise poplib.error_proto(b'-ERR invalid password')
    def uidl(self):
        return b'+OK', [f'{index} uid-{index}'.encode() for index in range(1, self.count + 1)], 0
    def quit(self): self.quit_called = True
    def close(self): pass


class ConnectionCheckTests(unittest.TestCase):
    CONFIG = {'host': 'pop3s.hiworks.com', 'port': 995, 'email': 'user@example.com'}

    def test_reports_mailbox_size_and_always_quits(self):
        client = TestClient(count=2)
        message = check_connection(self.CONFIG, 'secret', factory=lambda *a, **k: client)
        self.assertIn('2건', message)
        self.assertTrue(client.quit_called)

    def test_authentication_failure_propagates(self):
        client = TestClient(fail_pass=True)
        with self.assertRaises(poplib.error_proto):
            check_connection(self.CONFIG, 'wrong', factory=lambda *a, **k: client)
        self.assertTrue(client.quit_called)

    def test_login_state_maps_the_three_outcomes(self):
        with patch('mail_assistant.services.check_login'):
            self.assertEqual(login_state()[0], '연결됨')
        with patch('mail_assistant.services.check_login', side_effect=RuntimeError('codex login 실행')):
            state, detail = login_state()
            self.assertEqual(state, '로그인 필요')
            self.assertIn('codex login', detail)
        with patch('mail_assistant.services.check_login', side_effect=OSError('없음')):
            self.assertEqual(login_state()[0], '확인 실패')


class ConnectionStepTests(unittest.TestCase):
    """The 진단 page's step list. It stops at the first failure and never echoes the password."""

    CONFIG = {'host': 'pop3s.hiworks.com', 'port': 995, 'email': 'me@corp.example'}

    def test_every_step_passes_on_a_healthy_mailbox(self):
        client = TestClient()
        steps = connection_steps(self.CONFIG, 'secret', factory=lambda *a, **k: client,
                                 resolve=lambda host: '203.0.113.7')
        self.assertTrue(all(ok for _, ok, _ in steps))
        self.assertEqual([label for label, _, _ in steps],
                         ['주소 조회', 'SSL 연결 995', '계정 전송', '비밀번호 인증', '사서함 조회'])
        self.assertIn('203.0.113.7', steps[0][2])

    def test_a_dns_failure_stops_before_connecting(self):
        def refuse(host):
            raise OSError('이름을 확인할 수 없습니다')
        steps = connection_steps(self.CONFIG, 'secret', factory=lambda *a, **k: 1 / 0,
                                 resolve=refuse)
        self.assertEqual(len(steps), 1)
        self.assertFalse(steps[0][1])

    def test_a_connect_failure_stops_before_authenticating(self):
        def refuse(*args, **kwargs):
            raise OSError('연결이 거부되었습니다')
        steps = connection_steps(self.CONFIG, 'secret', factory=refuse,
                                 resolve=lambda host: '203.0.113.7')
        self.assertEqual([label for label, _, _ in steps], ['주소 조회', 'SSL 연결 995'])
        self.assertFalse(steps[-1][1])

    def test_a_wrong_password_is_reported_without_printing_it(self):
        class Refuses(TestClient):
            def pass_(self, password):
                raise poplib.error_proto(f'-ERR invalid password {password}')
        steps = connection_steps(self.CONFIG, 'secret', factory=lambda *a, **k: Refuses(),
                                 resolve=lambda host: '203.0.113.7')
        self.assertFalse(steps[-1][1])
        detail = steps[-1][2]
        self.assertNotIn('secret', detail)
        self.assertIn('***', detail)


class ChatReplyTests(unittest.TestCase):
    """A chat turn reuses analyze()'s hardened invocation, schema and all."""

    def run_chat(self, written=None, returncode=0):
        seen = {}

        def fake_run(command, **kwargs):
            seen['command'] = command
            seen['input'] = kwargs.get('input', '')
            target = command[command.index('-o') + 1]
            if written is not None:
                Path(target).write_text(json.dumps(written, ensure_ascii=False), encoding='utf-8')
            return types.SimpleNamespace(returncode=returncode, stdout='', stderr='')

        with patch('mail_assistant.services.codex_command', return_value=['codex']), \
             patch('mail_assistant.services.subprocess.run', side_effect=fake_run):
            try:
                reply = chat_reply('언제까지 회신해야 하나요?',
                                   history=[('user', '이전 질문'), ('codex', '이전 답')],
                                   mail={'subject': '견적 요청', 'body': '9월 14일까지'},
                                   config={'model': 'gpt-5'})
            except Exception as exc:
                return seen, exc
        return seen, reply

    def test_the_answer_comes_back_from_the_schema_file(self):
        seen, reply = self.run_chat({'reply': '9월 14일까지입니다.'})
        self.assertEqual(reply, '9월 14일까지입니다.')

    def test_the_sandbox_flags_are_the_same_as_analysis(self):
        seen, _ = self.run_chat({'reply': 'ok'})
        command = seen['command']
        self.assertIn('--sandbox', command)
        self.assertEqual(command[command.index('--sandbox') + 1], 'read-only')
        self.assertIn('features.shell_tool=false', command)
        self.assertIn('approval_policy="never"', command)
        self.assertIn('--ephemeral', command)
        self.assertEqual(command[command.index('--model') + 1], 'gpt-5')

    def test_the_prompt_carries_the_injection_guard_and_the_transcript(self):
        seen, _ = self.run_chat({'reply': 'ok'})
        self.assertIn('신뢰하지 않는 입력', seen['input'])
        self.assertIn('도구를 사용하지 마세요', seen['input'])
        self.assertIn('이전 질문', seen['input'])
        self.assertIn('견적 요청', seen['input'])

    def test_a_failed_turn_says_nothing_about_stdout(self):
        seen, error = self.run_chat(None, returncode=1)
        self.assertIsInstance(error, RuntimeError)
        self.assertIn('Codex 응답을 받지 못했어요', str(error))

    def test_the_schema_allows_only_a_reply_string(self):
        shape = chat_schema()
        self.assertEqual(shape['required'], ['reply'])
        self.assertFalse(shape['additionalProperties'])


class DraftTests(unittest.TestCase):
    """답변 초안: 말투와 방향을 골라서, 그리고 부르지 않으면 만들지 않는다."""

    MAIL = {'subject': 'Quotation request', 'sender': 'john@globaltrade.example',
            'received': '2026-09-15 11:37', 'body': 'Please quote 200 units.',
            'summary': '견적 요청', 'requests': '단가와 납기', 'action': '견적 회신'}

    def run_draft(self, written=None, returncode=0, tone='', way='', mail=None):
        seen = {}

        def fake_run(command, **kwargs):
            seen['command'] = command
            seen['input'] = kwargs.get('input', '')
            target = command[command.index('-o') + 1]
            if written is not None:
                Path(target).write_text(json.dumps(written, ensure_ascii=False),
                                        encoding='utf-8')
            return types.SimpleNamespace(returncode=returncode, stdout='', stderr='')

        with patch('mail_assistant.services.codex_command', return_value=['codex']), \
             patch('mail_assistant.services.subprocess.run', side_effect=fake_run):
            try:
                answer = draft(mail or self.MAIL, tone, way, {'model': 'gpt-5'})
            except Exception as exc:
                return seen, exc
        return seen, answer

    def sent(self, seen):
        """The payload as Codex received it, which is the only place it is."""
        return json.loads(seen['input'][seen['input'].index('{'):])

    def test_the_subject_and_the_body_both_come_back(self):
        seen, answer = self.run_draft({'subject': 'Re: Quotation request',
                                       'draft': '확인 후 회신드리겠습니다.'})
        self.assertEqual(answer, ('Re: Quotation request', '확인 후 회신드리겠습니다.'))

    def test_the_chosen_tone_and_direction_reach_the_prompt(self):
        seen, _ = self.run_draft({'subject': 's', 'draft': 'd'},
                                 tone='간결하게', way='거절')
        payload = self.sent(seen)
        self.assertIn('간결하게', payload['tone'])
        # 거절 goes as a sentence, not as the bare word: a model reads '거절' alone
        # as what the mail is about rather than as what the reply should do.
        self.assertIn('받아들이기 어렵다', payload['way'])

    def test_the_free_choices_ask_for_nothing_at_all(self):
        """'기본 말투로 쓰세요' is a constraint a model will invent a meaning for."""
        seen, _ = self.run_draft({'subject': 's', 'draft': 'd'},
                                 tone=DRAFT_TONES[0], way=DRAFT_WAYS[0])
        payload = self.sent(seen)
        self.assertEqual((payload['tone'], payload['way']), ('', ''))

    def test_every_direction_the_screen_offers_has_a_sentence(self):
        for way in DRAFT_WAYS[1:]:
            seen, _ = self.run_draft({'subject': 's', 'draft': 'd'}, way=way)
            self.assertTrue(self.sent(seen)['way'], way)

    def test_the_reply_is_written_in_the_language_the_mail_came_in(self):
        seen, _ = self.run_draft({'subject': 's', 'draft': 'd'})
        self.assertIn('받은 메일과 같은 언어로', seen['input'])
        self.assertIn('서명과 연락처는 넣지 마세요', seen['input'])

    def test_the_body_goes_in_on_stdin_and_is_clipped(self):
        mail = dict(self.MAIL, body='a' * (DRAFT_BODY_LIMIT + 500))
        seen, _ = self.run_draft({'subject': 's', 'draft': 'd'}, mail=mail)
        self.assertEqual(len(self.sent(seen)['body']), DRAFT_BODY_LIMIT)
        self.assertNotIn('Quotation request', ' '.join(seen['command']))

    def test_a_failure_says_nothing_about_stdout(self):
        seen, error = self.run_draft(None, returncode=1)
        self.assertIsInstance(error, RuntimeError)
        self.assertIn('초안을 만들지 못했어요', str(error))

    def test_the_schema_allows_only_a_subject_and_a_draft(self):
        shape = draft_schema()
        self.assertEqual(shape['required'], ['subject', 'draft'])
        self.assertFalse(shape['additionalProperties'])


class AnalysisDraftTests(unittest.TestCase):
    """분석은 더 이상 초안을 쓰지 않는다 — 말투를 아무도 고르지 않았기 때문."""

    def test_the_analysis_is_told_to_leave_the_draft_empty(self):
        seen = {}

        def execute(command, **kwargs):
            seen['input'] = kwargs['input']
            Path(command[command.index('-o') + 1]).write_text(json.dumps(RESULT),
                                                              encoding='utf-8')
            return subprocess.CompletedProcess(command, 0)

        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            store.add(account_key(CONFIG), 'new', mail())
            row = store.pending(account_key(CONFIG), 0)[0]
            with patch('mail_assistant.services.codex_command', return_value=['codex']), \
                 patch('mail_assistant.services.subprocess.run', side_effect=execute):
                analyze(row, CONFIG)
            store.db.close()
        self.assertIn('reply_draft는 항상 빈 문자열', seen['input'])
        # 답장이 필요한지와 제목은 그대로 판단한다 — 큐가 그것으로 서 있다.
        self.assertIn('reply_needed=false', seen['input'])


class ReplyDraftStoreTests(unittest.TestCase):
    """만든 초안은 분석 안으로 들어간다. draft_edit은 사람의 자리로 남는다."""

    def store(self, folder):
        store = Store(Path(folder) / 'mail.db')
        account = account_key(CONFIG)
        ident = store.add(account, 'uid-1', mail())
        store.analyzed(ident, {'sender': 'a@b.c', 'subject': '견적', 'attachments': []},
                       dict(RESULT, reply_draft='', reply_needed=False))
        return store, ident

    def test_the_draft_lands_in_the_analysis_and_not_in_the_edit(self):
        with tempfile.TemporaryDirectory() as folder:
            store, ident = self.store(folder)
            self.assertTrue(store.set_reply_draft(ident, 'Re: 견적', '회신드리겠습니다.'))
            row = store.detail(ident)
            result = json.loads(row['result'])
            self.assertEqual(result['reply_draft'], '회신드리겠습니다.')
            self.assertEqual(result['reply_subject'], 'Re: 견적')
            # 사람이 손댔다는 표시가 아니므로 검토 큐에 그대로 남는다.
            self.assertEqual(row['draft_edit'], '')
            store.db.close()

    def test_a_new_draft_replaces_the_one_in_the_box(self):
        """The screen shows draft_edit or reply_draft, so an edit left in place would
        store a draft nobody could see. 초안 새로 만들기 is what the button says."""
        with tempfile.TemporaryDirectory() as folder:
            store, ident = self.store(folder)
            store.set_draft(ident, '사람이 고쳐 쓴 초안')
            store.set_reply_draft(ident, '', '새로 만든 초안')
            row = store.detail(ident)
            self.assertEqual(row['draft_edit'], '')
            self.assertEqual(json.loads(row['result'])['reply_draft'], '새로 만든 초안')
            store.db.close()

    def test_asking_for_a_draft_is_the_answer_to_답장이_필요한가(self):
        with tempfile.TemporaryDirectory() as folder:
            store, ident = self.store(folder)
            store.set_reply_draft(ident, '', '회신드리겠습니다.')
            self.assertTrue(json.loads(store.detail(ident)['result'])['reply_needed'])
            store.db.close()

    def test_a_mail_with_no_analysis_yet_is_told_so_rather_than_written_to(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            ident = store.add(account_key(CONFIG), 'uid-1', mail())
            self.assertFalse(store.set_reply_draft(ident, 'Re:', '초안'))
            self.assertFalse(store.set_reply_draft('없는 id', 'Re:', '초안'))
            store.db.close()


class TranslateTests(unittest.TestCase):
    """해외 메일 한 통을 한국어로, on the same hardened invocation analysis uses."""

    def run_translate(self, written=None, returncode=0, body='Dear Sir, please quote.'):
        seen = {}

        def fake_run(command, **kwargs):
            seen['command'] = command
            seen['input'] = kwargs.get('input', '')
            target = command[command.index('-o') + 1]
            if written is not None:
                Path(target).write_text(json.dumps(written, ensure_ascii=False),
                                        encoding='utf-8')
            return types.SimpleNamespace(returncode=returncode, stdout='', stderr='')

        with patch('mail_assistant.services.codex_command', return_value=['codex']), \
             patch('mail_assistant.services.subprocess.run', side_effect=fake_run):
            try:
                answer = translate(body, 'Quotation request', {'model': 'gpt-5'})
            except Exception as exc:
                return seen, exc
        return seen, answer

    def test_the_language_and_the_korean_both_come_back(self):
        seen, answer = self.run_translate({'language': '영어', 'korean': '견적을 요청드립니다.'})
        self.assertEqual(answer, ('영어', '견적을 요청드립니다.'))

    def test_the_sandbox_flags_are_the_same_as_analysis(self):
        seen, _ = self.run_translate({'language': '영어', 'korean': '번역'})
        command = seen['command']
        self.assertEqual(command[command.index('--sandbox') + 1], 'read-only')
        self.assertIn('features.shell_tool=false', command)
        self.assertIn('--ephemeral', command)
        self.assertEqual(command[command.index('--model') + 1], 'gpt-5')

    def test_the_body_goes_in_on_stdin_and_never_on_the_command_line(self):
        seen, _ = self.run_translate({'language': '영어', 'korean': '번역'})
        self.assertNotIn('Dear Sir', ' '.join(seen['command']))
        self.assertIn('Dear Sir', seen['input'])
        self.assertIn('신뢰하지 않는', seen['input'])
        self.assertIn('도구를 사용하지 마세요', seen['input'])

    def test_the_figures_are_told_to_stay_as_they_were_written(self):
        """A translated price or part number is the one thing worth nothing at all."""
        seen, _ = self.run_translate({'language': '영어', 'korean': '번역'})
        for kept in ('금액', '품번', '날짜'):
            self.assertIn(kept, seen['input'])

    def test_a_failure_says_nothing_about_stdout(self):
        seen, error = self.run_translate(None, returncode=1)
        self.assertIsInstance(error, RuntimeError)
        self.assertIn('번역하지 못했어요', str(error))

    def test_a_body_too_long_is_refused_rather_than_clipped(self):
        """Half a mail translated is worse than none: nothing would say which half."""
        seen, error = self.run_translate({'language': '영어', 'korean': 'x'},
                                         body='a' * (TRANSLATE_LIMIT + 1))
        self.assertIsInstance(error, RuntimeError)
        self.assertIn('초과', str(error))
        self.assertNotIn('command', seen)          # Codex was never asked

    def test_an_empty_body_is_not_sent_either(self):
        seen, error = self.run_translate({'language': '영어', 'korean': 'x'}, body='   ')
        self.assertIsInstance(error, RuntimeError)
        self.assertNotIn('command', seen)

    def test_the_schema_allows_only_the_two_fields(self):
        shape = translate_schema()
        self.assertEqual(shape['required'], ['language', 'korean'])
        self.assertFalse(shape['additionalProperties'])


class ReasoningEffortTests(unittest.TestCase):
    """생각의 깊이는 앱이 정한다 — 모델이 저마다 들고 있는 기본값에 맡기지 않는다."""

    def effort_of(self, call):
        seen = {}

        def fake_run(command, **kwargs):
            seen['command'] = command
            Path(command[command.index('-o') + 1]).write_text('{}', encoding='utf-8')
            return types.SimpleNamespace(returncode=1, stdout='', stderr='')

        with patch('mail_assistant.services.codex_command', return_value=['codex']), \
             patch('mail_assistant.services.subprocess.run', side_effect=fake_run):
            with self.assertRaises(RuntimeError):
                call()
        settings = [value for value in seen['command']
                    if str(value).startswith('model_reasoning_effort=')]
        self.assertEqual(len(settings), 1, seen['command'])
        return settings[0].split('=', 1)[1].strip('"')

    def test_the_call_that_reads_a_mail_for_the_first_time_thinks(self):
        """틀린 마감은 달력과 엑셀 일정 시트까지 흘러간다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            store.add(account_key(CONFIG), 'one', mail())
            row = store.pending(account_key(CONFIG), time.time())[0]
            self.assertEqual(self.effort_of(lambda: analyze(row, CONFIG)), EFFORT_THINK)
            store.db.close()

    def test_a_question_and_a_draft_think_too(self):
        self.assertEqual(self.effort_of(lambda: chat_reply('질문', config={})),
                         EFFORT_THINK)
        self.assertEqual(self.effort_of(lambda: draft({'body': 'hello'}, config={})),
                         EFFORT_THINK)

    def test_work_over_what_is_already_analysed_does_not(self):
        """브리핑의 입력은 analyze()가 이미 값을 치른 요약이고, 번역은 옮기는 일이다."""
        self.assertEqual(self.effort_of(lambda: briefing({}, {})), EFFORT_READ)
        self.assertEqual(self.effort_of(lambda: translate('hello', '', {})), EFFORT_READ)

    def test_both_levels_are_ones_every_listed_model_supports(self):
        """그 위는 240초 타임아웃에 걸릴 수 있고, 타임아웃은 곧 전체 백오프다."""
        self.assertEqual((EFFORT_READ, EFFORT_THINK), ('low', 'medium'))

    def test_the_clis_own_config_is_still_ignored(self):
        """--ignore-user-config가 있으므로 이 -c가 유일한 지시다."""
        self.assertEqual(self.effort_of(lambda: briefing({}, {})), EFFORT_READ)
        with patch('mail_assistant.services.codex_command', return_value=['codex']), \
             patch('mail_assistant.services.subprocess.run') as run:
            run.return_value = types.SimpleNamespace(returncode=1, stdout='', stderr='')
            with self.assertRaises(RuntimeError):
                briefing({}, {})
            self.assertIn('--ignore-user-config', run.call_args[0][0])


class CodexInvocationTests(unittest.TestCase):
    """The one `codex exec` every call in this app makes, and what each says when it fails."""

    def test_every_caller_goes_through_the_same_flags(self):
        seen = []

        def fake_run(command, **kwargs):
            seen.append(command)
            Path(command[command.index('-o') + 1]).write_text('{}', encoding='utf-8')
            return types.SimpleNamespace(returncode=1, stdout='', stderr='')

        callers = ((lambda: translate('hello', '', {}), '번역하지 못했어요'),
                   (lambda: chat_reply('질문', config={}), 'Codex 응답을 받지 못했어요'),
                   (lambda: briefing({}, {}), '브리핑을 만들지 못했어요'))
        with patch('mail_assistant.services.codex_command', return_value=['codex']), \
             patch('mail_assistant.services.subprocess.run', side_effect=fake_run):
            for call, said in callers:
                with self.assertRaises(RuntimeError) as caught:
                    call()
                # Each caller names what it was doing; none of them quotes stdout.
                self.assertIn(said, str(caught.exception))
        for command in seen:
            self.assertEqual(command[command.index('--sandbox') + 1], 'read-only')
            self.assertIn('--ignore-user-config', command)
            self.assertIn('--output-schema', command)


class MailtoTests(unittest.TestCase):
    def test_address_is_taken_from_the_display_form(self):
        self.assertEqual(address_of('김철수 <kim@a.com>'), 'kim@a.com')
        self.assertEqual(address_of('kim@a.com'), 'kim@a.com')
        self.assertEqual(address_of('이름만 있음'), '')

    def test_subject_and_body_are_encoded(self):
        link = mailto('김철수 <kim@a.com>', 'Re: 견적', '안녕하세요.\n회신드립니다.')
        self.assertTrue(link.startswith('mailto:kim@a.com?subject='))
        self.assertIn('&body=', link)
        self.assertNotIn(' ', link)
        self.assertNotIn('\n', link)

    def test_long_drafts_stay_under_the_hyperlink_limit(self):
        self.assertLessEqual(len(mailto('a <a@b.c>', 'x' * 300, 'y' * 9000)), 1800)

    def test_no_address_means_no_link(self):
        sheet = MagicMock()
        self.assertEqual(mailto('이름만', '제목', '초안'), '')
        link_to_draft(sheet, '답변 초안', 5, ['m1', '이름만', '제목', '초안', '', '검토 전'])
        sheet.Hyperlinks.Add.assert_not_called()
        link_to_draft(sheet, '답변 초안', 5, ['m1', 'a <a@b.c>', '제목', '초안', '', '검토 전'])
        arguments = sheet.Hyperlinks.Add.call_args.kwargs
        self.assertTrue(arguments['Address'].startswith('mailto:a@b.c'))
        sheet.Cells.assert_called_with(5, 3)     # 답변 제목 column


ROW_VALUES = ['m1', '김철수 <kim@example.com>', 'Re: 견적', '확인 후 회신드리겠습니다.', '', '검토 전']


class FirstColumnTests(unittest.TestCase):
    """COM's Range.Value has three shapes and only one of them is a tuple of rows."""

    def test_a_multi_cell_range_is_a_tuple_of_rows(self):
        self.assertEqual(first_column((('a',), ('b',), (None,))), ['a', 'b', None])

    def test_a_single_cell_range_is_a_bare_scalar(self):
        self.assertEqual(first_column('a'), ['a'])
        self.assertEqual(first_column(7), [7])

    def test_a_single_empty_cell_is_none(self):
        self.assertEqual(first_column(None), [])

    def test_a_one_cell_string_is_not_read_character_by_character(self):
        """The bug this replaced compared mail ids against the letters of one id."""
        self.assertEqual(first_column('abc123'), ['abc123'])

    def test_a_flat_tuple_is_accepted_too(self):
        self.assertEqual(first_column(('a', 'b')), ['a', 'b'])


class LinkTests(unittest.TestCase):
    def test_mail_rows_maps_ids_to_row_numbers(self):
        sheet = Sheet({(1, 1): '메일 ID', (2, 1): 'm1', (3, 1): 'm2', (4, 1): 'm3'})
        self.assertEqual(mail_rows(sheet), {'m1': 2, 'm2': 3, 'm3': 4})
        self.assertEqual(mail_rows(Sheet({(1, 1): '메일 ID'})), {})

    def test_link_points_at_the_mail_id_column_of_each_sheet(self):
        sheet = MagicMock()
        link_to_mail(sheet, '일정', 7, 5)
        arguments = sheet.Hyperlinks.Add.call_args.kwargs
        self.assertEqual(arguments['SubAddress'], "'메일 목록'!A5")
        self.assertEqual(arguments['Address'], '')
        sheet.Cells.assert_called_with(7, 2)      # 일정 keeps 메일 ID in column B
        link_to_mail(sheet, '우선순위', 7, 5)
        sheet.Cells.assert_called_with(7, 1)      # 우선순위 keeps it in column A

    def test_write_links_new_rows_but_not_the_mail_list(self):
        sheets = {}
        for name, columns in HEADERS.items():
            sheet = MagicMock()
            sheet.Range.return_value.Value = (tuple(columns),)
            sheet.ListObjects.Count = 0
            sheets[name] = sheet
        book = MagicMock()
        book.Worksheets.side_effect = sheets.__getitem__
        row = {'id': 'm1', 'received': '2026-09-11T00:00:00+00:00',
               'parsed': json.dumps({'sender': 'a@b.c', 'subject': 's', 'attachments': []}),
               'result': json.dumps(RESULT)}
        with patch('mail_assistant.excel.append_missing', return_value=[(7, ROW_VALUES)]), \
             patch('mail_assistant.excel.mail_rows', return_value={'m1': 5}), \
             patch('mail_assistant.excel.link_to_mail') as linked, \
             patch('mail_assistant.excel.sheet_events', return_value={}), \
             patch('mail_assistant.excel.apply_style'), \
             patch('mail_assistant.excel.update_calendar'), \
             patch('mail_assistant.excel.update_dashboard'):
            Excel._write(book, [row], {})
        self.assertEqual({call.args[1] for call in linked.call_args_list},
                         {'일정', '우선순위', '답변 초안'})
        self.assertTrue(all(call.args[2:] == (7, 5) for call in linked.call_args_list))

    def test_unknown_mail_id_is_left_unlinked(self):
        sheets = {}
        for name, columns in HEADERS.items():
            sheet = MagicMock()
            sheet.Range.return_value.Value = (tuple(columns),)
            sheet.ListObjects.Count = 0
            sheets[name] = sheet
        book = MagicMock()
        book.Worksheets.side_effect = sheets.__getitem__
        row = {'id': 'm1', 'received': '2026-09-11T00:00:00+00:00',
               'parsed': json.dumps({'sender': 'a@b.c', 'subject': 's', 'attachments': []}),
               'result': json.dumps(RESULT)}
        with patch('mail_assistant.excel.append_missing', return_value=[(7, ROW_VALUES)]), \
             patch('mail_assistant.excel.mail_rows', return_value={}), \
             patch('mail_assistant.excel.link_to_mail') as linked, \
             patch('mail_assistant.excel.sheet_events', return_value={}), \
             patch('mail_assistant.excel.apply_style'), \
             patch('mail_assistant.excel.update_calendar'), \
             patch('mail_assistant.excel.update_dashboard'):
            Excel._write(book, [row], {})
        linked.assert_not_called()


class CalendarTests(unittest.TestCase):
    TODAY = datetime.date(2026, 9, 10)
    ROWS = [
        ['m1:0', 'm1', '견적 회신 마감', '', '2026-09-16T18:00:00+09:00', '9월 16일까지', '', '미처리'],
        ['m1:1', 'm1', '사양 확정 회의', '2026-09-14T10:00:00+09:00', '', '', '', '미처리'],
        ['m2:0', 'm2', '유지보수 미팅', '2026-09-15', '', '', '필요', '미처리'],
        ['m3:0', 'm3', '도면 회신', '', '2026-08-25', '', '', '미처리'],
    ]

    def test_parse_day_accepts_iso_only(self):
        self.assertEqual(parse_day('2026-09-16'), (datetime.date(2026, 9, 16), ''))
        self.assertEqual(parse_day('2026-09-16T18:00:00+09:00'), (datetime.date(2026, 9, 16), '18:00'))
        self.assertEqual(parse_day(''), (None, ''))
        self.assertEqual(parse_day('다음 주 화요일'), (None, ''))
        self.assertEqual(parse_day('2026-02-30'), (None, ''))

    def test_collect_splits_start_deadline_and_review(self):
        events = collect(self.ROWS)
        deadline = events[datetime.date(2026, 9, 16)][0]
        self.assertEqual((deadline.label, deadline.kind), ('■ 18:00 견적 회신 마감', '마감'))
        self.assertEqual(deadline.mail_id, 'm1')
        self.assertEqual(deadline.row, 2)          # first data row of the 일정 sheet
        start = events[datetime.date(2026, 9, 14)][0]
        self.assertEqual((start.label, start.kind, start.row), ('▶ 10:00 사양 확정 회의', '시작', 3))
        # 확인 필요 wins over the start/deadline colour.
        review = events[datetime.date(2026, 9, 15)][0]
        self.assertEqual((review.label, review.kind), ('◆ 유지보수 미팅', '확인 필요'))

    def test_cell_text_colour_spans_line_up(self):
        entries = [Entry('■ 마감', '마감', 'm1', 2), Entry('▶ 시작', '시작', 'm1', 3)]
        text, spans = cell_text(datetime.date(2026, 9, 16), entries)
        self.assertEqual(text.split('\n'), ['16', '■ 마감', '▶ 시작'])
        for (start, length, color), entry in zip(spans, entries):
            self.assertEqual(text[start - 1:start - 1 + length], entry.label)
        self.assertEqual(spans[0][2], URGENT)

    def test_cell_text_caps_long_days(self):
        entries = [Entry(f'■ 일정 {index}', '마감', 'm1', 2 + index) for index in range(5)]
        text, spans = cell_text(datetime.date(2026, 9, 16), entries)
        self.assertEqual(len(spans), 3)
        self.assertTrue(text.endswith('…외 2건'))

    def test_month_grid_places_weeks_monday_first(self):
        text, cells = month_grid(2026, 9, collect(self.ROWS))
        self.assertEqual(text[0][0], '')          # 2026-09-01 is a Tuesday
        self.assertEqual(text[0][1], '1')
        self.assertEqual(text[2][2].split('\n')[0], '16')
        self.assertNotIn('도면 회신', ''.join(''.join(row) for row in text))  # August stays out
        self.assertEqual(cells[2, 0][0], datetime.date(2026, 9, 14))

    def test_overdue_lists_only_past_deadlines(self):
        self.assertEqual(overdue(collect(self.ROWS), self.TODAY),
                         [(datetime.date(2026, 8, 25), '■ 도면 회신')])

    def test_next_month_crosses_the_year(self):
        self.assertEqual(next_month(datetime.date(2026, 12, 31)), datetime.date(2027, 1, 1))
        self.assertEqual(next_month(self.TODAY), datetime.date(2026, 10, 1))

    def book(self):
        schedule = Sheet({(row, col + 1): value for row, values in enumerate(self.ROWS, 2)
                          for col, value in enumerate(values)})
        sheets = {'일정': schedule, SHEET: MagicMock()}
        book = MagicMock()
        book.Worksheets.side_effect = sheets.__getitem__
        return book, sheets

    def test_sheet_events_reads_the_schedule_rows(self):
        book, sheets = self.book()
        self.assertEqual(sheet_events(sheets['일정'], 8), collect(self.ROWS))

    def test_calendar_is_redrawn_once_per_change(self):
        book, sheets = self.book()
        events = collect(self.ROWS)
        update_calendar(book, events, self.TODAY)
        sheets[SHEET].Cells.Clear.assert_called_once()
        # A matching signature means the grid is already current.
        book.CustomDocumentProperties.return_value.Value = signature(events, self.TODAY)
        update_calendar(book, events, self.TODAY)
        sheets[SHEET].Cells.Clear.assert_called_once()

    def test_calendar_skipped_before_any_analysis(self):
        book = MagicMock()
        book.Worksheets.side_effect = KeyError('일정')
        update_calendar(book, {}, self.TODAY)   # must not raise
        book.Worksheets.Add.assert_not_called()


class StyleTests(unittest.TestCase):
    def sheets(self):
        result = {}
        for name in HEADERS:
            sheet = MagicMock()
            sheet.ListObjects.Count = 1
            result[name] = sheet
        return result

    def test_rules_use_named_arguments_and_run_once(self):
        book, sheets = MagicMock(), self.sheets()
        apply_style(book, sheets)
        conditions = sheets['우선순위'].Range.return_value.FormatConditions
        conditions.Delete.assert_called_once()
        self.assertEqual(conditions.Add.call_args_list[0], call(Type=2, Formula1='=$B2="긴급"'))
        self.assertEqual(conditions.Add.call_count, 3)
        self.assertEqual(sheets['메일 목록'].ListObjects(1).TableStyle, 'TableStyleLight8')
        # The version marker keeps the same workbook from being restyled every export.
        book.CustomDocumentProperties.return_value.Value = STYLE_VERSION
        apply_style(book, sheets)
        self.assertEqual(conditions.Add.call_count, 3)

    def test_deadline_rules_compare_iso_text(self):
        book, sheets = MagicMock(), self.sheets()
        apply_style(book, sheets)
        formulas = [call.kwargs['Formula1']
                    for call in sheets['일정'].Range.return_value.FormatConditions.Add.call_args_list]
        self.assertEqual(len(formulas), 2)
        self.assertTrue(all('LEFT($E2,10)' in formula for formula in formulas))
        self.assertIn('<TEXT(TODAY(),"yyyy-mm-dd")', formulas[0])

    def test_status_sheet_gets_no_table_or_freeze(self):
        book, sheets = MagicMock(), self.sheets()
        apply_style(book, sheets)
        status = sheets['실행 상태']
        self.assertNotIsInstance(status.ListObjects(1).TableStyle, str)
        status.Range.assert_not_called()     # no conditional formats
        status.Activate.assert_not_called()  # no freeze panes, so the view stays put
        self.assertEqual(status.Columns(1).Font.Bold, True)

    def test_cosmetic_failure_is_reported_not_raised(self):
        status = Sheet({(1, 1): '항목', (1, 2): '값'})
        sheets = {}
        for name, columns in HEADERS.items():
            if name == '실행 상태':
                sheets[name] = status
                continue
            sheet = MagicMock()
            sheet.Range.return_value.Value = (tuple(columns),)
            sheet.ListObjects.Count = 0
            sheets[name] = sheet
        book = MagicMock()
        book.Worksheets.side_effect = sheets.__getitem__
        with patch('mail_assistant.excel.update_calendar', side_effect=OSError('달력 실패')), \
             patch('mail_assistant.excel.update_dashboard', side_effect=OSError('대시보드 실패')), \
             patch('mail_assistant.excel.sheet_events', return_value={}), \
             patch('mail_assistant.excel.mail_rows', return_value={}), \
             patch('mail_assistant.excel.apply_style', side_effect=OSError('서식 실패')):
            Excel._write(book, [], {'안내': '새 메일 1건 수집'})
        written = {status.data[row, 1]: status.data[row, 2]
                   for row in range(2, status.data and max(r for r, _ in status.data) + 1)}
        self.assertEqual(written['안내'], '새 메일 1건 수집')
        self.assertIn('달력 실패', written['일정 달력 갱신'])
        self.assertIn('대시보드 실패', written['대시보드 갱신'])
        self.assertIn('서식 실패', written['서식 적용'])


class DashboardTests(unittest.TestCase):
    TODAY = datetime.date(2026, 9, 11)

    def formulas(self):
        return ' '.join(str(value) for *_, kind, value, _ in blocks(self.TODAY)
                        if str(value).startswith('='))

    def test_formulas_point_at_the_right_columns(self):
        joined = self.formulas()
        self.assertIn("""COUNTIF('메일 목록'!I:I,"미처리")""", joined)    # 처리 상태 is column I
        self.assertIn("""COUNTIF('우선순위'!B:B,"긴급")""", joined)
        self.assertIn("""COUNTIF('답변 초안'!F:F,"검토 전")""", joined)
        self.assertIn("'일정'!$E$2:$E$20000", joined)                   # 마감 is column E

    def test_layout_cells_never_collide(self):
        # The deadline list is filled in later, so reserve its cells up front.
        taken = {(row, column) for row in range(UPCOMING_TOP, UPCOMING_TOP + UPCOMING_ROWS)
                 for column in range(8, 14)}
        for row, first, width, *_ in blocks(self.TODAY):
            for column in range(first, first + width):
                self.assertNotIn((row, column), taken, f'{row},{column} 중복')
                taken.add((row, column))

    def test_upcoming_lists_nearest_deadlines_only(self):
        events = {datetime.date(2026, 9, 9): [Entry('■ 지난 건', '마감', 'm1', 2)],
                  datetime.date(2026, 9, 11): [Entry('■ 오늘 마감', '마감', 'm2', 3)],
                  datetime.date(2026, 9, 14): [Entry('▶ 착수', '시작', 'm3', 4),
                                               Entry('■ 보고서', '마감', 'm3', 5)]}
        self.assertEqual([(day, entry.label) for day, entry in upcoming(events, self.TODAY)],
                         [(datetime.date(2026, 9, 11), '■ 오늘 마감'),
                          (datetime.date(2026, 9, 14), '■ 보고서')])
        self.assertEqual(describe(datetime.date(2026, 9, 9), self.TODAY), '2일 지남')
        self.assertEqual(describe(datetime.date(2026, 9, 11), self.TODAY), '오늘')

    def test_drawn_once_but_deadlines_refresh_every_export(self):
        sheet, book = MagicMock(), MagicMock()
        cells = {}
        # One mock per cell, or every write lands on the same object.
        sheet.Cells.side_effect = lambda row, col: cells.setdefault((row, col), MagicMock())
        book.Worksheets.side_effect = {DASHBOARD: sheet}.__getitem__
        events = {datetime.date(2026, 9, 14): [Entry('■ 보고서', '마감', 'm3', 5)]}
        update_dashboard(book, events, self.TODAY, mail_rows={'m3': 7})
        sheet.Cells.Clear.assert_called_once()
        book.CustomDocumentProperties.return_value.Value = '1'
        update_dashboard(book, events, self.TODAY, mail_rows={'m3': 7})
        sheet.Cells.Clear.assert_called_once()                     # layout stays put
        self.assertEqual(cells[10, 8].Value, '2026-09-14')         # list still refilled
        self.assertEqual(cells[10, 9].Value, '■ 보고서')
        self.assertEqual(cells[10, 12].Value, '3일 뒤')
        self.assertEqual(cells[11, 9].Value, '')                   # empty rows cleared
        target = sheet.Hyperlinks.Add.call_args.kwargs
        self.assertEqual(target['SubAddress'], "'메일 목록'!A7")     # click jumps to the mail
        self.assertEqual(target['Address'], '')


class SampleFileTests(unittest.TestCase):
    def test_sample_workbook_has_every_sheet_and_the_calendar(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest('openpyxl이 없습니다')
        import sample
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'sample.xlsx'
            with patch.object(sys, 'argv', ['sample.py', str(path)]):
                sample.main()
            book = openpyxl.load_workbook(path)
            self.assertEqual(set(HEADERS) | {'일정 달력', '대시보드'}, set(book.sheetnames))
            self.assertEqual(book.sheetnames[0], '대시보드')
            dashboard = book['대시보드']
            self.assertFalse(dashboard.sheet_view.showGridLines)
            self.assertEqual(dashboard['B2'].value, '메일 업무 도우미')
            self.assertTrue(str(dashboard['B6'].value).startswith('=COUNTIF'))
            self.assertIn('견적 회신 마감', str(dashboard['I11'].value))  # nearest deadlines listed
            for name, headers in HEADERS.items():
                self.assertEqual([cell.value for cell in book[name][1]][:len(headers)], list(headers))
            self.assertEqual(book['메일 목록'].max_row, 6)   # header + 5 sample mails
            calendar_sheet = book['일정 달력']
            self.assertEqual(calendar_sheet['A1'].value, f'{sample.TODAY.year}년 {sample.TODAY.month}월')
            grid = '\n'.join(str(cell.value) for row in calendar_sheet.iter_rows() for cell in row if cell.value)
            self.assertIn('견적 회신 마감', grid)
            self.assertIn('지난 마감', grid)


    def test_no_whitespace_run_without_preserve(self):
        """Excel repairs the file when a text run holds only whitespace."""
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest('openpyxl이 없습니다')
        import re
        import zipfile
        import sample
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'sample.xlsx'
            with patch.object(sys, 'argv', ['sample.py', str(path)]):
                sample.main()
            with zipfile.ZipFile(path) as archive:
                parts = [name for name in archive.namelist() if name.startswith('xl/worksheets/sheet')]
                self.assertTrue(parts)
                for name in parts:
                    body = archive.read(name).decode('utf-8')
                    for tag, text in re.findall(r'<t( [^>]*)?>(.*?)</t>', body, re.S):
                        if text and text != text.strip():
                            self.assertIn('xml:space="preserve"', tag, f'{name}: {text!r}')


class WorkerTests(unittest.TestCase):
    def test_export_failure_keeps_result_and_reports_detail(self):
        from mail_assistant.worker import run
        with tempfile.TemporaryDirectory() as folder:
            directory = Path(folder)
            store = Store(directory / 'mail.db')
            ident = store.add(account_key(CONFIG), 'new', mail())
            store.analyzed(ident, parse_mail(mail()), RESULT)
            store.db.close()
            stop = MagicMock()
            # One pass through the loop. The briefing check at the end of the cycle
            # asks as well, which is why this is three and not two.
            stop.is_set.side_effect = [False, False, True]
            messages = []
            with patch('mail_assistant.worker.fetch_mail', return_value='새 메일 0건 수집'), patch('mail_assistant.worker.read_password', return_value='test'), patch('mail_assistant.worker.write_briefing'), patch('mail_assistant.worker.Excel') as excel:
                excel.return_value.update.side_effect = ExcelUpdateError('표 만들기', RuntimeError('invalid argument'))
                run({**CONFIG, 'workbook': str(directory / 'test.xlsx'), 'interval': 180}, directory, stop, messages.append)
            self.assertIn('표 만들기', messages[-1])
            self.assertIn('invalid argument', messages[-1])
            snapshot = json.loads((directory / 'status.json').read_text(encoding='utf-8'))
            self.assertIn('invalid argument', snapshot['message'])
            store = Store(directory / 'mail.db')
            self.assertEqual(len(store.unexported(account_key(CONFIG))), 1)
            store.db.close()


class AnalyzingMarkerTests(unittest.TestCase):
    """The worker's own end of '분석 중': set before the call, gone after every exit."""

    def run_once(self, directory, answer):
        """`answer` is analyze_one(sent, config)'s: the worker prepares for itself now."""
        from mail_assistant.worker import run
        stop = MagicMock()
        # One whole cycle: counting is_set() calls breaks the moment the loop grows a
        # branch, so the end of the cycle — the interval wait — is what stops it.
        stop.is_set.return_value = False
        stop.wait.side_effect = lambda *_: stop.is_set.configure_mock(return_value=True)
        messages = []
        # write_briefing은 반드시 막는다: 분석에 성공하면 대기열이 비고 브리핑이 due가
        # 되므로, 패치하지 않은 테스트는 진짜 Codex를 부르고 진짜 사용량을 쓴다.
        with patch('mail_assistant.worker.fetch_mail', return_value='새 메일 0건 수집'), \
                patch('mail_assistant.worker.read_password', return_value='test'), \
                patch('mail_assistant.worker.check_login'), \
                patch('mail_assistant.worker.analyze_one', side_effect=answer), \
                patch('mail_assistant.worker.write_briefing'), \
                patch('mail_assistant.worker.Excel'):
            run({**CONFIG, 'workbook': str(directory / 'test.xlsx'), 'interval': 180},
                directory, stop, messages.append)
        return messages

    def prepared(self, folder):
        directory = Path(folder)
        store = Store(directory / 'mail.db')
        ident = store.add(account_key(CONFIG), 'new', mail())
        store.db.close()
        return directory, ident

    def test_the_mail_being_analysed_carries_the_marker_while_it_is(self):
        with tempfile.TemporaryDirectory() as folder:
            directory, ident = self.prepared(folder)
            seen = []

            def answer(sent, config=None):
                store = Store(directory / 'mail.db')
                seen.append(store.analyzing(account_key(CONFIG)))
                store.db.close()
                return RESULT

            self.run_once(directory, answer)
            self.assertEqual(seen, [ident])
            store = Store(directory / 'mail.db')
            self.assertEqual(store.analyzing(account_key(CONFIG)), '')
            store.db.close()

    def test_an_analysis_that_fails_still_gives_the_marker_back(self):
        """Otherwise the mail reads '분석 중' for ever and 다시 분석 refuses to touch it."""
        with tempfile.TemporaryDirectory() as folder:
            directory, ident = self.prepared(folder)
            self.run_once(directory, RuntimeError('한도 초과'))
            store = Store(directory / 'mail.db')
            self.assertEqual(store.analyzing(account_key(CONFIG)), '')
            self.assertEqual(state_of(store.detail(ident)), '1회 실패')
            store.db.close()

    def test_a_marker_left_by_a_crash_is_cleared_at_the_next_start(self):
        with tempfile.TemporaryDirectory() as folder:
            directory, ident = self.prepared(folder)
            store = Store(directory / 'mail.db')
            store.mark_analyzing(account_key(CONFIG), ident)
            store.db.close()
            self.run_once(directory, lambda sent, config=None: RESULT)
            store = Store(directory / 'mail.db')
            self.assertEqual(store.analyzing(account_key(CONFIG)), '')
            store.db.close()


class BodyLimitTests(unittest.TestCase):
    """본문 상한은 일부러 있는 것이고, 거절의 *종류*가 무엇이냐가 요점이다."""

    def oversize(self):
        msg = EmailMessage()
        msg['Subject'] = '긴 메일'
        msg['From'] = 'sender@example.com'
        msg['Date'] = 'Thu, 10 Sep 2026 10:00:00 +0900'
        msg.set_content('가' * (BODY_LIMIT + 1))
        return msg.as_bytes()

    def analyzed(self, raw):
        """analyze() against a fake Codex, returning (what was sent, parsed, result)."""
        seen = {}

        def execute(command, **kwargs):
            seen['input'] = kwargs['input']
            Path(command[command.index('-o') + 1]).write_text(json.dumps(RESULT),
                                                              encoding='utf-8')
            return subprocess.CompletedProcess(command, 0)

        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            store.add(account_key(CONFIG), 'one', raw)
            row = store.pending(account_key(CONFIG), time.time())[0]
            with patch('mail_assistant.services.codex_command', return_value=['codex']), \
                 patch('mail_assistant.services.subprocess.run', side_effect=execute):
                parsed, result = analyze(row, CONFIG)
            store.db.close()
        text = seen['input']
        return text, json.loads(text[text.index('{'):]), parsed, result

    def test_only_the_first_part_is_sent_and_the_model_is_told_so(self):
        """잘라 놓고 아무 말이 없는 것이 이 상한의 유일한 위험이다."""
        text, payload, _, _ = self.analyzed(self.oversize())
        self.assertEqual(len(payload['body']), BODY_LIMIT)
        self.assertIn('clipped', payload)
        # 뒤쪽을 보지 못한 모델이 '일정 없음'이라고 쓰면 잘린 절반이 조용히 사라진다.
        self.assertIn('단정하지 말고', text)
        self.assertIn('needs_review=true', text)

    def test_an_ordinary_mail_is_told_none_of_that(self):
        text, payload, parsed, result = self.analyzed(mail())
        self.assertNotIn('앞부분만 잘라', text)
        self.assertNotIn('clipped', payload)
        self.assertNotIn('clipped', parsed)
        self.assertEqual(result, RESULT)

    def test_what_is_stored_is_the_whole_body_and_the_original_length(self):
        """원문은 전체가 남아야 한다 — 화면이 보여주는 것이 이 parsed이기 때문이다."""
        _, _, parsed, _ = self.analyzed(self.oversize())
        self.assertGreater(len(parsed['body']), BODY_LIMIT)
        self.assertEqual(parsed['clipped'], len(parsed['body']))

    def test_it_is_not_the_kind_of_failure_a_retry_fixes(self):
        """Unanalyzable subclasses RuntimeError, so an older except still catches it —
        but the worker's own branch has to come first, which is what this pins."""
        self.assertTrue(issubclass(Unanalyzable, RuntimeError))

    def test_a_mail_with_nothing_in_it_at_all_is_put_down(self):
        message = EmailMessage()
        message['From'] = 'sender@example.com'
        message['Date'] = 'Thu, 10 Sep 2026 10:00:00 +0900'
        message.set_content('   ')
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            store.add(account_key(CONFIG), 'empty', message.as_bytes())
            row = store.pending(account_key(CONFIG), time.time())[0]
            with patch('mail_assistant.services.subprocess.run') as run:
                with self.assertRaises(Unanalyzable):
                    analyze(row, CONFIG)
            run.assert_not_called()             # Codex was never asked
            store.db.close()

    def test_a_subject_with_no_body_is_still_worth_analysing(self):
        """'회의 12시'만 제목에 있는 메일을 거절하면 그것이 회귀다."""
        message = EmailMessage()
        message['Subject'] = '내일 12시 회의'
        message['From'] = 'sender@example.com'
        message['Date'] = 'Thu, 10 Sep 2026 10:00:00 +0900'
        message.set_content('')
        self.assertEqual(self.analyzed(message.as_bytes())[3], RESULT)

    def test_a_body_inside_the_limit_is_analysed_as_before(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            store.add(account_key(CONFIG), 'fine', mail())
            row = store.pending(account_key(CONFIG), time.time())[0]

            def execute(command, **kwargs):
                Path(command[command.index('-o') + 1]).write_text(json.dumps(RESULT),
                                                                  encoding='utf-8')
                return subprocess.CompletedProcess(command, 0)

            with patch('mail_assistant.services.codex_command', return_value=['codex']), \
                 patch('mail_assistant.services.subprocess.run', side_effect=execute):
                _, result = analyze(row, CONFIG)
            self.assertEqual(result, RESULT)
            store.db.close()


class UnanalyzableTests(unittest.TestCase):
    """다시 물어도 같은 답이 나오는 실패는, 다시 묻지 않는다."""

    def run_once(self, directory, reports, prepare=None, answer=None, asked=None):
        """Unanalyzable은 이제 prepare()에서 나온다 — worker가 Codex를 부르기 *전*이다.

        That is the split this class is about, so the two are patched separately: what
        refuses a mail and what fails while analysing it are no longer the same call.
        """
        from mail_assistant.worker import run
        stop = MagicMock()
        stop.is_set.return_value = False
        stop.wait.side_effect = lambda *_: stop.is_set.configure_mock(return_value=True)
        messages = []

        def analysed(sent, config=None):
            if asked is not None:
                asked.append(sent.get('subject'))
            if answer is None:
                return RESULT
            return answer(sent, config)

        def batched(items, config=None):
            return {ident: analysed(sent, config) for ident, sent in items}

        with patch('mail_assistant.worker.fetch_mail', return_value='새 메일 0건 수집'), \
                patch('mail_assistant.worker.read_password', return_value='test'), \
                patch('mail_assistant.worker.check_login'), \
                patch('mail_assistant.worker.analyze_one', side_effect=analysed), \
                patch('mail_assistant.worker.analyze_many', side_effect=batched), \
                patch('mail_assistant.worker.write_briefing'), \
                patch('mail_assistant.worker.report',
                      side_effect=lambda *args, **kw: reports.append(args)), \
                patch('mail_assistant.worker.Excel'):
            if prepare is not None:
                with patch('mail_assistant.worker.prepare', side_effect=prepare):
                    run({**CONFIG, 'workbook': str(directory / 'test.xlsx'), 'interval': 180},
                        directory, stop, messages.append)
                return messages
            run({**CONFIG, 'workbook': str(directory / 'test.xlsx'), 'interval': 180},
                directory, stop, messages.append)
        return messages

    # 다시 물어도 같은 답이 나오는 실패 하나, 두 통 모두에.
    REFUSAL = '본문을 읽지 못했어요. 원문은 그대로 보관돼요.'

    def refuse(self, row):
        raise Unanalyzable(self.REFUSAL)

    def two_mails(self, folder):
        directory = Path(folder)
        store = Store(directory / 'mail.db')
        first = store.add(account_key(CONFIG), 'long', mail())
        second = store.add(account_key(CONFIG), 'next', mail())
        store.db.close()
        return directory, first, second

    def test_an_oversize_body_is_put_down_rather_than_asked_again(self):
        with tempfile.TemporaryDirectory() as folder:
            directory, first, _ = self.two_mails(folder)
            reports, asked = [], []
            self.run_once(directory, reports, prepare=self.refuse, asked=asked)
            store = Store(directory / 'mail.db')
            row = store.detail(first)
            # No clock brings it back; 다시 분석 is the only way in.
            self.assertEqual(row['retry_at'], NO_RETRY)
            self.assertEqual(store.pending(account_key(CONFIG), time.time()), [])
            # And the mail itself says why, rather than '로그인·한도·본문 형식을 확인하세요'.
            self.assertIn(self.REFUSAL, row['error'])
            # Codex was never asked: prepare() refuses before a slot is taken at all.
            self.assertEqual(asked, [])
            store.db.close()

    def test_it_never_reaches_the_crash_channel(self):
        """A report that arrives again every hour for ever is one nobody reads."""
        with tempfile.TemporaryDirectory() as folder:
            directory, _, _ = self.two_mails(folder)
            reports = []
            self.run_once(directory, reports, prepare=self.refuse)
            self.assertEqual([stage for stage, *_ in reports], [])

    def test_the_queue_behind_it_carries_on_in_the_same_cycle(self):
        """The retry path breaks the loop and backs every mail off by up to an hour;
        one mail that can never be analysed must not cost the mailbox that."""
        with tempfile.TemporaryDirectory() as folder:
            directory, first, second = self.two_mails(folder)
            reports, asked = [], []
            from mail_assistant.services import prepare as real_prepare

            def prepare(row):
                if row['id'] == first:
                    raise Unanalyzable(self.REFUSAL)
                return real_prepare(row)

            self.run_once(directory, reports, prepare=prepare, asked=asked)
            store = Store(directory / 'mail.db')
            self.assertIsNotNone(store.detail(second)['result'])
            # The one that could never be analysed never reached Codex, and the one
            # behind it was analysed in the same cycle rather than an hour later.
            self.assertEqual(len(asked), 1)
            self.assertEqual(store.detail(first)['retry_at'], NO_RETRY)
            store.db.close()

    def test_다시_분석_is_the_way_back_in(self):
        with tempfile.TemporaryDirectory() as folder:
            directory, first, _ = self.two_mails(folder)
            self.run_once(directory, [], prepare=self.refuse)
            store = Store(directory / 'mail.db')
            self.assertEqual(store.reset([first]), 1)
            self.assertEqual([row['id'] for row in
                              store.pending(account_key(CONFIG), time.time())], [first])
            store.db.close()

    def test_a_failure_a_retry_could_fix_still_backs_off_and_reports(self):
        """The split is the point: a rate limit and a 60,000-character body are not
        the same thing, and before this they were handled as if they were."""
        with tempfile.TemporaryDirectory() as folder:
            directory, first, _ = self.two_mails(folder)
            reports = []
            self.run_once(directory, reports, answer=lambda sent, config=None: (
                _ for _ in ()).throw(RuntimeError('rate limited')))
            store = Store(directory / 'mail.db')
            row = store.detail(first)
            self.assertLess(row['retry_at'], NO_RETRY)
            self.assertGreater(row['retry_at'], time.time())
            store.db.close()
            self.assertEqual([stage for stage, *_ in reports], ['분석 실패'])


class SqueezeTests(unittest.TestCase):
    """Codex로 가는 사본에서만 공백을 접는다. 저장되는 원문은 화면이 읽는 그것이다."""

    HTML = ('<div style="padding:20px">\n      <p>\n        안녕하세요.\n      </p>\n'
            '      <p>\n        견적서를 보냅니다.\n      </p>\n    </div>')

    def test_the_indentation_html_mail_arrives_with_is_dropped(self):
        text = text_of_html(self.HTML)
        self.assertGreater(len(text), len(squeeze_body(text)) * 1.5)
        self.assertEqual(squeeze_body(text).splitlines()[0], '안녕하세요.')

    def test_a_run_of_blank_lines_becomes_one(self):
        self.assertEqual(squeeze_body('가\n\n\n\n나'), '가\n\n나')

    def test_a_gap_inside_a_line_is_left_alone(self):
        """평문 메일이 표를 그리는 자리다. 그것까지 접으면 숫자가 어느 칸의 것인지 사라진다."""
        self.assertEqual(squeeze_body('품번      수량\nA26090135      10'),
                         '품번      수량\nA26090135      10')

    def test_what_is_stored_still_has_every_space_in_it(self):
        message = EmailMessage()
        message['Subject'] = 'HTML 메일'
        message['From'] = 'sender@example.com'
        message['Date'] = 'Thu, 10 Sep 2026 10:00:00 +0900'
        message.set_content(self.HTML, subtype='html')
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            store.add(account_key(CONFIG), 'html', message.as_bytes())
            row = store.pending(account_key(CONFIG), time.time())[0]
            parsed, sent = prepare(row)
            store.db.close()
        self.assertGreater(len(parsed['body']), len(sent['body']))
        self.assertIn('\n      ', parsed['body'])

    def test_whitespace_alone_no_longer_costs_a_mail_its_ending(self):
        """상한을 먹는 것이 들여쓰기이면, 잘리는 것은 진짜 내용이다."""
        body = ('가' * 40 + ' ' * 200 + '\n') * 300        # 72,000자, 내용은 12,000자
        message = EmailMessage()
        message['Subject'] = '긴 메일'
        message['From'] = 'sender@example.com'
        message['Date'] = 'Thu, 10 Sep 2026 10:00:00 +0900'
        message.set_content(body)
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            store.add(account_key(CONFIG), 'padded', message.as_bytes())
            row = store.pending(account_key(CONFIG), time.time())[0]
            parsed, sent = prepare(row)
            store.db.close()
        self.assertGreater(len(parsed['body']), BODY_LIMIT)
        self.assertNotIn('clipped', sent)               # 접고 나니 상한 아래다
        self.assertNotIn('clipped', parsed)


class GroupMailsTests(unittest.TestCase):
    """무엇을 묶고 무엇을 혼자 보내는가. codex exec 한 번의 고정비를 나누는 규칙이다."""

    def entries(self, *specs):
        return [{'id': ident, 'attempts': attempts, 'size': size}
                for ident, attempts, size in specs]

    def test_new_mail_travels_together(self):
        rows = self.entries(*[(f'm{i}', 0, 100) for i in range(3)])
        self.assertEqual(group_mails(rows), [['m0', 'm1', 'm2']])

    def test_the_count_bounds_a_group(self):
        rows = self.entries(*[(f'm{i}', 0, 10) for i in range(BATCH_MAILS + 2)])
        groups = group_mails(rows)
        self.assertEqual(len(groups[0]), BATCH_MAILS)
        self.assertEqual(len(groups[1]), 2)

    def test_the_character_budget_bounds_a_group(self):
        rows = self.entries(('a', 0, BATCH_CHARS - 10), ('b', 0, 100))
        self.assertEqual(group_mails(rows), [['a'], ['b']])

    def test_one_mail_bigger_than_the_budget_goes_by_itself(self):
        rows = self.entries(('big', 0, BATCH_CHARS * 2), ('b', 0, 10))
        self.assertEqual(group_mails(rows), [['big'], ['b']])

    def test_a_mail_that_has_failed_before_goes_alone(self):
        """묶음이 통째로 실패했을 때 범인이 스스로 드러나게 하는 것이 이 규칙의 전부다."""
        rows = self.entries(('fresh', 0, 10), ('burnt', 2, 10), ('other', 0, 10))
        self.assertEqual(group_mails(rows), [['fresh'], ['burnt'], ['other']])

    def test_the_order_mail_arrived_in_survives(self):
        rows = self.entries(('a', 0, 10), ('b', 0, 10), ('c', 1, 10), ('d', 0, 10))
        self.assertEqual(group_mails(rows), [['a', 'b'], ['c'], ['d']])


class BatchAnalysisTests(unittest.TestCase):
    """여러 통이 codex exec 한 번으로 간다. 나머지는 답을 어느 메일의 것으로 돌리느냐다."""

    def sent(self, ident, subject):
        return (ident, {'subject': subject, 'body': '본문', 'date': '', 'attachments': []})

    def call(self, answer):
        """analyze_many() against a fake Codex, returning (what was sent, what came back)."""
        seen = {}

        def execute(command, **kwargs):
            seen['input'] = kwargs['input']
            Path(command[command.index('-o') + 1]).write_text(
                json.dumps(answer, ensure_ascii=False), encoding='utf-8')
            return subprocess.CompletedProcess(command, 0)

        with patch('mail_assistant.services.codex_command', return_value=['codex']), \
             patch('mail_assistant.services.subprocess.run', side_effect=execute) as run:
            found = analyze_many([self.sent('id-a', '첫 메일'), self.sent('id-b', '둘째')],
                                 CONFIG)
        return seen['input'], found, run

    def answer(self, *idents):
        return {'results': [{'mail_id': ident, **RESULT} for ident in idents]}

    def test_two_mails_are_one_codex_process(self):
        text, found, run = self.call(self.answer('id-a', 'id-b'))
        self.assertEqual(run.call_count, 1)
        self.assertEqual(set(found), {'id-a', 'id-b'})
        self.assertIn('첫 메일', text)
        self.assertIn('둘째', text)

    def test_the_model_is_told_not_to_mix_them(self):
        """한 메일의 마감이 옆 메일의 근거가 되면 달력과 엑셀까지 틀린 채로 흘러간다."""
        text, _, _ = self.call(self.answer('id-a', 'id-b'))
        self.assertIn('독립적으로', text)
        self.assertIn('다른 메일의 근거로 쓰지 마세요', text)

    def test_mail_id_never_reaches_the_stored_result(self):
        _, found, _ = self.call(self.answer('id-a', 'id-b'))
        self.assertEqual(found['id-a'], RESULT)

    def test_an_id_the_model_invented_is_dropped(self):
        _, found, _ = self.call(self.answer('id-a', 'id-made-up'))
        self.assertEqual(set(found), {'id-a'})

    def test_a_mail_left_out_of_the_answer_is_simply_absent(self):
        """부르는 쪽이 그것을 보고 다음 주기에 한 통씩 다시 보낸다."""
        _, found, _ = self.call(self.answer('id-a'))
        self.assertEqual(set(found), {'id-a'})

    def test_a_broken_date_costs_only_its_own_mail(self):
        broken = {'results': [{'mail_id': 'id-a', **RESULT},
                              {'mail_id': 'id-b', **RESULT,
                               'events': [{'title': '회신', 'start': '', 'deadline': '다음 주',
                                           'evidence': '', 'needs_review': False}]}]}
        _, found, _ = self.call(broken)
        self.assertEqual(set(found), {'id-a'})


class WorkerBatchTests(unittest.TestCase):
    """한 주기에 Codex를 몇 번 부르는가, 그리고 묶음이 깨졌을 때 무엇이 혼자 가는가."""

    def cycle(self, directory, one=None, many=None, reports=None):
        """한 주기. 부른 것을 그대로 돌려주므로 호출 횟수와 묶음 크기를 셀 수 있다."""
        from mail_assistant.worker import run
        stop = MagicMock()
        stop.is_set.return_value = False
        stop.wait.side_effect = lambda *_: stop.is_set.configure_mock(return_value=True)
        calls, messages = [], []

        def solo(sent, config=None):
            calls.append([sent.get('subject')])
            return RESULT if one is None else one(sent)

        def batch(items, config=None):
            calls.append([sent.get('subject') for _, sent in items])
            if many is not None:
                return many(items)
            return {ident: RESULT for ident, _ in items}

        with patch('mail_assistant.worker.fetch_mail', return_value='새 메일 0건 수집'), \
                patch('mail_assistant.worker.read_password', return_value='test'), \
                patch('mail_assistant.worker.check_login'), \
                patch('mail_assistant.worker.analyze_one', side_effect=solo), \
                patch('mail_assistant.worker.analyze_many', side_effect=batch), \
                patch('mail_assistant.worker.write_briefing'), \
                patch('mail_assistant.worker.report',
                      side_effect=lambda *args, **kw: (reports if reports is not None
                                                       else []).append(args)), \
                patch('mail_assistant.worker.Excel'):
            run({**CONFIG, 'workbook': str(directory / 'test.xlsx'), 'interval': 180},
                directory, stop, messages.append)
        return calls, messages

    def stocked(self, folder, count):
        directory = Path(folder)
        store = Store(directory / 'mail.db')
        idents = []
        for index in range(count):
            message = EmailMessage()
            message['Subject'] = f'메일 {index}'
            message['From'] = 'sender@example.com'
            message['Date'] = 'Thu, 10 Sep 2026 10:00:00 +0900'
            message.set_content('9월 11일까지 회신 부탁드립니다.')
            idents.append(store.add(account_key(CONFIG), f'uid-{index}',
                                    message.as_bytes()))
        store.db.close()
        return directory, idents

    def test_a_cycle_of_new_mail_is_one_codex_call(self):
        """이 변경의 전부다. 다섯 통이면 다섯 번이 아니라 한 번."""
        with tempfile.TemporaryDirectory() as folder:
            directory, idents = self.stocked(folder, BATCH_MAILS)
            calls, _ = self.cycle(directory)
            self.assertEqual(len(calls), 1)
            self.assertEqual(len(calls[0]), BATCH_MAILS)
            store = Store(directory / 'mail.db')
            for ident in idents:
                self.assertIsNotNone(store.detail(ident)['result'])
            store.db.close()

    def test_every_mail_in_the_batch_carries_분석_중_while_it_runs(self):
        with tempfile.TemporaryDirectory() as folder:
            directory, idents = self.stocked(folder, 3)
            seen = []

            def many(items):
                store = Store(directory / 'mail.db')
                seen.extend(row['id'] for row in store.page(account_key(CONFIG))
                            if row['analyzing'])
                store.db.close()
                return {ident: RESULT for ident, _ in items}

            self.cycle(directory, many=many)
            self.assertEqual(sorted(seen), sorted(idents))
            store = Store(directory / 'mail.db')
            self.assertEqual(store.analyzing(account_key(CONFIG)), '')
            store.db.close()

    def test_a_mail_the_answer_left_out_comes_back_alone(self):
        """빠진 한 통 때문에 나머지를 버리지 않는다. 그 한 통만 다음 주기에 혼자 간다."""
        with tempfile.TemporaryDirectory() as folder:
            directory, idents = self.stocked(folder, 3)
            missing = idents[1]
            self.cycle(directory, many=lambda items: {ident: RESULT for ident, _ in items
                                                      if ident != missing})
            store = Store(directory / 'mail.db')
            row = store.detail(missing)
            self.assertIsNone(row['result'])
            self.assertEqual(row['attempts'], 1)
            # 전체 백오프가 아니다: 다음 주기에 바로 다시 간다.
            self.assertEqual(row['retry_at'], 0)
            self.assertIsNotNone(store.detail(idents[0])['result'])
            store.db.close()
            # 그리고 혼자 간다 — group_mails()가 attempts로 가른다.
            calls, _ = self.cycle(directory)
            self.assertEqual(calls, [['메일 1']])

    def test_a_batch_that_fails_sends_every_one_of_them_alone_next_time(self):
        with tempfile.TemporaryDirectory() as folder:
            directory, idents = self.stocked(folder, 3)
            self.cycle(directory, many=lambda items: (_ for _ in ()).throw(
                RuntimeError('rate limited')))
            store = Store(directory / 'mail.db')
            rows = [store.detail(ident) for ident in idents]
            self.assertEqual([row['attempts'] for row in rows], [1, 1, 1])
            # 아무도 포기당하지 않는다: 묶음의 실패는 그 메일의 잘못이 아니다.
            self.assertTrue(all(row['retry_at'] < NO_RETRY for row in rows))
            for row in rows:
                store.db.execute('UPDATE mail SET retry_at=0 WHERE id=?', (row['id'],))
            store.db.commit()
            store.db.close()
            calls, _ = self.cycle(directory)
            self.assertEqual(calls, [['메일 0'], ['메일 1'], ['메일 2']])

    def test_one_wait_line_rather_than_one_for_each_mail(self):
        with tempfile.TemporaryDirectory() as folder:
            directory, _ = self.stocked(folder, 3)
            _, messages = self.cycle(directory, many=lambda items: (_ for _ in ()).throw(
                RuntimeError('rate limited')))
            waits = [line for line in ' / '.join(messages).split(' / ')
                     if line.startswith('분석 대기')]
            self.assertEqual(len(waits), 1)
            self.assertIn('3건', waits[0])


class GiveUpTests(unittest.TestCase):
    """영원한 재시도를 끊되, 한도 소진을 '메일 전체 포기'로 바꾸지는 않는다."""

    # 지난 실패의 시각. 진짜 순서에서는 실패와 그 뒤의 성공 사이에 백오프가 5분에서 한
    # 시간 흐르지만, 테스트는 전부 한 틱 안에서 끝난다 — 그리고 Windows의 시계는 약
    # 15ms마다 움직이므로 failed_at과 analyzed_at이 같은 문자열이 되어, 실패한 곳은
    # 리눅스가 아니라 CI의 Windows 레그였다. received가 rowid tie-break를 필요로 하는
    # 것과 같은 시계다. 그래서 지난 실패를 분명히 앞에 둔다.
    BEFORE = '2020-01-01T00:00:00+00:00'

    def workspace(self):
        """tmpdir도 테스트 정리에 맡긴다. 순서가 이 헬퍼의 전부다.

        `with tempfile.TemporaryDirectory()` 안에서 addCleanup으로 연결을 닫으면 순서가
        뒤집힌다 — 정리는 with 블록이 끝난 *뒤*에 돌고, Windows는 그때 이미 열린 mail.db를
        지우려다 WinError 32를 낸 뒤다. 폴더를 먼저 등록하면 LIFO로 연결이 먼저 닫힌다.
        """
        holder = tempfile.TemporaryDirectory()
        self.addCleanup(holder.cleanup)
        return Path(holder.name)

    def opened(self, directory):
        """열고 나서 반드시 닫는다. 단언이 먼저 터지면 Windows가 mail.db를 붙들고 있다."""
        store = Store(directory / 'mail.db')
        self.addCleanup(store.db.close)
        return store

    def prepared(self, attempts):
        """실패를 `attempts`번 쌓은 메일 하나, 그리고 두 번째 메일 하나."""
        directory = self.workspace()
        store = self.opened(directory)
        first = store.add(account_key(CONFIG), 'stuck', mail())
        second = store.add(account_key(CONFIG), 'other', mail())
        for _ in range(attempts):
            store.failed(first, '분석 실패', 0)
        if attempts:
            store.db.execute('UPDATE mail SET failed_at=? WHERE id=?', (self.BEFORE, first))
            store.db.commit()
        return directory, store, first, second

    def cycle(self, directory):
        from mail_assistant.worker import run
        stop = MagicMock()
        stop.is_set.return_value = False
        stop.wait.side_effect = lambda *_: stop.is_set.configure_mock(return_value=True)
        messages = []
        with patch('mail_assistant.worker.fetch_mail', return_value='새 메일 0건 수집'), \
                patch('mail_assistant.worker.read_password', return_value='test'), \
                patch('mail_assistant.worker.check_login'), \
                patch('mail_assistant.worker.analyze_one',
                      side_effect=RuntimeError('rate limited')), \
                patch('mail_assistant.worker.write_briefing'), \
                patch('mail_assistant.worker.report'), \
                patch('mail_assistant.worker.Excel'):
            run({**CONFIG, 'workbook': str(directory / 'test.xlsx'), 'interval': 180},
                directory, stop, messages.append)
        return messages

    def test_it_is_put_down_once_codex_is_known_to_be_working(self):
        from mail_assistant.worker import MAX_ATTEMPTS
        directory, store, first, second = self.prepared(MAX_ATTEMPTS - 1)
        # 이 메일이 마지막으로 실패한 *뒤에* 다른 메일이 분석에 성공했다.
        store.analyzed(second, parse_mail(mail()), RESULT)
        store.db.close()          # 워커가 같은 파일을 연다
        messages = self.cycle(directory)
        store = self.opened(directory)
        row = store.detail(first)
        self.assertEqual(row['retry_at'], NO_RETRY)
        self.assertIn('더 시도하지 않아요', row['error'])
        self.assertIn('분석 포기', ' / '.join(messages))

    def test_a_quota_outage_never_puts_the_queue_down(self):
        """아무것도 성공하지 못했다면 한도나 로그인 쪽이고, 그것은 메일의 잘못이 아니다."""
        from mail_assistant.worker import MAX_ATTEMPTS
        directory, store, first, _ = self.prepared(MAX_ATTEMPTS - 1)
        store.db.close()
        self.cycle(directory)
        store = self.opened(directory)
        row = store.detail(first)
        self.assertLess(row['retry_at'], NO_RETRY)
        self.assertGreater(row['retry_at'], time.time())

    def test_다시_분석_is_still_the_way_back_in(self):
        from mail_assistant.worker import MAX_ATTEMPTS
        directory, store, first, second = self.prepared(MAX_ATTEMPTS - 1)
        store.analyzed(second, parse_mail(mail()), RESULT)
        store.db.close()
        self.cycle(directory)
        store = self.opened(directory)
        self.assertEqual(store.reset([first]), 1)
        row = store.detail(first)
        self.assertEqual(row['attempts'], 0)
        self.assertEqual(row['failed_at'], '')
        self.assertEqual([one['id'] for one in
                          store.pending(account_key(CONFIG), time.time())], [first])

    def test_analyzed_since_is_what_tells_the_two_apart(self):
        directory, store, first, second = self.prepared(1)
        stamp = store.detail(first)['failed_at']
        self.assertEqual(stamp, self.BEFORE)
        self.assertFalse(store.analyzed_since(account_key(CONFIG), stamp))
        store.analyzed(second, parse_mail(mail()), RESULT)
        self.assertTrue(store.analyzed_since(account_key(CONFIG), stamp))
        # 한 번도 실패한 적 없는 메일에는 물어볼 것이 없다.
        self.assertFalse(store.analyzed_since(account_key(CONFIG), ''))


class ConsoleTests(unittest.TestCase):
    """Korean print() into a redirected pipe is the ANSI codepage on Windows."""

    def test_reconfigures_a_narrow_stream(self):
        import io
        from mail_assistant.console import use_utf8
        narrow = io.TextIOWrapper(io.BytesIO(), encoding='cp1252')
        with patch.object(sys, 'stdout', narrow), patch.object(sys, 'stderr', narrow):
            use_utf8()
            print('예시 파일 생성', file=sys.stdout)
        self.assertEqual(narrow.encoding, 'utf-8')

    def test_is_safe_when_a_stream_cannot_be_reconfigured(self):
        from mail_assistant.console import use_utf8
        with patch.object(sys, 'stdout', object()), patch.object(sys, 'stderr', None):
            use_utf8()


class SkipRuleTests(unittest.TestCase):
    """분석 전에 내려놓을 메일을 가리는 규칙 — 순수 함수라 여기서 전부 재진다.

    이 규칙이 category를 보지 않는 이유가 이 모듈이 있는 이유다: category는 분석이
    *만드는* 값이라, '공지면 건너뛴다'는 이미 Codex를 한 번 부른 뒤에야 할 수 있는 말이다.
    """

    def parsed(self, subject='견적서 회신 부탁드립니다',
               sender='김도현 <dhkim@daesung.co.kr>', **bulk):
        return {'subject': subject, 'sender': sender, 'body': '본문',
                'bulk': {name: bulk.get(name.replace('-', '_').lower(), '')
                         for name in BULK_HEADERS}}

    def test_an_ordinary_business_mail_is_analysed(self):
        self.assertEqual(skip_reason(self.parsed()), '')

    def test_a_newsletter_says_so_in_its_own_headers(self):
        """낱말 짐작이 아니라 보낸 쪽이 스스로 밝힌 사실이라는 점이 이 신호의 값이다."""
        self.assertEqual(skip_reason(self.parsed(list_unsubscribe='<https://x/u>')),
                         LETTER_REASON)
        self.assertEqual(skip_reason(self.parsed(list_id='<news.x.com>')), LETTER_REASON)

    def test_bulk_and_auto_senders_are_put_down_too(self):
        self.assertEqual(skip_reason(self.parsed(precedence='bulk')), BULK_REASON)
        self.assertEqual(skip_reason(self.parsed(auto_submitted='auto-generated')),
                         AUTO_REASON)
        # 'Precedence: first-class'는 무더기라는 뜻이 아니다.
        self.assertEqual(skip_reason(self.parsed(precedence='first-class')), '')

    def test_the_advertising_mark_is_the_one_word_worth_matching(self):
        """정보통신망법이 제목에 요구하는 표기라, 붙이는 쪽이 법 때문에 붙인다."""
        for subject in ('[광고] 9월 특가', '(광고) 세미나 안내', '[AD] Sale'):
            self.assertEqual(skip_reason(self.parsed(subject=subject)), AD_REASON)

    def test_a_notice_a_person_wrote_is_never_skipped(self):
        """거르는 것은 광고와 뉴스레터이지 공지가 아니다. 이 둘은 시드에 있는 실제
        제목이고, 둘 다 마감과 우선순위를 달고 나온다 — 이 앱이 있는 이유에 가깝다."""
        for subject in ('10월 정기 점검 일정 안내', '단가 인상 안내의 건',
                        '9월 정산 내역 확인 요청'):
            self.assertEqual(skip_reason(self.parsed(subject=subject)), '')

    def test_my_own_company_is_never_filtered_whatever_it_attaches(self):
        """사내 그룹웨어·인사 공지가 수신거부 헤더를 다는 일이 실제로 있고, 그때 걸러
        버리면 놓치면 안 되는 바로 그 메일을 놓친다. 오탐 하나가 정탐 백 개보다 비싸다."""
        inside = self.parsed(sender='인사팀 <hr@monitorapp.com>',
                             list_unsubscribe='<https://x/u>')
        self.assertEqual(skip_reason(inside, own_domain='monitorapp.com'), '')
        # 같은 메일이라도 바깥에서 왔으면 걸린다.
        self.assertEqual(skip_reason(inside, own_domain='other.co.kr'), LETTER_REASON)

    def test_the_switch_is_on_until_somebody_turns_it_off(self):
        """값이 없으면 켜짐 — 이 기능이 있는 이유가 곧 기본값이다."""
        self.assertTrue(skip_bulk({}))
        self.assertTrue(skip_bulk({'skip_bulk': '1'}))
        self.assertFalse(skip_bulk({'skip_bulk': ''}))
        self.assertFalse(skip_bulk({'skip_bulk': '0'}))

    def test_a_quiet_cycle_says_nothing(self):
        self.assertEqual(skipped_text(0), '')
        self.assertIn('3건', skipped_text(3))


class SkippedMailTests(unittest.TestCase):
    """건너뜀은 삭제가 아니다 — 목록에 남고, 왜인지 말하고, 다시 분석이 되돌린다.

    조용히 사라지는 필터가 이 프로젝트가 가장 싫어하는 실패라서, 아래 넷은 전부
    '내려놓은 뒤에도 보이는가'를 묻는다.
    """

    ACCOUNT = 'pop3s.hiworks.com:995/me@corp.example'

    def letter(self):
        message = EmailMessage()
        message['From'] = '소식지 <news@letters.example>'
        message['Subject'] = '9월 뉴스레터'
        message['Date'] = 'Fri, 11 Sep 2026 10:00:00 +0900'
        message['List-Unsubscribe'] = '<https://letters.example/u>'
        message.set_content('이번 달 소식입니다.')
        return message.as_bytes()

    def test_the_headers_survive_the_round_trip_through_raw(self):
        """raw를 보관하므로 마이그레이션이 필요 없다 — 이미 수집된 메일도 그대로 읽힌다."""
        parsed = parse_mail(self.letter())
        self.assertEqual(skip_reason(parsed), LETTER_REASON)

    def test_a_skipped_mail_leaves_the_queue_but_not_the_list(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                ident = store.add(self.ACCOUNT, 'u1', self.letter())
                self.assertEqual(len(store.pending(self.ACCOUNT, time.time())), 1)
                store.skip(ident, LETTER_REASON)
                # 분석 차례에서는 빠지고,
                self.assertEqual(store.pending(self.ACCOUNT, time.time()), [])
                # 목록에는 사유를 달고 남는다.
                row = store.detail(ident)
                self.assertEqual(state_of(row), SKIPPED)
                self.assertEqual(row_view(row)['skipped'], LETTER_REASON)
            finally:
                store.db.close()

    def test_it_is_neither_waiting_nor_failed(self):
        """대기로 두면 영원히 차례를 기다리는 것처럼 보이고, 실패로 세면 분석 실패
        카드가 시도하지도 않은 메일을 센다. 둘 다 화면이 거짓말을 하는 쪽이다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                ident = store.add(self.ACCOUNT, 'u1', self.letter())
                store.skip(ident, LETTER_REASON)
                rows = list(store.page(self.ACCOUNT))
                self.assertEqual(failures(rows), 0)
                self.assertEqual(store.detail(ident)['attempts'], 0)
                waiting, _ = store.search(self.ACCOUNT, state='분석 대기')
                self.assertEqual(list(waiting), [])
                gone, _ = store.search(self.ACCOUNT, state=SKIPPED)
                self.assertEqual([row['id'] for row in gone], [ident])
                # 목록이 읽는 컬럼으로도 건너뜀이라고 말해야 한다. SQL 필터만 맞고
                # state_of()가 '분석 대기'라고 답하는 상태가 실제로 있었다 —
                # LIST_COLUMNS 에 skipped 가 없어서, DB에는 사유가 있는데 화면만
                # 틀렸고 그 사이에도 이 테스트는 통과하고 있었다.
                self.assertEqual(state_of(gone[0]), SKIPPED)
                self.assertEqual(row_view(gone[0])['skipped'], LETTER_REASON)
                listed = list(store.page(self.ACCOUNT))
                self.assertEqual([state_of(row) for row in listed if row['id'] == ident],
                                 [SKIPPED])
            finally:
                store.db.close()

    def test_reanalyse_is_the_way_back_in(self):
        """잘못 걸렀다 싶을 때 되돌리는 길이 하나여야 하고, 그것은 이미 있는 길이다."""
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'mail.db')
            try:
                ident = store.add(self.ACCOUNT, 'u1', self.letter())
                store.skip(ident, LETTER_REASON)
                self.assertEqual(store.reset([ident], reanalyze=True), 1)
                self.assertEqual(store.detail(ident)['skipped'], '')
                self.assertEqual(state_of(store.detail(ident)), '분석 대기')
                self.assertEqual(len(store.pending(self.ACCOUNT, time.time())), 1)
            finally:
                store.db.close()


if __name__ == '__main__':
    unittest.main()
