"""도우미가 만드는 엑셀과 같은 모양의 예시 파일을 만듭니다. Excel 없이 동작합니다.

    MailAssistantTools.exe sample              바탕화면(없으면 현재 폴더)에 '메일 업무관리 예시.xlsx'
    MailAssistantTools.exe sample 경로.xlsx    위치를 직접 지정

메일 수신과 Codex 분석을 하지 않고, 샘플 분석 결과만 실제 코드 경로로 통과시킵니다.
"""
import json
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.hyperlink import Hyperlink
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:  # Bundled in the release; only a source checkout can miss it.
    raise SystemExit('예시 파일 생성에는 openpyxl이 필요합니다: python -m pip install openpyxl')

from mail_assistant.calendar_sheet import MARKERS, collect, month_grid, next_month, overdue
from mail_assistant.console import use_utf8
from mail_assistant.core import HEADERS, workbook_rows
from mail_assistant.excel import mailto
from mail_assistant.dashboard import (SIZES, UPCOMING_ROWS, UPCOMING_TOP, WIDTHS, blocks, cards,
                                      describe, upcoming)
from mail_assistant.style import (CALM, DASH_BG, DASH_CARD, DASH_LINE, DASH_LINK, DASH_MUTED,
                                  DASH_RED, DASH_TEXT, FONT, HEADER_FILL, LINE, LINK, RULES,
                                  SOON, TODAY_FILL, URGENT)

TODAY = date.today()
TABLE_NAMES = {'메일 목록': 'MailList', '일정': 'Schedule', '우선순위': 'Priority', '답변 초안': 'Reply'}


def rgb(bgr):
    """style.py keeps COM's BGR order; openpyxl wants RGB hex."""
    return f'{bgr & 0xFF:02X}{(bgr >> 8) & 0xFF:02X}{(bgr >> 16) & 0xFF:02X}'


def day(offset, clock=''):
    target = TODAY + timedelta(days=offset)
    return f'{target.isoformat()}T{clock}:00+09:00' if clock else target.isoformat()


def sample_rows():
    """수신 → 분석 결과 형태 그대로. 실제 반영과 같은 함수로 행을 만듭니다."""
    received = datetime.now(timezone.utc) - timedelta(hours=3)
    mails = [
        (dict(sender='대성정밀 김철수 <kim@daesung.example.com>', subject='[견적 요청] G-Tech 자동화 설비 2차 견적 회신 요청',
              attachments=['컨베이어_사양서_REV_C.pdf']),
         dict(category='견적·계약', summary='컨베이어 라인 3호기 개조 2차 견적을 요청했습니다. 사양서 3페이지 공차 항목 확인이 필요합니다.',
              requests='9월 16일까지 2차 견적 회신, 사양 확정 회의 참석',
              events=[dict(title='사양 확정 회의', start=day(3, '10:00'), deadline='', evidence='9월 14일(월) 오전 10시, 당사 회의실', needs_review=False),
                      dict(title='견적 회신 마감', start='', deadline=day(5, '18:00'), evidence='9월 16일(수) 오후 6시까지', needs_review=False)],
              priority='높음', priority_reason='회신 기한이 명시되어 있고 계약 진행에 직접 영향', next_action='공차 항목 검토 후 견적 산출',
              reply_needed=True, reply_subject='Re: [견적 요청] G-Tech 자동화 설비 2차 견적 회신 요청',
              reply_draft='안녕하세요, G-Tech 시스템입니다.\n\n요청하신 컨베이어 라인 3호기 개조 2차 견적은 9월 16일까지 회신드리겠습니다.\n사양서 3페이지 공차 항목은 검토 후 회의에서 확인하겠습니다.\n\n감사합니다.')),
        (dict(sender='총무팀 <admin@example.com>', subject='[공지] 9월 사내 네트워크 정기 점검 안내', attachments=[]),
         dict(category='공지', summary='9월 19일 새벽 네트워크 정기 점검으로 그룹웨어와 VPN이 중단됩니다.', requests='',
              events=[dict(title='네트워크 정기 점검', start=day(8, '02:00'), deadline='', evidence='9월 19일(토) 02:00 ~ 05:00', needs_review=False)],
              priority='낮음', priority_reason='업무 시간 외 작업이며 별도 조치 불필요', next_action='점검 시간 중 원격 접속 자제',
              reply_needed=False, reply_subject='', reply_draft='')),
        (dict(sender='박영희 <park@partner.example.com>', subject='다음 주에 잠깐 미팅 가능할까요?', attachments=[]),
         dict(category='회의·일정', summary='유지보수 계약 갱신 건으로 다음 주 미팅을 요청했습니다. 날짜가 확정되지 않았습니다.',
              requests='미팅 일정 회신',
              events=[dict(title='유지보수 계약 갱신 미팅', start=day(4), deadline='', evidence='다음 주 화요일쯤 오후', needs_review=True)],
              priority='보통', priority_reason='기한은 이달 말로 여유가 있으나 일정 조율 필요', next_action='가능한 시간 2개 제시',
              reply_needed=True, reply_subject='Re: 다음 주에 잠깐 미팅 가능할까요?',
              reply_draft='안녕하세요. 다음 주 화요일 오후 2시 또는 수요일 오후 3시에 방문 가능합니다.\n편하신 시간 알려주시면 일정 확정하겠습니다.')),
        (dict(sender='승인팀 <approval@daesung.example.com>', subject='RE: RE: [협의] 도면 승인 및 착수 일정', attachments=[]),
         dict(category='업무 요청', summary='도면 승인이 완료되었고 착수 보고서 제출만 남았습니다.', requests='9월 18일까지 착수 보고서 제출',
              events=[dict(title='착수 보고서 제출', start='', deadline=day(7), evidence='9월 18일(금)까지 착수 보고서', needs_review=False),
                      dict(title='도면 회신(지난 건)', start='', deadline=(TODAY - timedelta(days=17)).isoformat(), evidence='인용된 8월 메일의 회신 기한', needs_review=False)],
              priority='높음', priority_reason='승인 완료 후 착수 일정이 확정됨', next_action='착수 보고서 작성',
              reply_needed=True, reply_subject='Re: [협의] 도면 승인 및 착수 일정',
              reply_draft='승인 확인했습니다. 착수 보고서는 9월 18일까지 보내드리겠습니다.')),
        (dict(sender='이용문의 <user@example.com>', subject='[문의] 서비스 이용 관련 질문', attachments=[]),
         dict(category='문의', summary='월 이용료 결제일 변경이 가능한지 문의했습니다.', requests='결제일 변경 가능 여부 회신',
              events=[dict(title='결제일 변경 검토', start='', deadline=day(21), evidence='다음 결제일 전까지', needs_review=False)],
              priority='보통', priority_reason='단순 문의이나 결제일 전 회신 필요', next_action='결제 정책 확인 후 회신',
              reply_needed=True, reply_subject='Re: [문의] 서비스 이용 관련 질문',
              reply_draft='문의 주셔서 감사합니다. 결제일 변경은 다음 결제일 7일 전까지 요청하시면 가능합니다.')),
    ]
    rows = []
    for index, (parsed, result) in enumerate(mails):
        parsed = {**parsed, 'date': '', 'message_id': '', 'body': ''}
        rows.append({'id': f'sample-{index + 1:02d}', 'received': (received + timedelta(minutes=index * 7)).isoformat(),
                     'parsed': json.dumps(parsed, ensure_ascii=False), 'result': json.dumps(result, ensure_ascii=False)})
    return rows


def write_sheets(book, rows):
    sheets = {}
    for name, headers in HEADERS.items():
        sheet = book.create_sheet(name)
        sheet.append(list(headers))
        sheets[name] = sheet
    for row in rows:
        for name, values in workbook_rows(row).items():
            for values_row in values:
                sheets[name].append([str(value or '') for value in values_row])
    index = {str(row[0].value): row[0].row for row in sheets['메일 목록'].iter_rows(min_row=2)}
    for name in ('일정', '우선순위', '답변 초안'):
        column = HEADERS[name].index('메일 ID') + 1
        for row in sheets[name].iter_rows(min_row=2):
            cell = row[column - 1]
            target = index.get(str(cell.value))
            if target:
                # A plain string becomes an external link; internal jumps need `location`.
                cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'메일 목록'!A{target}",
                                           tooltip='메일 목록에서 이 메일 보기')
                cell.font = Font(name=FONT, size=10, color=rgb(LINK), underline='single')
    column = HEADERS['답변 초안'].index('답변 제목') + 1
    for row in sheets['답변 초안'].iter_rows(min_row=2):
        link = mailto(row[1].value, row[2].value, row[3].value)
        if link:
            cell = row[column - 1]
            cell.hyperlink = Hyperlink(ref=cell.coordinate, target=link,
                                       tooltip='기본 메일 앱에서 초안이 채워진 답장 창 열기')
            cell.font = Font(name=FONT, size=10, color=rgb(LINK), underline='single')
    # 사용자가 직접 관리하는 열의 예시.
    sheets['메일 목록']['I2'] = '처리'
    sheets['답변 초안']['E2'] = '금액은 검토 후 별도 안내로 수정'
    sheets['답변 초안']['F2'] = '검토 중'
    for key, value in {'마지막 메일 확인': datetime.now(timezone.utc).isoformat(), '이번 엑셀 반영': datetime.now(timezone.utc).isoformat(),
                       '분석 대기': 0, '엑셀 반영 대기(이번 반영 전)': 0,
                       '안내': f'예시 파일입니다. 메일 {len(rows)}건을 반영했습니다.'}.items():
        sheets['실행 상태'].append([key, str(value)])
    return sheets


def style_sheets(sheets):
    header_fill = PatternFill('solid', fgColor=rgb(HEADER_FILL))
    for name, sheet in sheets.items():
        for row in sheet.iter_rows():
            for cell in row:
                if cell.hyperlink is not None:
                    continue  # keep the link styling applied above
                cell.font = Font(name=FONT, size=10)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(name=FONT, size=10, bold=True)
            cell.fill = header_fill
        if name == '실행 상태':
            sheet.column_dimensions['A'].width = 26
            sheet.column_dimensions['B'].width = 62
            for cell in sheet['A']:
                cell.font = Font(name=FONT, size=10, bold=True)
            continue
        width = len(HEADERS[name])
        sheet.freeze_panes = 'A2'
        sheet.column_dimensions['A'].width = 27
        for column in range(2, width + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 24 if column < 5 else 48
        last = max(sheet.max_row, 2)
        table = Table(displayName=TABLE_NAMES[name], ref=f'A1:{get_column_letter(width)}{last}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleLight8', showRowStripes=True)
        sheet.add_table(table)
        for address, formula, color, bold in RULES.get(name, []):
            sheet.conditional_formatting.add(
                address, FormulaRule(formula=[formula.lstrip('=')], font=Font(color=rgb(color), bold=bold)))


def write_calendar(book, sheets):
    events = collect([[cell.value for cell in row] for row in sheets['일정'].iter_rows(min_row=2)])
    sheet = book.create_sheet('일정 달력', 1)
    border = Side('thin', color=rgb(LINE))
    for top, month in ((1, TODAY), (10, next_month(TODAY))):
        sheet.cell(top, 1, f'{month.year}년 {month.month}월').font = Font(name=FONT, size=13, bold=True)
        for index, label in enumerate(('월', '화', '수', '목', '금', '토', '일'), 1):
            cell = sheet.cell(top + 1, index, label)
            cell.font = Font(name=FONT, size=10, bold=True)
            cell.fill = PatternFill('solid', fgColor=rgb(HEADER_FILL))
            cell.alignment = Alignment(horizontal='center')
        text, cells = month_grid(month.year, month.month, events)
        for week in range(6):
            sheet.row_dimensions[top + 2 + week].height = 76
            for column in range(7):
                cell = sheet.cell(top + 2 + week, column + 1)
                cell.border = Border(left=border, right=border, top=border, bottom=border)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(name=FONT, size=9)
                entry = cells.get((week, column))
                if not entry:
                    continue
                current, spans = entry
                lines = text[week][column].split('\n')
                head = InlineFont(rFont=FONT, sz=9, b=current == TODAY,
                                  color=rgb(CALM) if column >= 5 else '000000')
                blocks = [TextBlock(head, lines[0])]
                for line, span in zip(lines[1:], spans):
                    # A run holding only '\n' is written without xml:space="preserve",
                    # and Excel repairs the file. Keep the newline with its line.
                    blocks.append(TextBlock(InlineFont(rFont=FONT, sz=9, color=rgb(span[2])), '\n' + line))
                for line in lines[1 + len(spans):]:
                    blocks.append(TextBlock(InlineFont(rFont=FONT, sz=9, color=rgb(CALM)), '\n' + line))
                cell.value = CellRichText(blocks)
                if current == TODAY:
                    cell.fill = PatternFill('solid', fgColor=rgb(TODAY_FILL))
    legend = sheet.cell(19, 1)
    legend.value = CellRichText([
        TextBlock(InlineFont(rFont=FONT, sz=10, color=rgb(URGENT)), f'{MARKERS["마감"]} 마감    '),
        TextBlock(InlineFont(rFont=FONT, sz=10, color=rgb(SOON)), f'{MARKERS["시작"]} 시작    '),
        TextBlock(InlineFont(rFont=FONT, sz=10, color=rgb(CALM)), f'{MARKERS["확인 필요"]} 확인 필요'),
        TextBlock(InlineFont(rFont=FONT, sz=10), '        일정 시트를 매 반영마다 다시 그립니다.')])
    past = overdue(events, TODAY)
    if past:
        heading = sheet.cell(21, 1, f'지난 마감 {len(past)}건 (달력 범위 밖)')
        heading.font = Font(name=FONT, size=10, bold=True)
        for index, (when, label) in enumerate(past, 22):
            sheet.cell(index, 1, f'{when.isoformat()}  {label}').font = Font(name=FONT, size=10, color=rgb(CALM))
    for column in range(1, 8):
        sheet.column_dimensions[get_column_letter(column)].width = 21


def write_dashboard(book, sheets):
    """Same layout as the live sheet, drawn with openpyxl instead of COM."""
    sheet = book.create_sheet('대시보드', 0)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = rgb(DASH_BG)
    fill = PatternFill('solid', fgColor=rgb(DASH_BG))
    card_fill = PatternFill('solid', fgColor=rgb(DASH_CARD))
    for row in range(1, 34):
        for column in range(1, 16):
            sheet.cell(row, column).fill = fill
    for index, width in WIDTHS.items():
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row, height in ((1, 10), (2, 30), (3, 18), (4, 8), (5, 16), (6, 34), (7, 14), (8, 10)):
        sheet.row_dimensions[row].height = height
    for row, first, width, kind, value, color in blocks(TODAY):
        cell = sheet.cell(row, first, value)
        size, bold = SIZES[kind]
        cell.font = Font(name=FONT, size=size, bold=bold, color=rgb(color))
        cell.alignment = Alignment(vertical='center', wrap_text=kind in ('subtitle', 'note'))
        if width > 1:
            sheet.merge_cells(start_row=row, start_column=first, end_row=row, end_column=first + width - 1)
        if kind in ('card_label', 'card_value', 'card_note'):
            for column in range(first, first + width):
                sheet.cell(row, column).fill = card_fill
    side = Side('thin', color=rgb(DASH_LINE))
    for index in range(len(cards())):
        left = 2 + index * 3
        for row in range(5, 8):
            for column in range(left, left + 3):
                cell = sheet.cell(row, column)
                cell.fill = card_fill
                cell.border = Border(left=side if column == left else None,
                                     right=side if column == left + 2 else None,
                                     top=side if row == 5 else None,
                                     bottom=side if row == 7 else None)
    events = collect([[cell.value for cell in row] for row in sheets['일정'].iter_rows(min_row=2)])
    for row in range(9, UPCOMING_TOP + UPCOMING_ROWS + 1):
        for column in range(8, 14):
            sheet.cell(row, column).fill = card_fill
    mail_rows = {str(row[0].value): row[0].row for row in sheets['메일 목록'].iter_rows(min_row=2)}
    for index, (when, entry) in enumerate(upcoming(events, TODAY)):
        row = UPCOMING_TOP + index
        sheet.cell(row, 8, when.isoformat()).font = Font(name=FONT, size=10, color=rgb(DASH_MUTED))
        cell = sheet.cell(row, 9, entry.label)
        target = mail_rows.get(entry.mail_id)
        location = f"'메일 목록'!A{target}" if target else f"'일정'!A{entry.row}"
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=location,
                                   tooltip='메일 목록에서 이 메일 보기' if target else '일정 시트에서 이 줄 보기')
        cell.font = Font(name=FONT, size=10, color=rgb(DASH_LINK), underline='single')
        urgent = (when - TODAY).days <= 3
        sheet.cell(row, 12, describe(when, TODAY)).font = Font(
            name=FONT, size=10, color=rgb(DASH_RED if urgent else DASH_MUTED))


def main():
    use_utf8()
    if len(sys.argv) > 1:
        path = Path(sys.argv[1])
    else:
        desktop = Path.home() / 'Desktop'
        path = (desktop if desktop.is_dir() else Path.cwd()) / '메일 업무관리 예시.xlsx'
    book = Workbook()
    book.remove(book.active)
    rows = sample_rows()
    sheets = write_sheets(book, rows)
    style_sheets(sheets)
    write_calendar(book, sheets)
    write_dashboard(book, sheets)
    book.save(path)
    print(f'예시 파일 생성: {path}')
    print(f'메일 {len(rows)}건, 시트 {len(book.sheetnames)}개: {", ".join(book.sheetnames)}')


if __name__ == '__main__':
    main()
